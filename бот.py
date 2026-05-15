import aiohttp
import asyncio
import logging
import time
import os
import random
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.image import MIMEImage
import requests
from aiogram import Bot, Dispatcher, types
from aiogram.contrib.middlewares.logging import LoggingMiddleware
from aiogram.dispatcher import FSMContext
from aiogram.dispatcher.filters.state import StatesGroup, State
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, ReplyKeyboardMarkup, KeyboardButton
from aiogram.utils import executor
from aiogram.contrib.fsm_storage.memory import MemoryStorage
from telethon import TelegramClient, errors
from aiogram.dispatcher import FSMContext
from telethon.tl.functions.channels import JoinChannelRequest
from datetime import datetime, timedelta
import re

##
from proxies import proxies
from user_agents import user_agents
from emails import mail, phone_numbers
###$
from states import ChannelDemolitionStates, option_mapping, reason_mapping, ReportStates, EmailTemplateStates, ComplaintStates, CreateAccountStates, RestoreAccountStates
####
from text import PAID_WELCOME_BANNER, WELCOME_BANNER, email_templates
from config import CHANNEL_ID, CHANNELS, bot_token, admin_chat_ids, CRYPTO_PAY_TOKEN, receivers, smtp_servers, clients, VIP_PRICE, CURRENCY_PRICES, private_channel_ids, SUBSCRIPTION_CHANNEL_ID, dp, bot, banned_users, private_channel_ids, banned_users_file, session_dir, script_dir, API_ID, API_HASH
from session import init_sessions

VK_TOKEN = "0af157510af157510af15751aa0a89e69600af10af157516a0bc15996e74fe2b440998c"
MAX_CAPTION_LENGTH = 1024
##

import inspect
from importlib import import_module
from config import dp, bot
from aiogram import executor

class DoxState(StatesGroup):
    waiting_for_input = State()

class BanState(StatesGroup):
    waiting_for_ban_user_id = State()
    waiting_for_unban_user_id = State()
    waiting_for_ban_confirmation = State()
    waiting_for_unban_confirmation = State()           
#------
import репорты
import вк
import session 
import se_me
for name, handler in inspect.getmembers(вк, inspect.iscoroutinefunction):
    dp.register_callback_query_handler(handler, state='*')
import vip
for name, handler in inspect.getmembers(vip, inspect.iscoroutinefunction):
    dp.register_callback_query_handler(handler, state='*')    
for name, handler in inspect.getmembers(session, inspect.iscoroutinefunction):
    dp.register_callback_query_handler(handler, state='*')    
for name, handler in inspect.getmembers(se_me, inspect.iscoroutinefunction):
    if name.startswith('process_') and 'message' in name:
        dp.register_message_handler(handler, state='*')
    elif name.startswith('callback_'):
        dp.register_callback_query_handler(handler, state='*')    
for name, handler in inspect.getmembers(репорты, inspect.iscoroutinefunction):
    if name.startswith('process_') and 'message' in name:
        dp.register_message_handler(handler, state='*')
    elif name.startswith('callback_'):
        dp.register_callback_query_handler(handler, state='*')
#------



def read_private_users():
    try:
        with open("База_Бота/private_users.txt", "r", encoding="utf-8") as file:
            lines = [line.strip() for line in file.readlines() if line.strip()]
    except FileNotFoundError:
        return {"ids": [], "usernames": []}
    
    ids = []
    usernames = []
    
    if lines:
        ids = [int(id_str) for id_str in lines[0].split(',') if id_str.strip().isdigit()]
    if len(lines) > 1:
        usernames = [uname.lstrip('@') for uname in lines[1].split(',') if uname.strip()]
    
    return {"ids": ids, "usernames": usernames}
        
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

class VkStates(StatesGroup):
    waiting_for_vk_query = State()

if not os.path.exists(session_dir):
    os.makedirs(session_dir)

class ChannelDemolitionStates(StatesGroup):
    channel_link = State()
    option = State()    

class SubStates(StatesGroup):
    add_sub_date = State()
    remove_sub = State()
    change_sub_date = State()

class SupportStates(StatesGroup):
    message = State()

def register_handlers_spam_code(dp: Dispatcher):
    dp.register_message_handler(process_spam_code, state=SpamCodeStates.phone_and_count)

class SendMessage(StatesGroup):
    text = State()
    media_type = State()
    media = State()

def add_user_to_file(user_id: int):
    try:
        users_file = os.path.join('База_Бота', 'users.txt')
        with open(users_file, 'r') as file:
            users = file.readlines()
        users = [line.strip() for line in users if line.strip()]
        user_ids = [line.split()[0] for line in users if line.split()]
        if str(user_id) not in user_ids:
            with open(users_file, 'a') as file:
                file.write(f"{user_id}\n")
    except Exception as e:
        print(f"Ошибка при добавлении пользователя в файл: {e}")

async def add_to_tracking_list(user_id, target_user_id):
    tracking_list = load_tracking_list()
    if user_id not in tracking_list:
        tracking_list[user_id] = []
    if target_user_id not in tracking_list[user_id]:
        tracking_list[user_id].append(target_user_id)
        save_tracking_list(tracking_list)


def save_tracking_list(tracking_list):
    tracking_file = os.path.join('База_Бота', 'tracking_list.txt')
    with open(tracking_file, 'w') as file:
        for user_id, target_user_ids in tracking_list.items():
            file.write(f"{user_id}:{','.join(map(str, target_user_ids))}\n")


def load_tracking_list():
    try:
        tracking_file = os.path.join('База_Бота', 'tracking_list.txt')
        with open(tracking_file, 'r') as file:
            tracking_list = {}
            for line in file:
                user_id, target_user_ids = line.strip().split(':')
                tracking_list[int(user_id)] = [int(uid) for uid in target_user_ids.split(',')]
            return tracking_list
    except FileNotFoundError:
        with open(tracking_file, 'w') as file:
            pass
        return {}
    except (ValueError, PermissionError, IsADirectoryError) as e:
        print(f"Error loading tracking list: {e}")
        return {}


async def notify_users_about_status():
    tracking_list = load_tracking_list()
    for user_id, target_user_ids in tracking_list.items():
        for target_user_id in target_user_ids:
            status, _ = await check_account_status(target_user_id)
            if status is False:
                await bot.send_message(user_id, f"✅ Аккаунт {target_user_id} был успешно удален.")
                tracking_list[user_id].remove(target_user_id)
                if not tracking_list[user_id]:
                    del tracking_list[user_id]
    save_tracking_list(tracking_list)


async def background_status_checker():
    while True:
        await notify_users_about_status()
        await asyncio.sleep(3600)

#----'z'

import aiohttp
import asyncio
import logging
import time
import os
import random
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.image import MIMEImage
import requests
from aiogram import Bot, Dispatcher, types
from aiogram.contrib.middlewares.logging import LoggingMiddleware
from aiogram.dispatcher import FSMContext
from aiogram.dispatcher.filters.state import State, StatesGroup
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, ReplyKeyboardMarkup, KeyboardButton
from aiogram.utils import executor
from aiogram.contrib.fsm_storage.memory import MemoryStorage
from telethon import TelegramClient, errors
from telethon.tl.functions.channels import JoinChannelRequest
from telethon.tl.functions.messages import ReportRequest
from telethon.tl.types import (
    InputReportReasonSpam,
    InputReportReasonViolence,
    InputReportReasonChildAbuse,
    InputReportReasonPornography,
    InputReportReasonCopyright,
    InputReportReasonPersonalDetails,
    InputReportReasonGeoIrrelevant,
    InputReportReasonFake,
    InputReportReasonIllegalDrugs,
    InputReportReasonOther
)
from telethon.tl.functions.channels import JoinChannelRequest
from datetime import datetime, timedelta
import re
from text import PAID_WELCOME_BANNER, WELCOME_BANNER, email_templates
from config import CHANNEL_ID, CHANNELS, bot_token, admin_chat_ids, CRYPTO_PAY_TOKEN, receivers, smtp_servers, clients, VIP_PRICE, CURRENCY_PRICES, private_channel_ids, SUBSCRIPTION_CHANNEL_ID, dp, bot

from telethon.tl.functions.channels import GetFullChannelRequest
from aiogram import types
from telethon.sessions import StringSession
from telethon import TelegramClient, errors
from telethon.tl.functions.channels import JoinChannelRequest, GetFullChannelRequest
from telethon.sessions import SQLiteSession
from aiogram import types
from aiogram.dispatcher import FSMContext
import asyncio
import sqlite3
from telethon import TelegramClient
from aiogram.dispatcher.filters.state import State, StatesGroup 

class ChannelDemolitionStates(StatesGroup):
    channel_link = State()
    option = State()   

from aiogram.dispatcher.filters.state import StatesGroup, State
from aiogram.dispatcher import FSMContext


    
class BanState(StatesGroup):
    waiting_for_ban_user_id = State()
    waiting_for_unban_user_id = State()
    waiting_for_ban_confirmation = State()
    waiting_for_unban_confirmation = State()     
    
class DoxState(StatesGroup):
    waiting_for_input = State()     

from config import dp
from aiogram import types

@dp.message_handler(commands=['start'], state='*')
async def cmd_start(message: types.Message, state: FSMContext):
    await state.finish()
    user_id = message.from_user.id
    user_id_str = str(user_id)
    not_subscribed_channels = await check_all_subscriptions(user_id)

    referrer_id = None
    if len(message.text.split()) > 1:
        start_payload = message.text.split()[1]
        if start_payload.startswith('ref_'):
            parts = start_payload[4:].split('_', 1)
            if len(parts) == 2:
                referrer_id_from_link, ref_code = parts
                
                try:
                    with open(os.path.join('База_Бота', 'referral_data.json'), 'r', encoding='utf-8') as f:
                        referral_data = json.load(f)
                    
                    if (referrer_id_from_link in referral_data and 
                        referral_data[referrer_id_from_link]['referral_code'] == ref_code and 
                        referrer_id_from_link != user_id_str):
                        referrer_id = referrer_id_from_link
                except (FileNotFoundError, json.JSONDecodeError):
                    pass
    
    await handle_referral_data(user_id, (await bot.me).username, referrer_id)
    
    if not not_subscribed_channels:
        has_access = False
        
        try:
            with open(os.path.join('База_Бота', 'paid_users.txt'), 'r', encoding='utf-8') as f:
                for line in f:
                    line = line.strip()
                    if line.startswith(f"{user_id_str},"):
                        if line.endswith(",forever"):
                            has_access = True
                            break
                        else:
                            try:
                                expiry_date_str = line.split(',')[1]
                                expiry_date = datetime.strptime(expiry_date_str, '%Y-%m-%d %H:%M:%S')
                                if expiry_date > datetime.now():
                                    has_access = True
                                    break
                            except (IndexError, ValueError):
                                continue
        except FileNotFoundError:
            pass
        if not has_access:
            try:
                chat_member = await bot.get_chat_member(chat_id=SUBSCRIPTION_CHANNEL_ID, user_id=user_id)
                if chat_member.status in ['member', 'administrator', 'creator']:
                    expiry_date = datetime.now() + timedelta(days=30)
                    with open(os.path.join('База_Бота', 'paid_users.txt'), 'a', encoding='utf-8') as f:
                        f.write(f"{user_id_str},{expiry_date.strftime('%Y-%m-%d %H:%M:%S')}\n")
                    
                    username = message.from_user.username or ""
                    try:
                        with open(os.path.join('База_Бота', 'private_users.txt'), 'r+', encoding='utf-8') as f:
                            lines = [line.strip() for line in f.readlines() if line.strip()]
                            ids = set()
                            if lines:
                                ids = set(filter(None, lines[0].split(',')))
                            usernames = set()
                            if len(lines) > 1:
                                usernames = set(filter(None, lines[1].split(',')))
                            ids.add(user_id_str)
                            if username:
                                usernames.add(username)
                            f.seek(0)
                            f.write(','.join(ids) + '\n')
                            if usernames:
                                f.write(','.join(usernames) + '\n')
                            f.truncate()
                    except FileNotFoundError:
                        with open(os.path.join('База_Бота', 'private_users.txt'), 'w', encoding='utf-8') as f:
                            f.write(f"{user_id_str}\n")
                            if username:
                                f.write(f"{username}\n")
                    except Exception:
                        pass
            except Exception:
                pass

        await handle_welcome(
            user_id=user_id,
            chat_id=message.chat.id,
            from_user=message.from_user,
            reply_photo_method=message.reply_photo
        )
    else:
        await UserStates.waiting_for_subscription.set()
        keyboard = InlineKeyboardMarkup()
        for channel_name, channel_url in not_subscribed_channels.items():
            keyboard.add(InlineKeyboardButton(f"🔗 {channel_name}", url=channel_url))
        keyboard.add(InlineKeyboardButton("✅ Проверить подписку", callback_data="check_subscription"))
        await message.answer(
            "📢 <b>Для получения доступа требуется подписка!</b>\n\n"
            "Необходимо подписаться на:\n"
            + "\n".join([f"▪️ {channel}" for channel in not_subscribed_channels.keys()]),
            reply_markup=keyboard,
            parse_mode="HTML"
        )


async def send_menu(chat_id: int, welcome_message: str):
    markup = InlineKeyboardMarkup(row_width=2)
    btn_support = InlineKeyboardButton('📩 Написать поддержку📩', callback_data='support')
    btn_demolition = InlineKeyboardButton('☠️Snos☠️', callback_data='demolition')  
    btn_restore_account = InlineKeyboardButton('🔄Pеснуть акк🔄', callback_data='restore_account')
    btn_my_time = InlineKeyboardButton('👤Профиль👤', callback_data='my_time')
    btn_spam_menu = InlineKeyboardButton('🔥Spam🔥', callback_data='spam_menu')
    btn_osint = InlineKeyboardButton('🔍Osint🔎', callback_data='osint')
    btn_instruction = InlineKeyboardButton('📚 Инструкция', callback_data='instruction') 
    
    markup.add(btn_demolition, btn_osint)
    markup.add(btn_spam_menu)
    markup.add(btn_support, btn_restore_account)
    markup.add(btn_my_time, btn_instruction)
    
    if str(chat_id) in admin_chat_ids:
        btn_admin_panel = InlineKeyboardButton('🛠 Админ панель', callback_data='admin_panel')
        markup.add(btn_admin_panel)
    photo_path = os.path.join('Фото', 'welcome_photo.jpg')
    
    await bot.send_photo(
        chat_id=chat_id,
        photo=open(photo_path, 'rb'),
        caption=welcome_message,
        reply_markup=markup,
        parse_mode="HTML"
    )


import zipfile
import rarfile


async def get_links_from_doxbot(target, message):
    sessions_di = "Не_основные"
    os.makedirs(sessions_di, exist_ok=True)
    status_msg = await message.answer("🔍 Идёт поиск доступных сессий...")
    
    sessions = []
    for file in os.listdir(sessions_di):
        if file.endswith('.session'):
            session_name = file.replace('.session', '')
            json_path = os.path.join(sessions_di, f"{session_name}.json")
            
            if not os.path.exists(json_path):
                with open(json_path, 'w') as f:
                    json.dump({"name": session_name, "status": "Свободна"}, f)
            
            sessions.append({
                "name": session_name,
                "path": os.path.join(sessions_di, file),
                "api_id": API_ID,
                "api_hash": API_HASH
            })
    
    if not sessions:
        await status_msg.edit_text("❌ Активные сессии не найдены!")
        return None
    for session in sessions:
        session_name = session["name"]
        json_path = os.path.join(sessions_di, f"{session_name}.json")
        
        try:
            with open(json_path, 'r') as f:
                session_data = json.load(f)
            
            if session_data.get("status") == "Свободна":
                temp_client = None
                try:
                    temp_client = TelegramClient(
                        session=SQLiteSession(session["path"]),
                        api_id=session["api_id"],
                        api_hash=session["api_hash"]
                    )
                    await temp_client.connect()
                    
                    if not await temp_client.is_user_authorized():
                        continue
                        
                except Exception as e:
                    continue
                finally:
                    if temp_client:
                        await temp_client.disconnect()
                
                session_data["status"] = "Занята"
                with open(json_path, 'w') as f:
                    json.dump(session_data, f)        
                client = None
                try:
                    client = TelegramClient(
                        session=SQLiteSession(session["path"]),
                        api_id=session["api_id"],
                        api_hash=session["api_hash"]
                    )
                    
                    await client.connect()
                    await status_msg.edit_text("🔍 Начинаю поиск сообщений...")
                    
                    class DoxState:
                        def __init__(self):
                            self.last_message_id = 0
                            self.collected_data = []
                            self.buttons_sequence = [["Сообщения"], ["Messages"], ["Все"], ["AII"], ["All"]]
                            self.current_stage = 0
                            self.done = False
                            self.max_retries = 3
                            self.retry_count = 0
                            self.all_links = set()
                            self.found_links = []
                            self.data_hidden = False
                            self.max_links = 200
                            self.current_msg = None

                    state = DoxState()

                    async def process_message(msg):
                        if not msg or state.done:
                            return False
                            
                        if msg.out or (msg.text and ("⏳ Bычисляю..." in msg.text or 
                                                    "Этoт юзернeйм упоминаeтcя в oпиcании чaтoв/каналов" in msg.text or 
                                                    "Этoт юзеpнейм испoльзoвaлcя в пpoшлом" in msg.text or
                                                    "Данныe cкpыты пoльзoвaтелем" in msg.text)):
                            if "Данныe cкpыты пoльзoвaтелем" in msg.text:
                                state.data_hidden = True
                                state.done = True
                            return False
                            
                        state.last_message_id = msg.id
                        state.current_msg = msg
                        
                        if msg.text:
                            
                            links = re.findall(
                                r'(?:https?://)?(?:t\.me/|telegram\.me/|telegram\.dog/)(?:[^/\s]+/\d+|c/\d+/\d+|s/[^/\s]+/\d+)', 
                                msg.text
                            )
                            
                            normalized_links = []
                            for link in links:
                                if not link.startswith('http'):
                                    link = 'https://' + link
                                link = link.replace('telegram.me/', 't.me/').replace('telegram.dog/', 't.me/')
                                
                                
                                if any(ignore in link for ignore in ['/c/', '/s/', '-100']):
                                    continue  
                                if re.match(r'https://t\.me/[^/\s]+/\d+$', link):
                                    normalized_links.append(link)
                            
                            if normalized_links:
                                new_links = []
                                for link in normalized_links:
                                    if link not in state.all_links and len(state.found_links) < state.max_links:
                                        state.all_links.add(link)
                                        state.found_links.append(link)
                                        new_links.append(link)
                                
                                if new_links:
                                    await status_msg.edit_text(
                                        f"🔍 Найдено доступных сообщений: {len(state.found_links)}\n"
                                    )
                                    
                                    if len(state.found_links) >= state.max_links:
                                        state.done = True
                                        return True
                            
                            if msg.text not in state.collected_data:
                                state.collected_data.append(msg.text)
                                return True
                            
                        return False

                    async def click_button(msg, button_texts):
                        if not hasattr(msg, 'reply_markup') or not msg.reply_markup:
                            return False, msg
                            
                        for i, row in enumerate(msg.reply_markup.rows):
                            for j, button in enumerate(row.buttons):
                                if any(kw.lower() in button.text.lower() for kw in button_texts):
                                    try:
                                        await msg.click(i, j)
                                        await asyncio.sleep(2)
                                        new_msg = await client.get_messages("@Dogizenrobot", ids=msg.id)
                                        if new_msg:
                                            return True, new_msg
                                    except Exception as e:
                                        print(f"Button click error: {e}")
                                    return False, msg
                        return False, msg

                    async def process_next_button(msg):
                        while not state.done:
                            clicked, new_msg = await click_button(msg, ["➡️", "›", ">", "далее", "next"])
                            if not clicked:
                                break
                                
                            if await process_message(new_msg):
                                msg = new_msg
                            else:
                                break
                        return msg

                    async def process_main_buttons(msg):
                        if state.current_stage < len(state.buttons_sequence):
                            buttons = state.buttons_sequence[state.current_stage]
                            clicked, new_msg = await click_button(msg, buttons)
                            
                            if clicked:
                                if await process_message(new_msg):
                                    new_msg = await process_next_button(new_msg)
                                    state.current_stage += 1
                                    return new_msg
                            else:
                                if state.retry_count < state.max_retries:
                                    state.retry_count += 1
                                    await asyncio.sleep(1)
                                    return await process_main_buttons(msg)
                                else:
                                    state.current_stage += 1
                                    return msg
                        else:
                            state.done = True
                            return None
                    
                    await client.send_message("@Dogizenrobot", target)
                    await asyncio.sleep(3)  
                    
                    async for msg in client.iter_messages("@Dogizenrobot", limit=1):
                        if await process_message(msg):
                            current_msg = msg
                            while not state.done and current_msg and not state.data_hidden:
                                current_msg = await process_main_buttons(current_msg)
                                await asyncio.sleep(1)

                    if state.data_hidden:
                        await status_msg.edit_text("❌ Не удалось извлечь ссылки: данные скрыты пользователем.")
                        await client.disconnect()
                        return None
                    elif state.found_links:
                        await status_msg.edit_text(
                            f"✅ Поиск завершен! Найдено {len(state.found_links)} сообщений\n"
                        )
                        await client.disconnect()
                        return {target: {"links": state.found_links[:state.max_links]}}
                    else:
                        await status_msg.edit_text("❌ Не удалось найти ссылки на сообщения.")
                        await client.disconnect()
                        return None

                except sqlite3.OperationalError as e:
                    if "database is locked" in str(e):
                        await status_msg.edit_text("❌ Ошибка: База данных заблокирована. Попробуйте позже.")
                    else:
                        await status_msg.edit_text(f"❌ Ошибка базы данных: {str(e)}")
                    if client:
                        await client.disconnect()
                    return None
                except Exception as e:
                    await status_msg.edit_text(f"❌ Ошибка при обработке запроса: {str(e)}")
                    if client:
                        await client.disconnect()
                    return None
                finally:
                    if client:
                        await client.disconnect()
                    session_data["status"] = "Свободна"
                    with open(json_path, 'w') as f:
                        json.dump(session_data, f)
                        
        except Exception as e:
            await status_msg.edit_text(f"⚠️ Ошибка проверки сессии {session_name}: {str(e)}")
            continue
    
    await status_msg.edit_text("❌ Нет свободных сессий в данный момент")
    return None    
@dp.message_handler(state=ReportStates.message_link)
async def process_message_link_step(message: types.Message, state: FSMContext):
    user_id = message.from_user.id
    target = message.text.strip()
    result_links = {}
    if '\n' in target:
        targets = [t.strip() for t in target.split('\n') if t.strip()]
    else:
        targets = [t.strip() for t in target.split() if t.strip()]
    
    all_links = True
    for t in targets:
        if not re.match(r'^https://t\.me/([^/]+/\d+|c/\d+/\d+|s/[^/]+/\d+)$', t):
            all_links = False
            break
    
    if all_links:
        for link in targets:
            username = link.split('/')[3] 
            if username not in result_links:
                result_links[username] = {"links": []}
            result_links[username]["links"].append(link)
    else:
        for t in targets:
            links_data = await get_links_from_doxbot(t, message)
            if links_data:
                for key, value in links_data.items():
                    if key not in result_links:
                        result_links[key] = {"links": []}
                    result_links[key]["links"].extend(value["links"][:200])
    
    if not result_links:
        await message.answer('❌ Не удалось получить ссылки. Введите ID, @username или ссылку')
        return
    
    async with state.proxy() as data:
        data['message_links'] = result_links            
    sessions_dir = "Не_основные"
    os.makedirs(sessions_dir, exist_ok=True)
    status_msg = await message.answer("🔍 Идёт поиск доступных сессий...")
    
    sessions = []
    for file in os.listdir(sessions_dir):
        if file.endswith('.session'):
            session_name = file.replace('.session', '')
            json_path = os.path.join(sessions_dir, f"{session_name}.json")
            
            if not os.path.exists(json_path):
                with open(json_path, 'w') as f:
                    json.dump({"name": session_name, "status": "Свободна"}, f)
            
            sessions.append({
                "name": session_name,
                "path": os.path.join(sessions_dir, file),
                "api_id": API_ID,
                "api_hash": API_HASH,
                "json_path": json_path
            })
    
    if not sessions:
        await status_msg.edit_text("❌ Активные сессии не найдены!")
        await state.finish()
        return
    
    await status_msg.edit_text(f"✅ Найдено сессий: {len(sessions)}. Проверяем доступность...")
    
    client = None
    selected_session = None
    for session in sessions:
        try:
            with open(session["json_path"], 'r') as f:
                session_data = json.load(f)
            
            if session_data.get("status") != "Свободна":
                continue
                
            temp_client = None
            try:
                temp_client = TelegramClient(
                    session=SQLiteSession(session["path"]),
                    api_id=session["api_id"],
                    api_hash=session["api_hash"],
                    system_version="4.16.30-vxCUSTOM"
                )
                await temp_client.connect()
                
                if not await temp_client.is_user_authorized():
                    continue
                session_data["status"] = "Занята"
                with open(session["json_path"], 'w') as f:
                    json.dump(session_data, f)
                
                selected_session = session
                client = temp_client
                temp_client = None  
                break
                
            except Exception as e:
                continue
            finally:
                if temp_client:
                    await temp_client.disconnect()
                    
        except Exception as e:
            continue
    
    if not client:
        await status_msg.edit_text("❌ Нет доступных сессий для подключения")
        await state.finish()
        return
    
    try:
        users_info = {}
        await status_msg.edit_text("<i>🔄 Обработка сообщений</i>", parse_mode="HTML")
        
        for username, link_data in result_links.items():
            if not link_data["links"]:
                continue
                
            message_link = link_data["links"][0]
            
            try:
                chat_id = None
                chat_title = None
                if 't.me/c/' in message_link:
                    chat_id = int('-100' + message_link.split('/')[-2])
                    message_id = int(message_link.split('/')[-1])
                    try:
                        chat = await client.get_entity(PeerChannel(chat_id))
                        chat_title = chat.title
                        await asyncio.sleep(2)  
                    except:
                        continue
                else:
                    username_part = message_link.split('/')[3]
                    message_id = int(message_link.split('/')[-1])
                    try:
                        chat = await client.get_entity(username_part)
                        chat_title = chat.title if hasattr(chat, 'title') else chat.username
                        await asyncio.sleep(2) 
                    except:
                        continue

                try:
                    await client(JoinChannelRequest(chat))
                    await asyncio.sleep(3)  
                except:
                    continue

                try:
                    target_message = await client.get_messages(chat, ids=message_id)
                    if not target_message:
                        continue
                    
                    user = await client.get_entity(target_message.sender_id)
                    await asyncio.sleep(1)  
                    
                    if user.id in users_info:
                        continue
                    
                    user_info = f"@{user.username}" if user.username else f"ID: {user.id}"                                        
                    if user.id in private_users["ids"] or (user.username and user.username in private_users["usernames"]):
                        continue
                    
                    premium_status = "✅" if user.premium else "❌"
                    is_bot = "🤖 Бот" if user.bot else "👤 Человек"
                    user_phone = user.phone if user.phone else "Не указан"
                    user_first_name = html.escape(user.first_name) if user.first_name else "Не указано"
                    user_last_name = html.escape(user.last_name) if user.last_name else "Не указано"
                    
                    users_info[user.id] = {
                        "user_info": user_info,
                        "premium_status": premium_status,
                        "is_bot": is_bot,
                        "user_phone": user_phone,
                        "user_first_name": user_first_name,
                        "user_last_name": user_last_name
                    }
                    
                    break

                except Exception as e:
                    print(f"Ошибка обработки пользователя: {e}")
                    continue

            except Exception as e:
                print(f"Ошибка обработки: {e}")
                continue

        if not users_info:
            await status_msg.edit_text(
                "❌ Не удалось получить информацию из предоставленных ссылок",
                parse_mode="HTML"
            )
            await state.finish()
            return
            
        user_info_text = "───── ⋆⋅☆⋅⋆ ─────\n<blockquote>"
        user_info_text += "<b>👤 НАЙДЕННЫЕ ПОЛЬЗОВАТЕЛИ:</b>\n\n"
        
        for i, (user_id, user_data) in enumerate(users_info.items(), 1):
            user_info_text += (
                f"<b>🔹 Пользователь {i}:</b> <code>{user_data['user_info']}</code>\n"
                f"<b>   ├ 👑 Премиум:</b> <code>{user_data['premium_status']}</code>\n"
                f"<b>   ├ Тип:</b> <code>{user_data['is_bot']}</code>\n"
                f"<b>   ├ 📱 Телефон:</b> <code>{user_data['user_phone']}</code>\n"
                f"<b>   ├ 👤 Имя:</b> <code>{user_data['user_first_name']}</code>\n"
                f"<b>   └ 👤 Фамилия:</b> <code>{user_data['user_last_name']}</code>\n\n"
            )

        user_info_text += "</blockquote>───── ⋆⋅☆⋅⋆ ─────\n\n"
        user_info_text += "<b>🚨 ВЫБЕРИТЕ ПРИЧИНУ РЕПОРТА:</b>"
        
        markup = InlineKeyboardMarkup(row_width=2)
        markup.add(
            InlineKeyboardButton('🚫 Спам', callback_data='option_1'),
            InlineKeyboardButton('🔪 Насилие', callback_data='option_2'),
            InlineKeyboardButton('👶 Детское порно', callback_data='option_3'),
            InlineKeyboardButton('🔞 Порно', callback_data='option_4'),
            InlineKeyboardButton('©️ Авторство', callback_data='option_5'),
            InlineKeyboardButton('👤 Данные', callback_data='option_6'),
            InlineKeyboardButton('🌍 Гео', callback_data='option_7'),
            InlineKeyboardButton('🎭 Фейк', callback_data='option_8'),
            InlineKeyboardButton('💊 Наркотики', callback_data='option_9')
        )

        await status_msg.edit_text(
            user_info_text,
            parse_mode="HTML",
            reply_markup=markup
        )
        
        async with state.proxy() as data:
            data['users_info'] = users_info
            data['found_users'] = list(users_info.keys())
            
        await ReportStates.option.set()

    except Exception as e:
        await message.answer(f'<b>❌ Ошибка:</b> <code>{str(e)}</code>', parse_mode="HTML")
        await state.finish()
    finally:
        try:
            if client:
                await client.disconnect()
            if selected_session:
                try:
                    with open(selected_session["json_path"], 'r') as f:
                        session_data = json.load(f)
                    session_data["status"] = "Свободна"
                    with open(selected_session["json_path"], 'w') as f:
                        json.dump(session_data, f)
                except:
                    pass
        except:
            pass

import json




        








import html

#---'zzz'
async def on_startup(dp):
    asyncio.create_task(background_status_checker())
    
from datetime import datetime, timedelta

async def check_payment(user_id):
    paid_users_file = os.path.join('База_Бота', 'paid_users.txt')
    if not os.path.exists(paid_users_file):
        return False
    with open(paid_users_file, 'r') as file:
        lines = file.readlines()
    for line in lines:
        line = line.strip()
        if not line:
            continue
        try:
            paid_user_id, expiry_info = line.split(',', 1)
            if paid_user_id == str(user_id):
                if expiry_info == "forever":
                    return True
                expiry_time = datetime.strptime(expiry_info, '%Y-%m-%d %H:%M:%S')
                return expiry_time > datetime.now()
        except ValueError:
            continue
    return False


async def save_paid_user(user_id, duration_days):
    paid_users_file = os.path.join('База_Бота', 'paid_users.txt')
    user_id_str = str(user_id)
    if duration_days == 36500:
        expiry_info = "forever"
    else:
        expiry_time = datetime.now() + timedelta(days=duration_days)
        expiry_info = expiry_time.strftime('%Y-%m-%d %H:%M:%S')
    lines = []
    if os.path.exists(paid_users_file):
        with open(paid_users_file, 'r') as file:
            lines = file.readlines()
    updated = False
    new_lines = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        try:
            paid_user_id, current_expiry = line.split(',', 1)
            if paid_user_id == user_id_str:
                if current_expiry == "forever":
                    new_lines.append(line + '\n')
                    updated = True
                    continue
                if duration_days != 36500:
                    try:
                        current_time = datetime.strptime(current_expiry, '%Y-%m-%d %H:%M:%S')
                        if current_time > datetime.now():
                            new_time = current_time + timedelta(days=duration_days)
                            expiry_info = new_time.strftime('%Y-%m-%d %H:%M:%S')
                    except ValueError:
                        pass
                new_lines.append(f"{user_id_str},{expiry_info}\n")
                updated = True
            else:
                new_lines.append(line + '\n')
        except ValueError:
            new_lines.append(line + '\n')
    if not updated:
        new_lines.append(f"{user_id_str},{expiry_info}\n")
    with open(paid_users_file, 'w') as file:
        file.writelines(new_lines)


async def update_time():
    if not os.path.exists('База_Бота/paid_users.tx'):
        return
    
    with open('База_Бота/paid_users.txt', 'r') as file:
        lines = file.readlines()
    
    updated_lines = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        try:
            user_id, expiry_info = line.split(',', 1)
            if expiry_info == "forever":
                updated_lines.append(line + '\n')
                continue
            
            expiry_time = datetime.strptime(expiry_info, '%Y-%m-%d %H:%M:%S')
            if expiry_time > datetime.now():
                expiry_time -= timedelta(seconds=1)
                expiry_info = expiry_time.strftime('%Y-%m-%d %H:%M:%S')
            updated_lines.append(f"{user_id},{expiry_info}\n")
        except ValueError as e:
            print(f"Ошибка при обработке строки '{line}': {e}")
            updated_lines.append(line + '\n')
    
    with open('База_Бота/paid_users.txt', 'w') as file:
        file.writelines(updated_lines)

async def check_and_notify():
    if not os.path.exists('База_Бота/paid_users.txt'):
        return
    with open('База_Бота/paid_users.txt', 'r') as file:
        lines = file.readlines()
    for line in lines:
        user_id, expiry_time_str = line.strip().split(',')
        expiry_time = datetime.strptime(expiry_time_str, '%Y-%m-%d %H:%M:%S')
        if expiry_time <= datetime.now():
            markup = InlineKeyboardMarkup()
            markup.add(InlineKeyboardButton("Купить время", callback_data="go_to_payment"))
            await bot.send_message(user_id, "⏳ Ваше время истекло. Пожалуйста, купите дополнительное время.", reply_markup=markup)
import uuid

def create_invoice(asset, amount, description):
    url = "https://pay.crypt.bot/api/createInvoice"
    headers = {
        "Crypto-Pay-API-Token": CRYPTO_PAY_TOKEN,
        "Content-Type": "application/json"
    }
    data = {
        "asset": asset,
        "amount": str(amount),
        "description": description,
        "payload": "custom_payload"
    }
    response = requests.post(url, headers=headers, json=data)
    if response.status_code == 200:
        return response.json()
    else:
        logging.error(f"Ошибка при создании счета: {response.status_code} - {response.text}")
        return None

def check_invoice_status(invoice_id):
    url = "https://pay.crypt.bot/api/getInvoices"
    headers = {
        "Crypto-Pay-API-Token": CRYPTO_PAY_TOKEN,
        "Content-Type": "application/json"
    }
    params = {"invoice_ids": [invoice_id]}
    response = requests.get(url, headers=headers, params=params)
    if response.status_code == 200:
        return response.json()
    else:
        logging.error(f"Ошибка при проверке статуса счета: {response.status_code} - {response.text}")
        return None

@dp.callback_query_handler(lambda c: c.data == 'buy_vip')
async def process_callback_buy_vip(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    
    markup = InlineKeyboardMarkup(row_width=1)
    
    for crypto, price in VIP_PRICE.items():
        markup.add(InlineKeyboardButton(
            f"{crypto}: {price}",
            callback_data=f"vip_pay_{crypto.lower()}"
        ))
    
    markup.add(InlineKeyboardButton("🔙 Назад", callback_data="my_time"))
    
    await bot.answer_callback_query(callback_query.id)
    if callback_query.message.photo:
        photo = callback_query.message.photo[-1].file_id  
        media = InputMediaPhoto(media=photo, caption="💸 Выберите криптовалюту для оплаты VIP-статуса:")
        
        await bot.edit_message_media(
            chat_id=user_id,
            message_id=callback_query.message.message_id,
            media=media,
            reply_markup=markup
        )
    else:
        await bot.edit_message_text(
            chat_id=user_id,
            message_id=callback_query.message.message_id,
            text="💸 Выберите криптовалюту для оплаты VIP-статуса:",
            reply_markup=markup
        )

@dp.callback_query_handler(lambda c: c.data.startswith('vip_pay_'))
async def process_callback_vip_pay(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    crypto_type = callback_query.data.split('_')[-1]
    crypto_type_upper = crypto_type.upper()
    
    price = VIP_PRICE.get(crypto_type_upper)
    
    if price is None:
        await bot.answer_callback_query(callback_query.id, "❌ Ошибка: криптовалюта не найдена.")
        return
    
    invoice = create_invoice(asset=crypto_type_upper, amount=price, description="Оплата VIP-статуса") 
    
    if invoice and 'result' in invoice:
        invoice_id = invoice['result']['invoice_id']
        pay_url = invoice['result']['pay_url']
        
        markup = InlineKeyboardMarkup(row_width=2)
        markup.add(
            InlineKeyboardButton("💳 Оплатить", url=pay_url),
            InlineKeyboardButton("✅ Проверить оплату", callback_data=f"check_vip_{invoice_id}")
        )
        markup.add(InlineKeyboardButton("🔙 Назад", callback_data="buy_vip"))
        
        await bot.answer_callback_query(callback_query.id)
        
        text = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🌟 Вайтлист 🌟</b>

📝 Добавьте себя в данный список, чтобы избежать сноса от других пользователей проекта и получить VIP статус

💸 <b>Стоимость добавления:</b> {price} {crypto_type_upper}
🛡️ Оплатите счет, и вы будете защищены.</blockquote>
<b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
        
        if callback_query.message.photo:
            await bot.edit_message_caption(
                chat_id=user_id,
                message_id=callback_query.message.message_id,
                caption=text,
                reply_markup=markup,
                parse_mode="HTML"
            )
        else:
            await bot.edit_message_text(
                chat_id=user_id,
                message_id=callback_query.message.message_id,
                text=text,
                reply_markup=markup,
                parse_mode="HTML"
            )
    else:
        await bot.answer_callback_query(callback_query.id, "❌ Ошибка при создании счета.")

@dp.callback_query_handler(lambda c: c.data.startswith('check_vip_'))
async def process_callback_check_vip(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    user_username = callback_query.from_user.username

    logger.info(f"Processing callback with data: {callback_query.data}")
    
    if not callback_query.data.startswith('check_vip_'):
        logger.error(f"Invalid callback data format: {callback_query.data}")
        await callback_query.answer("❌ Ошибка: неверный формат данных.", show_alert=True)
        return

    invoice_id = callback_query.data.split('_')[2]
    logger.info(f"Checking invoice status for ID: {invoice_id}")
    
    invoice_status = check_invoice_status(invoice_id)
    logger.info(f"Invoice status response: {invoice_status}")

    if not invoice_status or not invoice_status.get('ok'):
        await callback_query.answer("❌ Чек не найден. Пожалуйста, попробуйте снова.", show_alert=True)
        return

    items = invoice_status.get('result', {}).get('items', [])
    if not items:
        await callback_query.answer("❌ Чек не найден. Пожалуйста, попробуйте снова.", show_alert=True)
        return

    status = items[0].get('status')
    logger.info(f"Invoice status: {status}")

    if status == 'paid':
        private_users = read_private_users()
        if user_id not in private_users["ids"]:
            private_users["ids"].append(user_id)
        if user_username and user_username not in private_users["usernames"]:
            private_users["usernames"].append(user_username)
        write_private_users(private_users)

        await callback_query.answer("✅ Оплата прошла успешно! VIP-статус активирован.", show_alert=True)
        await process_callback_my_time(callback_query) 
    elif status == 'active':
        await callback_query.answer("ℹ️ Счет ожидает оплаты. Пожалуйста, завершите оплату.", show_alert=True)
    elif status in ['pending', 'unpaid', 'processing']:
        await callback_query.answer("⏳ Оплата еще обрабатывается. Пожалуйста, подождите.", show_alert=True)
    elif status in ['expired', 'failed', 'cancelled']:
        await callback_query.answer("❌ Оплата не найдена или отклонена. Пожалуйста, попробуйте снова.", show_alert=True)
    else:
        await callback_query.answer(f"❌ Неизвестный статус оплаты: {status}", show_alert=True)

        
@dp.callback_query_handler(lambda c: c.data == 'return_to_welcome')
async def return_to_welcome_handler(callback_query: types.CallbackQuery):
    user = callback_query.from_user
    message = callback_query.message

    async def reply_adapter(photo=None, caption=None, reply_markup=None, parse_mode="HTML"):
        if photo:
            if message.photo:
                media = InputMediaPhoto(media=photo, caption=caption, parse_mode=parse_mode)
                await bot.edit_message_media(
                    chat_id=message.chat.id,
                    message_id=message.message_id,
                    media=media,
                    reply_markup=reply_markup
                )
            else:
                await bot.delete_message(message.chat.id, message.message_id)
                await bot.send_photo(
                    chat_id=message.chat.id,
                    photo=photo,
                    caption=caption,
                    reply_markup=reply_markup,
                    parse_mode=parse_mode
                )
        else:
            if message.photo:
                await bot.delete_message(message.chat.id, message.message_id)
                await bot.send_message(
                    chat_id=message.chat.id,
                    text=caption,
                    reply_markup=reply_markup,
                    parse_mode=parse_mode
                )
            else:
                await bot.edit_message_text(
                    chat_id=message.chat.id,
                    message_id=message.message_id,
                    text=caption,
                    reply_markup=reply_markup,
                    parse_mode=parse_mode
                )

    await handle_welcome(
        user_id=user.id,
        chat_id=message.chat.id,
        from_user=user,
        reply_photo_method=reply_adapter
    )



async def handle_welcome(user_id: int, chat_id: int, from_user: types.User, reply_photo_method):
    add_user_to_file(user_id)
    paid_users_path = os.path.join('База_Бота', 'paid_users.txt')

    if not os.path.exists(paid_users_path):
        os.makedirs('База_Бота', exist_ok=True)  
        with open(paid_users_path, 'w') as file:
            pass

    if not await check_payment(user_id) and str(user_id) not in admin_chat_ids:  
        markup = InlineKeyboardMarkup()
        markup.add(InlineKeyboardButton("💳 Перейти к оплате", callback_data="extend_subscription"))
        markup.add(InlineKeyboardButton("🔑 Активировать Промокод", callback_data="activate_promo"))
        markup.add(InlineKeyboardButton("📩 Написать поддержку📩", callback_data='support'))
        markup.add(InlineKeyboardButton("👤Профиль👤", callback_data='my_time'))
        photo_path = os.path.join('Фото', 'unnamed.jpg')
        
        await reply_photo_method(
            photo=open(photo_path, 'rb'),
            caption=WELCOME_BANNER,
            reply_markup=markup,
            parse_mode="HTML"
        )
        return
    
    first_name = from_user.first_name if from_user.first_name else ''
    last_name = from_user.last_name if from_user.last_name else ''
    username = f"@{from_user.username}" if from_user.username else f"id{from_user.id}"
    
    welcome_message = PAID_WELCOME_BANNER.format(
        first_name=first_name,
        last_name=last_name,
        username=username
    )
    
    await send_menu(chat_id, welcome_message)

async def is_user_subscribed(user_id, channel_id):
    try:
        member = await bot.get_chat_member(chat_id=channel_id, user_id=user_id)
        return member.status not in ["left", "kicked"]
    except Exception as e:
        print(f"Ошибка при проверке подписки на канал {channel_id}: {e}")
        return False

async def get_channel_name(channel_id):
    try:
        chat = await bot.get_chat(chat_id=channel_id)
        return chat.title  
    except Exception as e:
        print(f"Ошибка при получении названия канала {channel_id}: {e}")
        return f"Канал {channel_id}"  

async def check_all_subscriptions(user_id):
    not_subscribed_channels = {}
    for channel_id, channel_url in CHANNELS.items():
        if not await is_user_subscribed(user_id, channel_id):
            channel_name = await get_channel_name(channel_id)
            not_subscribed_channels[channel_name] = channel_url
    return not_subscribed_channels  

import string



class UserStates(StatesGroup):
    waiting_for_subscription = State()

def generate_referral_code(length=16):
    characters = string.ascii_letters + string.digits
    return ''.join(random.choice(characters) for _ in range(length))

import json
from datetime import datetime, timedelta
from aiogram import types
from aiogram.dispatcher import FSMContext

async def handle_referral_data(user_id: int, bot_username: str, referrer_id: str = None):
    referral_file = os.path.join('База_Бота', 'referral_data.json')
    try:
        with open(referral_file, 'r', encoding='utf-8') as f:
            referral_data = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        referral_data = {}
    user_id_str = str(user_id)
    if user_id_str not in referral_data:
        referral_code = generate_referral_code()
        referral_link = f"https://t.me/{bot_username}?start=ref_{user_id_str}_{referral_code}"
        referral_data[user_id_str] = {
            "referral_link": referral_link,
            "referral_code": referral_code,
            "referral_count": 0,
            "balance": 0.0,
            "referrals": []
        }
        if referrer_id and referrer_id in referral_data and referrer_id != user_id_str:
            referral_data[referrer_id]['referral_count'] += 1
            referral_data[referrer_id]['balance'] += 0.2
            referral_data[referrer_id]['referrals'].append(user_id_str)
            try:
                await bot.send_message(int(referrer_id), f"🎉 Новый реферал!\n💵 Баланс: ${referral_data[referrer_id]['balance']:.2f}")
            except Exception:
                pass
        with open(referral_file, 'w', encoding='utf-8') as f:
            json.dump(referral_data, f, indent=4, ensure_ascii=False)
    return referral_data.get(user_id_str, {})


from typing import Union
@dp.callback_query_handler(lambda c: c.data == 'my_time')
@dp.message_handler(commands=['profile'])
async def show_user_profile(update: Union[types.CallbackQuery, types.Message]):
    if isinstance(update, types.CallbackQuery):
        message = update.message
        user = update.from_user
        is_callback = True
    else:
        message = update
        user = update.from_user
        is_callback = False
    
    user_id = user.id
    user_name = html.escape(user.first_name)
    user_username = html.escape(user.username) if user.username else "отсутствует"
    current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    subscription_end = await get_subscription_end_time(user_id)
    remaining_time = await get_remaining_time(user_id)
    has_active_subscription = subscription_end == "forever" or (subscription_end and subscription_end > datetime.now())
    is_admin = str(user_id) in admin_chat_ids
    if subscription_end == "forever":
        subscription_status = "✅ <b>Активна (Навсегда)</b>"
        subscription_end_formatted = "∞"
    elif subscription_end and subscription_end > datetime.now():
        subscription_status = "✅ <b>Активна</b>"
        subscription_end_formatted = subscription_end.strftime("%Y-%m-%d %H:%M:%S")
    else:
        subscription_status = "❌ <b>Не активна</b>"
        subscription_end_formatted = "<i>Нет активной подписки</i>"
    private_users = read_private_users()
    is_vip = user_id in private_users["ids"] or user_username in private_users["usernames"]
    vip_status = "✅ <b>Да</b>" if is_vip else "❌ <b>Нет</b>"
    try:
        with open('База_Бота/referral_data.json', 'r', encoding='utf-8') as f:
            referral_data = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        referral_data = {}
    
    user_id_str = str(user_id)
    user_data = referral_data.get(user_id_str, {})
    if not user_data:
        referral_info = "🔗 Реферальная система: ❌ Не настроена"
    else:
        referral_info = f"""
🔗 <b>Реферальная система</b>
├ 💰 Баланс: ${user_data.get('balance', 0):.2f}
├ 👥 Приглашено: {user_data.get('referral_count', 0)}
└ 🔗 Ссылка: {user_data.get('referral_link', 'Не создана')}
        """

    profile_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>⚡️ Профиль пользователя ⚡️</b>

<b>🆔 ID:</b> <code>{user_id}</code>
<b>👤 Имя:</b> <code>{user_name}</code>
<b>👤 Юзернейм:</b> @{user_username}
<b>🕐 Текущее время:</b> <code>{current_time}</code>
<b>🔐 Подписка:</b> {subscription_status}
<b>💰 Подписка до:</b> <code>{subscription_end_formatted}</code>
<b>⏳ Оставшееся время:</b> <code>{remaining_time}</code>
<b>🌟 VIP статус:</b> {vip_status}

{referral_info}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    markup = InlineKeyboardMarkup()
    if has_active_subscription and subscription_end != "forever" and not is_admin:
        markup.add(InlineKeyboardButton("🔄 Продлить подписку", callback_data="extend_subscription"))
    
    if has_active_subscription or is_admin:
        markup.add(InlineKeyboardButton("🔙 Назад", callback_data="to_start"))  
    else:
        markup.add(InlineKeyboardButton("🔙 Назад", callback_data="return_to_welcome"))
    
    if not is_vip:
        markup.add(InlineKeyboardButton("🌟 Купить VIP", callback_data="buy_vip"))
    if user_data:
        markup.add(InlineKeyboardButton("🔄 Обновить реферальную ссылку", callback_data="update_referral"))
    if is_callback:
        await bot.answer_callback_query(update.id)
        
        if message.photo:
            photo = message.photo[-1].file_id
            media = InputMediaPhoto(
                media=photo, 
                caption=profile_message,
                parse_mode="HTML"
            )
            await bot.edit_message_media(
                chat_id=user_id,
                message_id=message.message_id,
                media=media,
                reply_markup=markup
            )
        else:
            await bot.edit_message_text(
                chat_id=user_id,
                message_id=message.message_id,
                text=profile_message,
                reply_markup=markup,
                parse_mode="HTML"
            )
    else:
        await bot.send_message(
            chat_id=user_id,
            text=profile_message,
            reply_markup=markup,
            parse_mode="HTML"
        )

        
@dp.callback_query_handler(lambda c: c.data == 'update_referral')  
async def handle_update_referral(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    bot_username = (await bot.me).username
    
    try:
        with open('База_Бота/referral_data.json', 'r', encoding='utf-8') as f:
            referral_data = json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        referral_data = {}
    
    user_id_str = str(user_id)
    
    if user_id_str in referral_data:
        referral_code = generate_referral_code()
        referral_data[user_id_str].update({
            "referral_link": f"https://t.me/{bot_username}?start=ref_{user_id_str}_{referral_code}",
            "referral_code": referral_code
        })
    
    with open('База_Бота/referral_data.json', 'w', encoding='utf-8') as f:
        json.dump(referral_data, f, indent=4, ensure_ascii=False)

    await show_user_profile(callback_query)                        
@dp.callback_query_handler(lambda c: c.data == 'buy_subscription')
async def process_buy_subscription(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    
    try:
        with open('База_Бота/referral_data.json', 'r', encoding='utf-8') as f:
            referral_data = json.load(f)
        user_balance = referral_data.get(str(user_id), {}).get('balance', 0)
    except (FileNotFoundError, json.JSONDecodeError):
        user_balance = 0
    
    markup = InlineKeyboardMarkup(row_width=1)
    markup.add(InlineKeyboardButton("💳 Криптовалюта", callback_data="go_to_payment"))
    markup.add(InlineKeyboardButton(f"💰 С баланса (Доступно: ${user_balance:.2f})", callback_data="balance_payment"))
    markup.add(InlineKeyboardButton("🔙 Назад", callback_data="to_start"))
    
    caption_text = f"<b>💸 Выберите способ покупки подписки</b>\n\n<b>Ваш текущий баланс:</b> <code>${user_balance:.2f}</code>"
    
    await bot.answer_callback_query(callback_query.id)
    
    if callback_query.message.photo:
        await callback_query.message.edit_caption(caption=caption_text, reply_markup=markup, parse_mode="HTML")
    else:
        await callback_query.message.edit_text(text=caption_text, reply_markup=markup, parse_mode="HTML")

@dp.callback_query_handler(lambda c: c.data == 'extend_subscription')
async def process_extend_subscription(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    is_admin = user_id in admin_chat_ids  
    

    subscription_result = await check_user_subscription(user_id) if asyncio.iscoroutinefunction(check_user_subscription) else check_user_subscription(user_id)
    has_active_subscription = subscription_result if isinstance(subscription_result, bool) else subscription_result[0]  
    
    try:
        with open('База_Бота/referral_data.json', 'r', encoding='utf-8') as f:
            referral_data = json.load(f)
        user_balance = referral_data.get(str(user_id), {}).get('balance', 0)
    except (FileNotFoundError, json.JSONDecodeError):
        user_balance = 0
    
    markup = InlineKeyboardMarkup(row_width=1)
    markup.add(InlineKeyboardButton("💳 Криптовалюта", callback_data="go_to_payment"))
    markup.add(InlineKeyboardButton(f"💰 С баланса (Доступно: ${user_balance:.2f})", callback_data="extend_balance_payment"))
    
    back_button_callback = "my_time" if has_active_subscription or is_admin else "return_to_welcome"
    markup.add(InlineKeyboardButton("🔙 Назад", callback_data=back_button_callback))
    
    caption_text = f"<b>💸 Выберите способ продления подписки</b>\n\n<b>Ваш текущий баланс:</b> <code>${user_balance:.2f}</code>"
    
    await bot.answer_callback_query(callback_query.id)
    
    if callback_query.message.photo:
        await callback_query.message.edit_caption(caption=caption_text, reply_markup=markup, parse_mode="HTML")
    else:
        await callback_query.message.edit_text(text=caption_text, reply_markup=markup, parse_mode="HTML")

@dp.callback_query_handler(lambda c: c.data == "go_to_payment")
async def process_go_to_payment(callback_query: types.CallbackQuery):
    await callback_query.answer()
    markup = InlineKeyboardMarkup()
    markup.add(InlineKeyboardButton("1 день⚡️", callback_data="period_1_day"))
    markup.add(InlineKeyboardButton("2 дня⚡️", callback_data="period_2_days"))
    markup.add(InlineKeyboardButton("5 дней⚡️", callback_data="period_5_days"))
    markup.add(InlineKeyboardButton("Месяц⚡️", callback_data="period_30_days"))
    markup.add(InlineKeyboardButton("Год⚡️", callback_data="period_1_year"))
    markup.add(InlineKeyboardButton("Навсегда⚡️", callback_data="period_forever"))
    markup.add(InlineKeyboardButton("🔙 Назад", callback_data="back_to_start"))
    
    if callback_query.message.photo:
        await callback_query.message.edit_caption(
            caption="💸 *Выберите период доступа:* 💸",
            reply_markup=markup,
            parse_mode="Markdown"
        )
    else:
        await callback_query.message.edit_text(
            text="💸 *Выберите период доступа:* 💸",
            reply_markup=markup,
            parse_mode="Markdown"
        )

@dp.callback_query_handler(lambda c: c.data.startswith('period_'))
async def process_callback_period(callback_query: types.CallbackQuery):
    try:
        parts = callback_query.data.split('_')
        if len(parts) == 3:  
            period = parts[1] + "_" + parts[2]
        elif callback_query.data == "period_forever":
            period = "forever"
        else:
            raise ValueError("Invalid period format")
            
        keyboard = InlineKeyboardMarkup(row_width=2)
        for currency, price in CURRENCY_PRICES[period].items():
            keyboard.add(InlineKeyboardButton(f"{currency} 💳 ({price})", callback_data=f"pay_{period}_{currency}"))
        keyboard.add(InlineKeyboardButton("🔙 Назад", callback_data="back_to_periods"))
        
        await bot.answer_callback_query(callback_query.id)
        caption_text = f"💸 *Выберите валюту для оплаты* ({'Навсегда' if period == 'forever' else period.replace('_', ' ')}) 💸"
        
        if callback_query.message.photo:
            await callback_query.message.edit_caption(
                caption=caption_text,
                reply_markup=keyboard,
                parse_mode="Markdown"
            )
        else:
            await callback_query.message.edit_text(
                text=caption_text,
                reply_markup=keyboard,
                parse_mode="Markdown"
            )
            
    except Exception as e:
        logger.error(f"Error in process_callback_period: {e}")
        await bot.answer_callback_query(callback_query.id, "❌ Произошла ошибка, попробуйте снова")

@dp.callback_query_handler(lambda c: c.data.startswith('pay_'))
async def process_callback_currency(callback_query: types.CallbackQuery):
    try:
        parts = callback_query.data.split('_')
        if len(parts) == 4:
            period = parts[1] + "_" + parts[2]
            asset = parts[3]
        elif len(parts) == 3:
            period = parts[1]
            asset = parts[2]
        else:
            raise ValueError("Invalid payment format")

        amount = CURRENCY_PRICES[period].get(asset, 0)
        
        if period == "forever":
            duration_days = 36500
        else:
            duration_days = int(period.split('_')[0])
        
        invoice = create_invoice(asset=asset, amount=amount, description=f"Оплата через CryptoBot на {'вечный доступ' if period == 'forever' else f'{duration_days} дней'}")

        if invoice and 'result' in invoice:
            invoice_id = invoice['result']['invoice_id']
            pay_url = invoice['result']['pay_url']
            markup = InlineKeyboardMarkup(row_width=1)
            markup.add(InlineKeyboardButton("💳 Оплатить", url=pay_url))
            markup.add(InlineKeyboardButton("✅ Проверить оплату", callback_data=f"check_{invoice_id}_{duration_days}"))
            markup.add(InlineKeyboardButton("🔙 Назад", callback_data=f"back_to_currencies_{period}"))
            
            await bot.answer_callback_query(callback_query.id)
            message_text = "💸 *Оплатите по кнопке ниже и нажмите кнопку 'Проверить оплату'* 💸"
            
            if callback_query.message.photo:
                await callback_query.message.edit_caption(
                    caption=message_text,
                    reply_markup=markup,
                    parse_mode="Markdown"
                )
            else:
                await callback_query.message.edit_text(
                    text=message_text,
                    reply_markup=markup,
                    parse_mode="Markdown"
                )
        else:
            await bot.answer_callback_query(callback_query.id, "❌ Ошибка при создании счета")
            
    except Exception as e:
        logger.error(f"Error in process_callback_currency: {e}")
        await bot.answer_callback_query(callback_query.id, "❌ Произошла ошибка, попробуйте снова")

@dp.callback_query_handler(lambda c: c.data.startswith('back_to_'))
async def process_callback_back(callback_query: types.CallbackQuery):
    try:
        data = callback_query.data.split('_')
        
        if data[2] == "periods":
            markup = InlineKeyboardMarkup()
            markup.add(InlineKeyboardButton("1 день⚡️", callback_data="period_1_day"))
            markup.add(InlineKeyboardButton("2 дня⚡️", callback_data="period_2_days"))
            markup.add(InlineKeyboardButton("5 дней⚡️", callback_data="period_5_days"))
            markup.add(InlineKeyboardButton("Месяц⚡️", callback_data="period_30_days"))
            markup.add(InlineKeyboardButton("Год⚡️", callback_data="period_1_year"))
            markup.add(InlineKeyboardButton("Навсегда⚡️", callback_data="period_forever"))
            markup.add(InlineKeyboardButton("🔙 Назад", callback_data="back_to_start"))
            
            if callback_query.message.photo:
                await callback_query.message.edit_caption(
                    caption="💸 *Выберите период доступа:* 💸",
                    reply_markup=markup,
                    parse_mode="Markdown"
                )
            else:
                await callback_query.message.edit_text(
                    text="💸 *Выберите период доступа:* 💸",
                    reply_markup=markup,
                    parse_mode="Markdown"
                )
                
        elif data[2] == "currencies":
            if len(data) >= 4:
                period = data[3] + ("_" + data[4] if len(data) > 4 else "")
                keyboard = InlineKeyboardMarkup(row_width=2)
                for currency, price in CURRENCY_PRICES[period].items():
                    keyboard.add(InlineKeyboardButton(f"{currency} 💳 ({price})", callback_data=f"pay_{period}_{currency}"))
                keyboard.add(InlineKeyboardButton("🔙 Назад", callback_data="back_to_periods"))
                
                await bot.answer_callback_query(callback_query.id)
                period_name = "Навсегда" if period == "forever" else period.replace('_', ' ')
                
                if callback_query.message.photo:
                    await callback_query.message.edit_caption(
                        caption=f"💸 *Выберите валюту для оплаты* ({period_name}) 💸",
                        reply_markup=keyboard,
                        parse_mode="Markdown"
                    )
                else:
                    await callback_query.message.edit_text(
                        text=f"💸 *Выберите валюту для оплаты* ({period_name}) 💸",
                        reply_markup=keyboard,
                        parse_mode="Markdown"
                    )
            else:
                await bot.answer_callback_query(callback_query.id, "❌ Ошибка в данных")
                
        elif data[2] == "start":
            markup = InlineKeyboardMarkup()
            markup.add(InlineKeyboardButton("💳Перейти к оплате", callback_data="extend_subscription"))
            markup.add(InlineKeyboardButton("🔑 Активировать Промокод", callback_data="activate_promo"))
            markup.add(InlineKeyboardButton("📩 Написать поддержку📩", callback_data='support'))
            markup.add(InlineKeyboardButton("👤Профиль👤", callback_data='my_time'))    
            if callback_query.message.photo:
                await callback_query.message.edit_caption(
                    caption=WELCOME_BANNER,
                    reply_markup=markup,
                    parse_mode="HTML"
                )
            else:
                await callback_query.message.edit_text(
                    text=WELCOME_BANNER,
                    reply_markup=markup,
                    parse_mode="HTML"
                )
    except Exception as e:
        logger.error(f"Error in process_callback_back: {e}")
        await bot.answer_callback_query(callback_query.id, "❌ Произошла ошибка, попробуйте снова")

@dp.callback_query_handler(lambda c: c.data.startswith('check_'))
async def process_callback_check(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    logger.info(f"Processing callback from user {user_id} with data: {callback_query.data}")

    try:
        _, invoice_id, duration = callback_query.data.split('_')
        is_permanent = duration == 'forever'
    except ValueError:
        logger.error(f"Invalid callback data format: {callback_query.data}")
        await callback_query.answer("❌ Ошибка: неверный формат данных.", show_alert=True)
        return

    logger.info(f"Checking invoice {invoice_id} for {'permanent' if is_permanent else duration + ' days'} access")
    
    invoice_status = check_invoice_status(invoice_id)
    if not invoice_status or not invoice_status.get('ok'):
        await callback_query.answer("❌ Чек не найден. Пожалуйста, попробуйте снова.", show_alert=True)
        return

    items = invoice_status.get('result', {}).get('items', [])
    if not items:
        await callback_query.answer("❌ Чек не найден. Пожалуйста, попробуйте снова.", show_alert=True)
        return

    status = items[0].get('status')
    logger.info(f"Invoice status: {status}")

    if status == 'paid':
        await handle_successful_payment(user_id, is_permanent, callback_query)
    elif status == 'active':
        await callback_query.answer("ℹ️ Счет ожидает оплаты. Пожалуйста, завершите оплату.", show_alert=True)
    elif status in ['pending', 'unpaid', 'processing']:
        await callback_query.answer("⏳ Оплата еще обрабатывается. Пожалуйста, подождите.", show_alert=True)
    elif status in ['expired', 'failed', 'cancelled']:
        await callback_query.answer("❌ Оплата не найдена или отклонена. Пожалуйста, попробуйте снова.", show_alert=True)
    else:
        await callback_query.answer(f"❌ Неизвестный статус оплаты: {status}", show_alert=True)


async def handle_successful_payment(user_id: int, is_permanent: bool, callback_query: types.CallbackQuery):
    try:
        current_sub = await get_user_subscription(user_id)
        
        if is_permanent:
            new_expiry = datetime.max
            await update_user_subscription(user_id, new_expiry)
        else:
            duration_days = int(callback_query.data.split('_')[2])
            if current_sub:
                if current_sub['expiry_date'] > datetime.now():
                    new_expiry = current_sub['expiry_date'] + timedelta(days=duration_days)
                else:
                    new_expiry = datetime.now() + timedelta(days=duration_days)
                await update_user_subscription(user_id, new_expiry)
            else:
                new_expiry = datetime.now() + timedelta(days=duration_days)
                await create_user_subscription(user_id, new_expiry)
        
        await callback_query.answer("✅ Оплата успешно подтверждена!", show_alert=True)
        await send_success_message(user_id, is_permanent)
        
    except Exception as e:
        logger.error(f"Error processing payment for user {user_id}: {str(e)}")
        await callback_query.answer("❌ Ошибка при обработке оплаты. Пожалуйста, свяжитесь с поддержкой.", show_alert=True)


async def send_success_message(user_id: int, is_permanent: bool):
    if is_permanent:
        expiry_date = "∞ (Навсегда)"
        duration_text = "∞ (Навсегда)"
    else:
        duration_days = int(callback_query.data.split('_')[2])
        expiry_date = (datetime.now() + timedelta(days=duration_days)).strftime("%d.%m.%Y")
        duration_text = f"{duration_days} дней"
    
    success_message = (
        "<b>───── ⋆⋅☆⋅⋆ ─────</b>\n<blockquote>"
        "<b>🎉 Оплата прошла успешно!</b>\n\n"
        f"<b>🔹 Ваш доступ активирован на</b> <code>{duration_text}</code>\n"
        f"<b>🔹 Дата окончания:</b> <code>{expiry_date}</code>\n\n"
        "<i>Теперь вам доступны все функции бота:</i>\n"
        "• Snos\n"
        "• Spam\n"
        "• Востановить акк\n"
        "<b>👇 Нажмите кнопку ниже чтобы начать</b>"
        "</blockquote>\n<b>───── ⋆⋅☆⋅⋆ ─────</b>"
    )
    
    await bot.send_message(
        user_id,
        success_message,
        parse_mode="HTML",
        reply_markup=InlineKeyboardMarkup().add(
            InlineKeyboardButton("🚀 НАЧАТЬ ИСПОЛЬЗОВАНИЕ", callback_data="send_welcome")
        )
    )
    
@dp.callback_query_handler(lambda c: c.data == "balance_payment" or c.data == "extend_balance_payment")
async def process_balance_payment(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    is_extend = callback_query.data.startswith('extend_')
    
    try:
        with open('База_Бота/referral_data.json', 'r', encoding='utf-8') as f:
            referral_data = json.load(f)
        user_balance = referral_data.get(str(user_id), {}).get('balance', 0)
    except (FileNotFoundError, json.JSONDecodeError):
        user_balance = 0
    
    PRICES = {
        '1_day': 5,
        '2_days': 9,
        '5_days': 20,
        '30_days': 50,
        '1_year': 200,
        'forever': 500
    }
    
    markup = InlineKeyboardMarkup(row_width=1)
    for period, price in PRICES.items():
        period_name = period.replace('_', ' ')
        btn_text = f"{period_name} (${price}) {'✅' if user_balance >= price else '❌'}"
        callback_data = f"{'extend_' if is_extend else ''}balance_{period}"
        markup.add(InlineKeyboardButton(btn_text, callback_data=callback_data))
    
    markup.add(InlineKeyboardButton("🔙 Назад", callback_data="extend_subscription" if is_extend else "buy_subscription"))
    
    caption_text = f"💸 *Выберите период ({'продление' if is_extend else 'покупка'} за баланс)*\n\n💰 *Доступно:* ${user_balance:.2f} 💸"
    
    await bot.answer_callback_query(callback_query.id)
    
    if callback_query.message.photo:
        await callback_query.message.edit_caption(caption=caption_text, reply_markup=markup, parse_mode="Markdown")
    else:
        await callback_query.message.edit_text(text=caption_text, reply_markup=markup, parse_mode="Markdown")

@dp.callback_query_handler(lambda c: c.data.startswith('balance_') or c.data.startswith('extend_balance_'))
async def process_balance_period(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    is_extend = callback_query.data.startswith('extend_')
    period = callback_query.data.replace('extend_', '').split('_')[1:]
    
    if len(period) == 2:
        period = f"{period[0]}_{period[1]}"
    else:
        period = period[0]
    price = PRICES.get(period, 0)
    
    try:
        with open('База_Бота/referral_data.json', 'r', encoding='utf-8') as f:
            referral_data = json.load(f)
        user_data = referral_data.get(str(user_id), {})
        user_balance = user_data.get('balance', 0)
    except (FileNotFoundError, json.JSONDecodeError):
        await callback_query.answer("❌ Ошибка при проверке баланса", show_alert=True)
        return
    
    if user_balance < price:
        await callback_query.answer(f"❌ Недостаточно средств. Нужно ${price:.2f}", show_alert=True)
        return
    
    user_data['balance'] = user_balance - price
    referral_data[str(user_id)] = user_data
    
    with open('База_Бота/referral_data.json', 'w', encoding='utf-8') as f:
        json.dump(referral_data, f, ensure_ascii=False, indent=4)
    
    duration_days = 36500 if period == "forever" else int(period.split('_')[0])
    await save_paid_user(user_id, duration_days, is_extend=is_extend)
    
    expiry_date = "∞ (Навсегда)" if period == "forever" else (datetime.now() + timedelta(days=duration_days)).strftime("%d.%m.%Y")
    
    success_message = (
        f"<b>🎉 {'Продление' if is_extend else 'Покупка'} подписки успешно!</b>\n\n"
        f"<b>🔹 Сумма:</b> <code>${price:.2f}</code>\n"
        f"<b>🔹 Новый баланс:</b> <code>${user_balance - price:.2f}</code>\n"
        f"<b>🔹 Срок:</b> <code>{'Навсегда' if period == 'forever' else duration_days} дней</code>\n"
        f"<b>🔹 Окончание:</b> <code>{expiry_date}</code>"
    )
    
    await callback_query.answer("✅ Оплата прошла успешно!", show_alert=True)
    
    if callback_query.message.photo:
        await callback_query.message.edit_caption(caption=success_message, parse_mode="HTML")
    else:
        await callback_query.message.edit_text(text=success_message, parse_mode="HTML")

async def save_paid_user(user_id, duration_days, is_extend=False):
    if duration_days == 36500:
        expiry_time_str = "forever"
    else:
        if is_extend:
            current_end = await get_subscription_end_time(user_id)
            if current_end == "forever":
                expiry_time_str = "forever"
            else:
                if isinstance(current_end, str):
                    current_end = datetime.strptime(current_end, '%Y-%m-%d %H:%M:%S')
                expiry_time = current_end + timedelta(days=duration_days)
                expiry_time_str = expiry_time.strftime('%Y-%m-%d %H:%M:%S')
        else:
            expiry_time = datetime.now() + timedelta(days=duration_days)
            expiry_time_str = expiry_time.strftime('%Y-%m-%d %H:%M:%S')
    
    with open('База_Бота/paid_users.txt', 'a+') as f:
        f.seek(0)
        lines = f.readlines()
        
        updated = False
        new_lines = []
        for line in lines:
            if line.strip():
                paid_user_id, paid_expiry = line.strip().split(',')
                if paid_user_id == str(user_id):
                    new_lines.append(f"{paid_user_id},{expiry_time_str}\n")
                    updated = True
                else:
                    new_lines.append(line)
        
        if not updated:
            new_lines.append(f"{user_id},{expiry_time_str}\n")
        
        f.seek(0)
        f.truncate()
        f.writelines(new_lines)

async def get_subscription_end_time(user_id):
    try:
        with open('База_Бота/paid_users.txt', 'r') as f:
            for line in f:
                if line.strip():
                    paid_user_id, paid_expiry = line.strip().split(',')
                    if paid_user_id == str(user_id):
                        return "forever" if paid_expiry == "forever" else datetime.strptime(paid_expiry, '%Y-%m-%d %H:%M:%S')
    except FileNotFoundError:
        pass
    return None        
        
    
        
async def check_user_exists_in_db(user_id: int) -> bool:
    if not os.path.exists('База_Бота/paid_users.txt'):
        return False
        
    with open('База_Бота/paid_users.txt', 'r', encoding='utf-8') as f:
        for line in f:
            if line.startswith(f"{user_id},"):
                return True
    return False


@dp.callback_query_handler(lambda c: c.data == "check_subscription", state=UserStates.waiting_for_subscription)
async def check_subscription(callback_query: types.CallbackQuery, state: FSMContext):
    not_subscribed_channels = await check_all_subscriptions(callback_query.from_user.id)
    
    if not_subscribed_channels:
        keyboard = InlineKeyboardMarkup()
        for channel_name, channel_url in not_subscribed_channels.items():
            keyboard.add(InlineKeyboardButton(f"🔗 {channel_name}", url=channel_url))
        keyboard.add(InlineKeyboardButton("🔁 Проверить снова", callback_data="check_subscription"))
        
        await callback_query.answer("⚠️ Вы не подписаны на все каналы!", show_alert=True)
        await callback_query.message.edit_text(
            "📛 <b>Требуется подписка!</b>\n\n"
            + "\n".join([f"▪️ {channel}" for channel in not_subscribed_channels.keys()]),
            reply_markup=keyboard,
            parse_mode="HTML"
        )
    else:
        await callback_query.answer("🎉 Подписка подтверждена!", show_alert=False)
        await callback_query.message.edit_text(
            "<b>───── ⋆⋅☆⋅⋆ ─────</b>\n<blockquote>🎉 <b>Доступ открыт!</b>\n\n"
            "Теперь вы можете пользоваться всеми функциями бота!</blockquote>\n<b>───── ⋆⋅☆⋅⋆ ─────</b>",
            parse_mode="HTML"
        )
        await handle_welcome(
            user_id=callback_query.from_user.id,
            chat_id=callback_query.message.chat.id,
            from_user=callback_query.from_user,
            reply_photo_method=callback_query.message.reply_photo
        )
        
@dp.callback_query_handler(lambda c: c.data == 'send_welcome', state='*')
async def process_callback_send_welcome(callback_query: types.CallbackQuery, state: FSMContext):
    await state.finish()
    await handle_welcome(
        user_id=callback_query.from_user.id,
        chat_id=callback_query.message.chat.id,
        from_user=callback_query.from_user,
        reply_photo_method=callback_query.message.reply_photo
    )
    await callback_query.answer()

from aiogram.dispatcher import Dispatcher, FSMContext
from telethon.sessions import SQLiteSession

class VKOsintState:
    waiting_for_id = "waiting_for_vk_id"


@dp.callback_query_handler(lambda call: call.data in [
    'email_complaint', 'website_complaint', 'create_account', 'report_message',
    'channel_demolition', 'restore_account', 'spam_code',
    'email_spam', 'vk_osint', 'dox', 'regular_search', 'extended_search'
])
async def handle_callbacks(call: types.CallbackQuery, state: FSMContext):
    user_id = call.from_user.id
    banned_users = load_banned_users()
    
    if str(user_id) not in admin_chat_ids:
        if user_id in banned_users:
            await call.answer("🚨 Вы забанены", show_alert=True)
            return
        
        if (call.data not in ['spam_code', 'email_spam', 'vk_osint', 'dox', 'regular_search', 'extended_search'] and 
            not await check_payment(user_id)):
            await call.answer("⏳ Время доступа истекло", show_alert=True)
            await call.message.answer(
                "⏳ Оплатите подписку",
                reply_markup=InlineKeyboardMarkup().add(
                    InlineKeyboardButton("Оплатить", callback_data="go_to_payment")
                )
            )
            return
    
    if call.data == 'email_complaint':
        await call.message.answer('📧 Введите тему письма:')
        await ComplaintStates.subject.set()
    elif call.data == 'website_complaint':
        await call.message.answer('🌐 Введите текст для отправки на сайт:')
        await ComplaintStates.text_for_site.set()
    elif call.data == 'create_account':
        await call.message.answer('📱 Введите ваш номер телефона:')
        await CreateAccountStates.phone.set()
    elif call.data == 'report_message':
        await call.message.answer('🔗 Введите ID, @username или ссылку на сообщение:')
        await ReportStates.message_link.set()
    elif call.data == 'channel_demolition':
        await call.message.answer('🔗 Введите ссылку на канал:')
        await ChannelDemolitionStates.channel_link.set()
    elif call.data == 'restore_account':
        await call.message.answer('📱 Введите номер телефона для восстановления аккаунта:')
        await RestoreAccountStates.phone.set()
    elif call.data == 'go_to_payment':
        await call.message.answer("ℹ️ Выберите способ оплаты:", reply_markup=payment_keyboard)
    elif call.data == 'spam_code':
        await call.message.answer(
            '📞 Введите номера телефонов и количество отправлений для каждого\n'
            'Формат: +79991234567 10 (по одному номеру на строку)\n\n'
            'Пример:\n'
            '+79161234567 5\n'
            '+79031234567 3\n'
            '+79261234567 7'
        )
        await SpamCodeStates.waiting_for_numbers.set()
    elif call.data == 'email_spam':
        await call.message.answer("📧 Введите email получателя:")
        await EmailSpamStates.waiting_for_receiver.set()
    elif call.data == 'vk_osint':
        msg = await call.message.answer("🔍 Введите ID или ссылку на профиль ВКонтакте:")
        await state.update_data(
            vk_chain_start=msg.message_id,
            messages_to_delete=[msg.message_id]
        )
        await state.set_state(VKOsintState.waiting_for_id)
    elif call.data == 'dox':
        msg = await call.message.answer("🔎 Введите номер телефона (79637829051) или email (ceo@vkontakte.ru):")
        await state.update_data(
            dox_chain_start=msg.message_id,
            messages_to_delete=[msg.message_id]
        )
        await state.set_state(DoxState.waiting_for_input)
    elif call.data == 'regular_search':
        await call.answer()
        text = """<b>🔎 Введите данные для поиска:</b>        
<blockquote>
👤 <b>Поиск по имени/ФИО</b>
├  Иван 2000
├  Иван Иванов 01.01
├  Иванов Иван Иванович 01.01.2000
└  Иванов Иван Иванович Москва 2000

🚗 <b>Поиск по авто</b>
├  А001АА77 - поиск по Гос номеру
└  XTA212130T1186583 - поиск по VIN

🌐 <b>Поиск по контактным данным</b>
├  79221110500 - поиск по Телефону
├  ivanov@mail.ru - поиск по Почте
├  @username - поиск по Телеграм
└  Igrok777 - поиск по Логину

🏛 <b>Поиск по документам</b>
├  4616233456 / 4616 233456 - поиск по Паспорту
├  7707083893 - поиск по ИНН (ЮЛ / ФЛ)
├  00461487830 - поиск по СНИЛС
├  1027739099772 - поиск по ОГРН
└  II-МЮ №723184 - поиск по Паспорту СССР (только в многоуровневом поиске)
</blockquote>
<i>Для запроса просто введите данные, которые у вас есть на человека в формате представленном выше и отправьте их боту (рекомендуется искать только по одному направлению)</i>"""
        await call.message.answer(text, parse_mode="HTML")
        await state.update_data(search_type='regular')
        await state.set_state('waiting_probe_query')
    elif call.data == 'extended_search':
        await call.answer()
        text = """<blockquote><b>⚠️ МНОГОУРОВНЕВЫЙ ПОИСК ⚠️</b>

В отличии от обычного поиска, многоуровневый поиск ищет совпадения слов из запроса в любых столбиках, поэтому бот зачастую может выдавать нерелевантные данные в отчете, мы рекомендуем использовать его только для опытных пользователей и специфических поисковых запросов.
</blockquote>
<b>Для запроса просто введите данные, которые у вас есть на человека</b>"""
        await call.message.answer(text, parse_mode="HTML")
        await state.update_data(search_type='extended')
        await state.set_state('waiting_probe_query')
    
    await call.answer()



@dp.message_handler(state=DoxState.waiting_for_input)
async def process_dox_input(message: types.Message, state: FSMContext):
    data = await state.get_data()
    messages_to_delete = data.get('messages_to_delete', [])
    messages_to_delete.append(message.message_id)

    user_input = message.text.strip()
    is_phone = user_input.replace('+', '').isdigit()
    is_email = '@' in user_input and '.' in user_input.split('@')[-1]
    
    if not (is_phone or is_email):
        error_msg = await message.reply("❌ Ошибка: Нужен номер телефона или email")
        messages_to_delete.append(error_msg.message_id)
        await state.update_data(messages_to_delete=messages_to_delete)
        return

    dots_message = await message.answer("🔎 Ищу информацию...")
    messages_to_delete.append(dots_message.message_id)
    
    try:
        result_messages = await get_dox_data(user_input, message, dots_message)
        
        try:
            await dots_message.delete()
            messages_to_delete.remove(dots_message.message_id)
        except:
            pass

        if result_messages:
            for msg_text in result_messages[:-1]:
                msg = await message.answer(msg_text)
                messages_to_delete.append(msg.message_id)
            
            markup = InlineKeyboardMarkup().add(
                InlineKeyboardButton('❌ Скрыть', callback_data='delete_dox_chain')
            )
            last_msg = await message.answer(
                result_messages[-1],
                reply_markup=markup
            )
            messages_to_delete.append(last_msg.message_id)
            
    except Exception as e:
        error_msg = await message.answer(f"⚠️ Ошибка: {str(e)}")
        messages_to_delete.append(error_msg.message_id)
    
    await state.update_data(messages_to_delete=messages_to_delete)

@dp.callback_query_handler(lambda c: c.data == 'delete_dox_chain', state='*')
async def delete_dox_chain(callback_query: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    messages_to_delete = data.get('messages_to_delete', [])
    
    for msg_id in sorted(messages_to_delete, reverse=True):
        try:
            await bot.delete_message(callback_query.from_user.id, msg_id)
        except:
            pass
    
    await callback_query.answer("Результаты скрыты")
    await state.finish()

async def get_dox_data(target, message, dots_message):
    sessions_dir = "Не_основные"
    os.makedirs(sessions_dir, exist_ok=True)
    
    sessions = []
    for file in os.listdir(sessions_dir):
        if file.endswith('.session'):
            session_name = file.replace('.session', '')
            json_path = os.path.join(sessions_dir, f"{session_name}.json")
            if not os.path.exists(json_path):
                with open(json_path, 'w') as f:
                    json.dump({"name": session_name, "status": "Свободна"}, f)
            
            sessions.append({
                "name": session_name,
                "path": os.path.join(sessions_dir, file),
                "json_path": json_path,
                "api_id": API_ID,
                "api_hash": API_HASH
            })
    
    if not sessions:
        await dots_message.edit_text("❌ Активные сессии не найдены!")
        return []

    result_messages = []
    
    for session in sessions:
        try:
            with open(session["json_path"], 'r') as f:
                session_data = json.load(f)
            
            if session_data.get("status") != "Свободна":
                continue
            session_data["status"] = "Занята"
            with open(session["json_path"], 'w') as f:
                json.dump(session_data, f)
                
            client = None
            try:
                client = TelegramClient(
                    session=SQLiteSession(session["path"]),
                    api_id=session["api_id"],
                    api_hash=session["api_hash"]
                )
                await client.connect()
                
                if not await client.is_user_authorized():
                    session_data["status"] = "Свободна"
                    with open(session["json_path"], 'w') as f:
                        json.dump(session_data, f)
                    continue

                await client.send_message('Diablo_Lair_Bot', target)
                
                first_msg = None
                for _ in range(10):
                    messages = await client.get_messages('Diablo_Lair_Bot', limit=1)
                    if messages and messages[0].text and "Поиск" not in messages[0].text:
                        first_msg = messages[0]
                        break
                    await asyncio.sleep(0.5)
                
                if not first_msg:
                    session_data["status"] = "Свободна"
                    with open(session["json_path"], 'w') as f:
                        json.dump(session_data, f)
                    continue
                    
                await asyncio.sleep(3)
                messages = await client.get_messages('Diablo_Lair_Bot', limit=14)
                cleaned_messages = []
                for msg in sorted(messages, key=lambda m: m.id):
                    if msg.text:
                        clean_text = msg.text.replace('*', '')
                        cleaned_messages.append(clean_text)
                filtered_messages = []
                for msg in cleaned_messages:
                    if "Поиск" not in msg and "завершён" not in msg:
                        filtered_messages.append(msg)
                
                if filtered_messages:
                    result_messages.extend(filtered_messages[:12])
                    session_data["status"] = "Свободна"
                    with open(session["json_path"], 'w') as f:
                        json.dump(session_data, f)
                    break
                        
            except Exception as e:
                session_data["status"] = "Свободна"
                with open(session["json_path"], 'w') as f:
                    json.dump(session_data, f)
                continue
            finally:
                if client:
                    await client.disconnect()
        
        except Exception as e:
            continue
    
    if not result_messages:
        await dots_message.edit_text("❌ Не удалось получить данные. Попробуйте позже.")
        return []
    
    return result_messages

@dp.message_handler(state=VKOsintState.waiting_for_id)
async def process_vk_id(message: types.Message, state: FSMContext):
    data = await state.get_data()
    messages_to_delete = data.get('messages_to_delete', [])
    messages_to_delete.append(message.message_id)

    user_input = message.text
    vk_id = extract_vk_id(user_input)
    
    if not vk_id:
        error_msg = await message.reply("❌ Ошибка: Нужен ID")
        messages_to_delete.append(error_msg.message_id)
        await state.update_data(messages_to_delete=messages_to_delete)
        return

    dots_message = await message.answer("🔎 Ищу информацию...")
    messages_to_delete.append(dots_message.message_id)
    
    try:
        history_messages = await get_vk_history_data(vk_id, message, dots_message)
        
        try:
            await dots_message.delete()
            messages_to_delete.remove(dots_message.message_id)
        except:
            pass

        if history_messages:
            for msg_text in history_messages[:-1]:
                msg = await message.answer(msg_text)
                messages_to_delete.append(msg.message_id)
            
            markup = InlineKeyboardMarkup().add(
                InlineKeyboardButton('❌ Скрыть', callback_data='delete_vk_chain')
            )
            last_msg = await message.answer(
                history_messages[-1],
                reply_markup=markup
            )
            messages_to_delete.append(last_msg.message_id)
            
    except Exception as e:
        error_msg = await message.answer(f"⚠️ Ошибка: {str(e)}")
        messages_to_delete.append(error_msg.message_id)
    
    await state.update_data(messages_to_delete=messages_to_delete)

@dp.callback_query_handler(lambda c: c.data == 'delete_vk_chain', state='*')
async def delete_vk_chain(callback_query: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    messages_to_delete = data.get('messages_to_delete', [])
    
    for msg_id in sorted(messages_to_delete, reverse=True):
        try:
            await bot.delete_message(callback_query.from_user.id, msg_id)
        except:
            pass
    
    await callback_query.answer("Результаты скрыты")
    await state.finish()

def extract_vk_id(input_str):
    input_str = input_str.strip()
    
    if 'vk.com/' in input_str:
        input_str = input_str.split('vk.com/')[-1].split('/')[0].split('?')[0]
    
    if input_str.isdigit():
        return f"id{input_str}"  
    elif input_str.startswith('id') and input_str[2:].isdigit():
        return input_str
    elif input_str.startswith('public') and input_str[6:].isdigit():
        return input_str
    else:
        return None
            
    
async def get_vk_history_data(target, message, dots_message):
    sessions_dir = "Не_основные"
    os.makedirs(sessions_dir, exist_ok=True)
    
    sessions = []
    for file in os.listdir(sessions_dir):
        if file.endswith('.session'):
            session_name = file.replace('.session', '')
            json_path = os.path.join(sessions_dir, f"{session_name}.json")
            if not os.path.exists(json_path):
                with open(json_path, 'w') as f:
                    json.dump({"name": session_name, "status": "Свободна"}, f)
            
            sessions.append({
                "name": session_name,
                "path": os.path.join(sessions_dir, file),
                "json_path": json_path,
                "api_id": API_ID,
                "api_hash": API_HASH
            })
    
    if not sessions:
        await dots_message.edit_text("❌ Активные сессии не найдены!")
        return []

    result_messages = []
    
    for session in sessions:
        try:
            with open(session["json_path"], 'r') as f:
                session_data = json.load(f)
            
            if session_data.get("status") != "Свободна":
                continue
            session_data["status"] = "Занята"
            with open(session["json_path"], 'w') as f:
                json.dump(session_data, f)
                
            client = None
            try:
                client = TelegramClient(
                    session=SQLiteSession(session["path"]),
                    api_id=session["api_id"],
                    api_hash=session["api_hash"]
                )
                await client.connect()
                
                if not await client.is_user_authorized():
                    session_data["status"] = "Свободна"
                    with open(session["json_path"], 'w') as f:
                        json.dump(session_data, f)
                    continue

                await client.send_message('VKHistoryRobot', target)
                
                first_msg = None
                for _ in range(10):
                    messages = await client.get_messages('VKHistoryRobot', limit=1)
                    if messages and messages[0].text and "🕓 Поиск займёт" not in messages[0].text:
                        first_msg = messages[0]
                        break
                    await asyncio.sleep(0.5)
                
                if not first_msg:
                    session_data["status"] = "Свободна"
                    with open(session["json_path"], 'w') as f:
                        json.dump(session_data, f)
                    continue
                    
                await asyncio.sleep(3)
                messages = await client.get_messages('VKHistoryRobot', limit=14)
                cleaned_messages = []
                for msg in sorted(messages, key=lambda m: m.id):
                    if msg.text:
                        clean_text = msg.text.replace('*', '')
                        cleaned_messages.append(clean_text)
                filtered_messages = []
                for msg in cleaned_messages:
                    if "🕓 Поиск займёт" not in msg and "✅ Поиск завершён" not in msg:
                        filtered_messages.append(msg)
                
                if filtered_messages:
                    result_messages.extend(filtered_messages[:12])
                    session_data["status"] = "Свободна"
                    with open(session["json_path"], 'w') as f:
                        json.dump(session_data, f)
                    break
                        
            except Exception as e:
                session_data["status"] = "Свободна"
                with open(session["json_path"], 'w') as f:
                    json.dump(session_data, f)
                continue
            finally:
                if client:
                    await client.disconnect()
        
        except Exception as e:
            continue
    
    if not result_messages:
        await dots_message.edit_text("❌ Не удалось получить данные. Попробуйте позже.")
        return []
    
    return result_messages
import io 
from aiogram.types import InputFile  

async def animate_searching(message: types.Message):
    dots = ["", ".", "..", "..."]
    i = 0
    try:
        while True:
            await asyncio.sleep(0.5)
            await message.edit_text(f"🔎 Ищу информацию{dots[i]}")
            i = (i + 1) % len(dots)
    except:
        pass        

MAX_MESSAGE_LENGTH = 4000

async def edit_or_send_message(callback_query, text, markup=None):
    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=text,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=text,
            reply_markup=markup,
            parse_mode="HTML"
        )
import aiohttp
from bs4 import BeautifulSoup
import ipaddress
from fake_useragent import UserAgent
@dp.callback_query_handler(lambda c: c.data == 'osint', state='*')
async def osint_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    markup = InlineKeyboardMarkup(row_width=2)
    markup.add(
        InlineKeyboardButton('🔍Osint🔎', callback_data='probe_osint'),
        InlineKeyboardButton('🔍Пробив🔎', callback_data='probe_menu')
    )
    markup.add(
        InlineKeyboardButton('🔙 Назад', callback_data='to_start')
    )
    
    text = """
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🔍 OSINT/Пробыв 🔎</b> - набор инструмендов для пробива и osinta
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    await edit_or_send_message(callback_query, text, markup)
    


@dp.callback_query_handler(lambda c: c.data == 'probe_osint', state='*')
async def osint_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    markup = InlineKeyboardMarkup(row_width=2)
    markup.add(
        InlineKeyboardButton('🌐 Osint IP', callback_data='osint_ip'),
        InlineKeyboardButton('👤 Osint VK', callback_data='osint_vk'),
        #InlineKeyboardButton('🔍Osint Фио🔎', callback_data='probe_fio'),
    )
    markup.add(
        InlineKeyboardButton('📷 Камеры', callback_data='kamera_onlain')
    )
    markup.add(
        InlineKeyboardButton('🔙 Назад', callback_data='osint')
    )
    
    text = """
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🔍 OSINT Tools 🔎</b>
<b>🌐 Osint IP</b> - Базовая Osint по IP
<b>👤 Osint VK</b> - Поиск по ВКонтакте
<b>📷 Камеры</b> - Камеры онлайн
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    await edit_or_send_message(callback_query, text, markup)
    
    
@dp.callback_query_handler(lambda c: c.data == 'kamera_onlain', state='*')
async def osint_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    markup = InlineKeyboardMarkup(row_width=2)
    markup.add(
        InlineKeyboardButton('📷 Insecam', url='http://insecam.org/'),
        InlineKeyboardButton('📷 Earthcam', url='http://www.earthcam.com/')
    )
    markup.add(
        InlineKeyboardButton('🔙 Назад', callback_data='probe_osint')
    )
    
    text = """
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📷 Insecam</b> - Просмотр публичных камер наблюдения 
<b>📷 Earthcam</b> - Просмотр публичных камер наблюдения 
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    await edit_or_send_message(callback_query, text, markup)    

class OSINTStates(StatesGroup):
    waiting_for_ip = State()

@dp.callback_query_handler(lambda c: c.data == 'osint_ip', state='*')
async def osint_ip_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await bot.send_message(
        callback_query.from_user.id,
        "Введите IP-адрес для OSINT-разведки:"
    )
    await OSINTStates.waiting_for_ip.set()

@dp.message_handler(state=OSINTStates.waiting_for_ip)
async def process_ip_for_osint(message: types.Message, state: FSMContext):
    ip_address = message.text.strip()
    
    ipv4_pattern = re.compile(r"^((25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)\.){3}(25[0-5]|2[0-4][0-9]|[01]?[0-9][0-9]?)$")
    
    if not ipv4_pattern.match(ip_address):
        await message.reply("❌ Некорректный формат IP-адреса. Пожалуйста, введите действительный IPv4-адрес.")
        return

    await bot.send_message(message.from_user.id, f"🔍 Выполняю OSINT-разведку для IP: <b>{ip_address}</b>...", parse_mode="HTML")

    collected_data = {
        'Геолокация': {},
        'Сетевые данные': {},
        'Дополнительно': {},
        'Безопасность': {}
    }
    services_checked = []

    async def safe_fetch(session, url, service_name, json_mode=True):
        try:
            async with session.get(url) as response:
                if response.status != 200:
                    return None, service_name
                if json_mode:
                    return await response.json(), service_name
                return await response.text(), service_name
        except Exception:
            return None, service_name

    async def add_if_not_exists(target_dict, key, value):
        if value and (key not in target_dict or not target_dict[key]):
            target_dict[key] = value
            return True
        return False

    try:
        async with aiohttp.ClientSession() as session:
            tasks = [
                safe_fetch(session, f"http://ip-api.com/json/{ip_address}?lang=ru&fields=status,message,continent,continentCode,country,countryCode,region,regionName,city,zip,lat,lon,timezone,offset,isp,org,as,asname,reverse,mobile,proxy,hosting,query", "IP-API"),
                safe_fetch(session, f"https://ipinfo.io/{ip_address}/json", "IPinfo"),
                safe_fetch(session, f"https://api.ipgeolocation.io/ipgeo?apiKey=demo&ip={ip_address}", "IPGeolocation"),
                safe_fetch(session, f"https://ipapi.co/{ip_address}/json/", "ipapi.co")
            ]
            
            results = await asyncio.gather(*tasks)
            
            for data, service_name in results:
                if data is None:
                    continue
                    
                services_checked.append(service_name)
                
                geo = collected_data['Геолокация']
                net = collected_data['Сетевые данные']
                other = collected_data['Дополнительно']
                sec = collected_data['Безопасность']
                
                if service_name == "IP-API" and data.get('status') == 'success':
                    await add_if_not_exists(geo, 'Страна', f"{data.get('country', '')} ({data.get('countryCode', '')})" if data.get('country') else None)
                    await add_if_not_exists(geo, 'Регион', data.get('regionName'))
                    await add_if_not_exists(geo, 'Город', data.get('city'))
                    await add_if_not_exists(geo, 'Индекс', data.get('zip'))
                    if data.get('lat') and data.get('lon'):
                        await add_if_not_exists(geo, 'Координаты', f"{data.get('lat')}, {data.get('lon')}")
                        await add_if_not_exists(geo, 'Карта', f"http://www.google.com/maps/place/{data.get('lat')},{data.get('lon')}")
                    
                    await add_if_not_exists(other, 'Континент', data.get('continent'))
                    await add_if_not_exists(other, 'Часовой пояс', data.get('timezone'))
                    await add_if_not_exists(other, 'Смещение UTC', data.get('offset'))
                    
                    await add_if_not_exists(net, 'Провайдер (ISP)', data.get('isp'))
                    await add_if_not_exists(net, 'Организация', data.get('org'))
                    if data.get('as'):
                        await add_if_not_exists(net, 'Автономная система (AS)', f"{data.get('as')} {data.get('asname', '')}".strip())
                    await add_if_not_exists(net, 'Обратная запись DNS', data.get('reverse'))
                    if data.get('mobile') is not None:
                        await add_if_not_exists(net, 'Мобильная сеть', 'Да' if data.get('mobile') else 'Нет')
                    if data.get('proxy') is not None:
                        await add_if_not_exists(sec, 'Прокси', 'Да' if data.get('proxy') else 'Нет')
                    if data.get('hosting') is not None:
                        await add_if_not_exists(sec, 'Хостинг', 'Да' if data.get('hosting') else 'Нет')
                
                elif service_name == "IPinfo" and not data.get('error'):
                    await add_if_not_exists(geo, 'Страна', data.get('country'))
                    await add_if_not_exists(geo, 'Регион', data.get('region'))
                    await add_if_not_exists(geo, 'Город', data.get('city'))
                    await add_if_not_exists(geo, 'Индекс', data.get('postal'))
                    if data.get('loc'):
                        await add_if_not_exists(geo, 'Координаты', data.get('loc'))
                        await add_if_not_exists(geo, 'Карта', f"http://www.google.com/maps/place/{data.get('loc')}")
                    await add_if_not_exists(other, 'Часовой пояс', data.get('timezone'))
                    
                    await add_if_not_exists(net, 'Провайдер (ISP)', data.get('org'))
                    await add_if_not_exists(net, 'Обратная запись DNS', data.get('hostname'))
                
                elif service_name == "IPGeolocation" and data.get('ip'):
                    await add_if_not_exists(geo, 'Страна', f"{data.get('country_name', '')} ({data.get('country_code2', '')})" if data.get('country_name') else None)
                    await add_if_not_exists(geo, 'Город', data.get('city'))
                    if data.get('latitude') and data.get('longitude'):
                        await add_if_not_exists(geo, 'Координаты', f"{data.get('latitude')}, {data.get('longitude')}")
                        await add_if_not_exists(geo, 'Карта', f"http://www.google.com/maps/place/{data.get('latitude')},{data.get('longitude')}")
                    
                    await add_if_not_exists(net, 'Провайдер (ISP)', data.get('isp'))
                    await add_if_not_exists(net, 'Организация', data.get('organization'))
                    if data.get('asn'):
                        await add_if_not_exists(net, 'Автономная система (AS)', data.get('asn'))
                    await add_if_not_exists(net, 'Обратная запись DNS', data.get('reverse_dns'))
                
                elif service_name == "ipapi.co" and not data.get('error'):
                    await add_if_not_exists(geo, 'Страна', f"{data.get('country_name', '')} ({data.get('country', '')})" if data.get('country_name') else None)
                    await add_if_not_exists(geo, 'Регион', data.get('region'))
                    await add_if_not_exists(geo, 'Город', data.get('city'))
                    await add_if_not_exists(geo, 'Индекс', data.get('postal'))
                    if data.get('latitude') and data.get('longitude'):
                        await add_if_not_exists(geo, 'Координаты', f"{data.get('latitude')}, {data.get('longitude')}")
                        await add_if_not_exists(geo, 'Карта', f"http://www.google.com/maps/place/{data.get('latitude')},{data.get('longitude')}")
                    
                    await add_if_not_exists(other, 'Часовой пояс', data.get('timezone'))
                    
                    await add_if_not_exists(net, 'Провайдер (ISP)', data.get('org'))
                    if data.get('asn'):
                        await add_if_not_exists(net, 'Автономная система (AS)', data.get('asn'))

    except Exception as e:
        print(f"Ошибка при сборе данных: {e}")

    final_message_parts = ["<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote><b>Собранная информация по IP-адресу:</b>"]
    
    for category, details in collected_data.items():
        if details:
            final_message_parts.append(f"\n<b>{category}:</b>")
            for key, value in details.items():
                if value:
                    final_message_parts.append(f"  {key}: {value}")
    
    final_message_parts.append("</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>")
    
    if services_checked:
        final_message_parts.append(f"\n<i>Проверенные сервисы: {', '.join(services_checked)}</i>")
    else:
        final_message_parts.append("\n⚠ Не удалось получить данные ни от одного сервиса")

    final_message = "\n".join(final_message_parts)
    await bot.send_message(message.from_user.id, final_message, parse_mode="HTML")
    await state.finish()


@dp.callback_query_handler(lambda c: c.data == 'probe_fio', state='*')
async def probe_fio_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer() 
    text = "🔍 Введите ФИО для поиска в формате: <b>Фамилия Имя Отчество</b>"
    await callback_query.message.answer(text, parse_mode='HTML')
    await StateProbeFIO.fio.set()
    
from urllib.parse import quote
import aiohttp
from bs4 import BeautifulSoup
import re

class StateProbeFIO(StatesGroup):
    fio = State()

async def fetch_search_results(fio: str):
    encoded_fio = quote(fio)
    url = f"https://www.social-searcher.com/google-social-search/?q={encoded_fio}"
    
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
        'Accept-Language': 'ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7'
    }
    
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(url, headers=headers) as response:
                if response.status == 200:
                    html = await response.text()
                    soup = BeautifulSoup(html, 'html.parser')
                    
                    iframes = {
                        'facebook': soup.find('div', id='facebook-iframe').find('iframe')['src'],
                        'twitter': soup.find('div', id='twitter-iframe').find('iframe')['src'],
                        'instagram': soup.find('div', id='instagram-iframe').find('iframe')['src'],
                        'tiktok': soup.find('div', id='tiktok-iframe').find('iframe')['src'],
                        'linkedin': soup.find('div', id='linkedin-iframe').find('iframe')['src'],
                        'pinterest': soup.find('div', id='pinterest-iframe').find('iframe')['src']
                    }
                    fb_url = f"https://www.social-searcher.com{iframes['facebook']}"
                    async with session.get(fb_url, headers=headers) as fb_response:
                        if fb_response.status == 200:
                            fb_html = await fb_response.text()
                            fb_soup = BeautifulSoup(fb_html, 'html.parser')
                            
                            results = []
                            for item in fb_soup.select('.g')[:3]: 
                                title = item.select_one('h3')
                                if not title:
                                    continue
                                
                                title_text = title.get_text(strip=True)
                                link = title.find_parent('a')['href'] if title.find_parent('a') else '#'
                                url_elem = item.select_one('.r > a')
                                display_url = url_elem.get_text(strip=True) if url_elem else link
                                
                                desc = item.select_one('.st')
                                description = desc.get_text(strip=True) if desc else 'Нет описания'
                                description = re.sub(r'\s+', ' ', description)
                                
                                results.append(
                                    f"<b>🔹 {title_text}</b>\n"
                                    f"<code>{display_url[:50]}...</code>\n"
                                    f"<i>{description[:120]}...</i>\n"
                                    f"<a href='{link}'>🔗 Открыть</a>\n"
                                )
                            
                            return results if results else ["<i>По вашему запросу ничего не найдено</i>"]
                        return ["<i>Не удалось загрузить результаты поиска</i>"]
                return [f"<i>Ошибка получения данных (код {response.status})</i>"]
    except Exception as e:
        print(f"Error fetching results: {e}")
        return ["<i>Произошла ошибка при обработке запроса</i>"]

@dp.message_handler(state=StateProbeFIO.fio)
async def process_fio(message: types.Message, state: FSMContext):
    fio = message.text.strip()
    msg = await message.answer("<i>Ищем информацию в социальных сетях... 🔍</i>", parse_mode='HTML')
    search_results = await fetch_search_results(fio)
    markup = InlineKeyboardMarkup(row_width=1)
    encoded_fio = quote(fio)
    markup.add(
        InlineKeyboardButton(
            text="🌐 Показать все результаты на social-searcher.com", 
            url=f"https://www.social-searcher.com/google-social-search/?q={encoded_fio}"
        ),
        InlineKeyboardButton('🔙 Назад', callback_data='osint')
    )
    
    result_text = "\n".join([
        f"<b>🔍 Результаты поиска:</b> <code>{fio}</code>",
        "<b>Социальные сети:</b>",
        "",
        *search_results,
        "",
        "<i>Для просмотра всех результатов нажмите кнопку ниже ⤵️</i>"
    ])
    
    await msg.edit_text(
        result_text,
        reply_markup=markup,
        disable_web_page_preview=True,
        parse_mode='HTML'
    )
    
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'api_settings', state='*')
async def api_settings_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    
    config = await load_or_create_api_config()
    api_statuses = []
    
    if not config["api_tokens"]:
        balance_info = "❌ Нет доступных API токенов"
    else:
        for api in config["api_tokens"]:
            try:
                
                response = await make_api_request("profile", token=api)
                if response and not response.get('error'):
                    profile = response.get('profile', {})
                    api_statuses.append(f"✅ {profile.get('name', 'N/A')} - {profile.get('balance', 0)}")
                else:
                    error_message = str(response.get('error', 'Неизвестная ошибка'))
                    if 'Unauthorized' in error_message or '401' in error_message:
                        api_statuses.append(f"❌ {api[:10]}... - Неверный API токен")
                    else:
                        api_statuses.append(f"❌ {api[:10]}... - Ошибка API: {error_message}")
            except Exception as e:
                error_msg = str(e)
                if 'Unauthorized' in error_msg or '401' in error_msg:
                    api_statuses.append(f"❌ {api[:10]}... - Неверный API токен")
                else:
                    api_statuses.append(f"❌ {api[:10]}... - Ошибка запроса: {error_msg}")
        
        balance_info = "\n".join(api_statuses)
    
    markup = InlineKeyboardMarkup(row_width=1)
    markup.add(
        InlineKeyboardButton('➕ Добавить API', callback_data='add_api'),
        InlineKeyboardButton('➖ Удалить API', callback_data='delete_api'),
        InlineKeyboardButton('🔙 Назад', callback_data='admin_panel')
    )
    
    text = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>⚙️ Настройки API</b>
Количество API - {len(config["api_tokens"])} 
{balance_info}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    await edit_or_send_message(callback_query, text, markup)

async def make_api_request(endpoint, token=None):
    config = await load_or_create_api_config()
    base_url = "https://infosearch54321.xyz/api"
    if not token and not config["api_tokens"]:
        return {"error": "No API tokens available"}
    api_token = token if token else random.choice(config["api_tokens"])
    url = f"{base_url}/{api_token}/{endpoint}"
    async with aiohttp.ClientSession() as session:
        async with session.get(url) as response:
            if response.status == 200:
                return await response.json()
            return {"error": f"API error: {response.status}"}

async def load_or_create_api_config():
    config_path = os.path.join('База_Бота', 'api_config.json')
    default_config = {"api_tokens": []}
    try:
        if not os.path.exists(config_path):
            with open(config_path, 'w', encoding='utf-8') as f:
                json.dump(default_config, f, indent=4, ensure_ascii=False)
            return default_config
        with open(config_path, 'r', encoding='utf-8') as f:
            config = json.load(f)
            if "api_tokens" not in config:
                with open(config_path, 'w', encoding='utf-8') as f:
                    json.dump(default_config, f, indent=4, ensure_ascii=False)
                return default_config
            return config
    except (json.JSONDecodeError, IOError):
        with open(config_path, 'w', encoding='utf-8') as f:
            json.dump(default_config, f, indent=4, ensure_ascii=False)
        return default_config


@dp.callback_query_handler(lambda c: c.data == 'delete_api', state='*')
async def delete_api_callback(callback_query: types.CallbackQuery, state: FSMContext):
    try:
        config = await load_or_create_api_config()
        if not config["api_tokens"]:
            await callback_query.answer("ℹ️ Нет сохраненных API токенов", show_alert=True)
            return
        
        await callback_query.answer()  
        
        markup = InlineKeyboardMarkup(row_width=1)
        for token in config["api_tokens"]:
            try:
                response = await make_api_request("profile", token=token)
                
                if response:
                    if response.get('error'):
                        if 'Unothorized' in str(response.get('error')) or '401' in str(response.get('error')):
                            btn_text = "❌ Ошибка не рабочый api"
                        else:
                            btn_text = f"❌ {response.get('error')}"
                    else:
                        btn_text = response.get('profile', {}).get('name', '❌ Неверный формат ответа')
                else:
                    btn_text = "❌ Нет ответа от сервера"
                    
            except Exception as e:
                if 'Unothorized' in str(e) or '401' in str(e):
                    btn_text = "❌ Ошибка не рабочый api"
                else:
                    btn_text = f"❌ Ошибка: {str(e)}"
            
            markup.add(InlineKeyboardButton(
                text=btn_text,
                callback_data=f"delete_token_{token}"
            ))

        markup.add(InlineKeyboardButton("🔙 Назад", callback_data="api_settings"))

        try:
            if callback_query.message.caption:
                await bot.edit_message_caption(
                    chat_id=callback_query.message.chat.id,
                    message_id=callback_query.message.message_id,
                    caption="🔴 Выберите API для удаления:",
                    reply_markup=markup
                )
            else:
                await bot.edit_message_text(
                    chat_id=callback_query.message.chat.id,
                    message_id=callback_query.message.message_id,
                    text="🔴 Выберите API для удаления:",
                    reply_markup=markup,
                    parse_mode="HTML"
                )
        except:
            await callback_query.message.answer(
                "🔴 Выберите API для удаления:",
                reply_markup=markup
            )

    except Exception as e:
        await callback_query.answer(f"❌ Ошибка: {str(e)}", show_alert=True)

@dp.callback_query_handler(lambda c: c.data.startswith('delete_token_'), state='*')
async def delete_specific_token(callback_query: types.CallbackQuery):
    token_to_delete = callback_query.data.replace('delete_token_', '')
    await callback_query.answer()
    
    try:
        config = await load_or_create_api_config()
        if token_to_delete not in config["api_tokens"]:
            await callback_query.answer("⚠️ Токен не найден", show_alert=True)
            return

        try:
            response = await make_api_request("profile", token=token_to_delete)
            if response:
                if response.get('error'):
                    if 'Unauthorized' in str(response.get('error')) or '401' in str(response.get('error')):
                        api_info = "❌ Нерабочий API (Ошибка авторизации)"
                    else:
                        api_info = f"❌ Нерабочий API ({response.get('error')})"
                else:
                    api_name = response.get('profile', {}).get('name', 'Без названия')
                    api_info = f"{api_name}"
            else:
                api_info = "❌ Нерабочий API (Нет ответа)"
        except Exception as e:
            if 'Unauthorized' in str(e) or '401' in str(e):
                api_info = "❌ Нерабочий API (Ошибка авторизации)"
            elif 'timed out' in str(e).lower():
                api_info = "❌ Нерабочий API (Таймаут соединения)"
            else:
                api_info = f"❌ Нерабочий API ({str(e)})"

        confirm_markup = InlineKeyboardMarkup(row_width=2)
        confirm_markup.add(
            InlineKeyboardButton("✅ Да", callback_data=f"confirm_delete_{token_to_delete}"),
            InlineKeyboardButton("❌ Нет", callback_data="delete_api")
        )

        try:
            if callback_query.message.caption:
                await bot.edit_message_caption(
                    chat_id=callback_query.message.chat.id,
                    message_id=callback_query.message.message_id,
                    caption=f"⚠️ Удалить этот API?\n{api_info}",
                    reply_markup=confirm_markup
                )
            else:
                await bot.edit_message_text(
                    chat_id=callback_query.message.chat.id,
                    message_id=callback_query.message.message_id,
                    text=f"⚠️ Удалить этот API?\n{api_info}",
                    reply_markup=confirm_markup,
                    parse_mode="HTML"
                )
        except:
            await callback_query.message.answer(
                f"⚠️ Удалить этот API?\n{api_info}",
                reply_markup=confirm_markup
            )

    except Exception as e:
        await callback_query.answer(f"❌ Ошибка: {str(e)}", show_alert=True)

@dp.callback_query_handler(lambda c: c.data.startswith('confirm_delete_'), state='*')
async def confirm_delete_token(callback_query: types.CallbackQuery):
    token = callback_query.data.replace('confirm_delete_', '')
    await callback_query.answer()
    
    try:
        config = await load_or_create_api_config()
        if token not in config["api_tokens"]:
            await callback_query.answer("⚠️ Токен не найден", show_alert=True)
            return
        api_status = ""
        api_name = "API"
        try:
            response = await make_api_request("profile", token=token)
            if response:
                if response.get('error'):
                    if 'Unauthorized' in str(response.get('error')) or '401' in str(response.get('error')):
                        api_status = " (Нерабочий API: Ошибка авторизации)"
                    else:
                        api_status = f" (Нерабочий API: {response.get('error')})"
                else:
                    api_name = response.get('profile', {}).get('name', 'API')
            else:
                api_status = " (Нерабочий API: Нет ответа от сервера)"
        except Exception as e:
            if 'Unauthorized' in str(e) or '401' in str(e):
                api_status = " (Нерабочий API: Ошибка авторизации)"
            elif 'timed out' in str(e).lower():
                api_status = " (Нерабочий API: Таймаут соединения)"
            else:
                api_status = f" (Нерабочий API: {str(e)})"
        config["api_tokens"].remove(token)
        with open("База_Бота/api_config.json", 'w', encoding='utf-8') as f:
            json.dump(config, f, indent=4, ensure_ascii=False)
        result_message = f"✅ {api_name}{api_status} успешно удален"
        back_markup = InlineKeyboardMarkup()
        back_markup.add(InlineKeyboardButton("🔙 Назад", callback_data="api_settings"))

        try:
            if callback_query.message.caption:
                await bot.edit_message_caption(
                    chat_id=callback_query.message.chat.id,
                    message_id=callback_query.message.message_id,
                    caption=result_message,
                    reply_markup=back_markup
                )
            else:
                await bot.edit_message_text(
                    chat_id=callback_query.message.chat.id,
                    message_id=callback_query.message.message_id,
                    text=result_message,
                    reply_markup=back_markup,
                    parse_mode="HTML"
                )
        except:
            await callback_query.message.answer(
                result_message, 
                reply_markup=back_markup
            )

    except Exception as e:
        await callback_query.answer(f"❌ Ошибка удаления: {str(e)}", show_alert=True)

@dp.callback_query_handler(lambda c: c.data == 'add_api', state='*')
async def add_api_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("Введите API токен:")
    await state.set_state("waiting_for_api")

@dp.message_handler(state="waiting_for_api")
async def process_api_input(message: types.Message, state: FSMContext):
    try:
        new_tokens = [t.strip() for t in message.text.split('\n') if t.strip()]
        config = await load_or_create_api_config()
        
        added = 0
        existing = 0
        
        for token in new_tokens:
            if token not in config["api_tokens"]:
                config["api_tokens"].append(token)
                added += 1
            else:
                existing += 1
        
        with open("База_Бота/api_config.json", "w", encoding='utf-8') as f:
            json.dump(config, f, indent=4, ensure_ascii=False)
        
        msg = f"✅ Добавлено новых токенов: {added}"
        if existing > 0:
            msg += f"\n⚠️ Пропущено (уже есть): {existing}"
            
        await message.answer(msg)
        await state.finish()   
    except Exception as e:
        await message.answer(f"❌ Ошибка: {e}")

@dp.callback_query_handler(lambda c: c.data == 'probe_menu', state='*')
async def probe_menu_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    markup = InlineKeyboardMarkup(row_width=2)
    markup.add(
        InlineKeyboardButton('🔍Обычный поиск🔎', callback_data='regular_search'),
        InlineKeyboardButton('🔍Многоуровневый поиск🔎', callback_data='extended_search'),
        InlineKeyboardButton('🔍Поиск по файлам🔎', callback_data='file_search'),
        InlineKeyboardButton('🕵️‍♂️ Щерлок', callback_data='sherlock_search'),
        InlineKeyboardButton('🔍VK Пробив🔎', callback_data='vk_osint'),
    )
    markup.add(InlineKeyboardButton('🔙 Назад', callback_data='osint'))
    
    text = """
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🔎 Выберите тип поиска</b>
├─ <b>Обычный поиск</b> - быстрый поиск
├─ <b>Многоуровневый</b> - глубокий анализ
├─ <b>Поиск по файлам</b> - поиск в базе данных
└─ <b>Щерлок</b> - поиск по username в соцсетях

</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>"""
    await edit_or_send_message(callback_query, text, markup)


@dp.callback_query_handler(lambda c: c.data == 'sherlock_search', state='*')
async def sherlock_search_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    response_text = """<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🕵️ Личность:</b>
<code>Навальный Алексей Анатольевич 04.06.1976</code> - ФИО

<b>📲 Контакты:</b>
<code>79637829051</code> – номер телефона
<code>ceo@vkontakte.ru</code> – email

<b>🚘 Транспорт:</b>
<code>В395ОК199</code> – номер автомобиля

<b>💬 Социальные сети:</b>
<code>vk.com/sherlock</code> – Вконтакте
<code>tiktok.com/@sherlock</code> – Tiktok
<code>instagram.com/sherlock</code> – Instagram
<code>ok.ru/profile/58460</code> – Одноклассники

<b>📟 Telegram:</b>
<code>@sherlock</code>, <code>tg123456</code> – логин или ID
<i>Можете переслать сообщение – попробую определить ID сам</i>

<b>📄 Документы:</b>
<code>/vu 1234567890</code> – водительские права
<code>/passport 1234567890</code> – паспорт
<code>/snils 12345678901</code> – СНИЛС
<code>/inn 123456789012</code> – ИНН

<b>🌐 Онлайн-следы:</b>
<code>/tag хирург москва</code> – поиск по телефонным книгам
<code>sherlock.com</code> / <code>1.1.1.1</code> – домен или IP
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    await callback_query.message.answer(response_text, parse_mode="HTML")
    await state.set_state("waiting_for_sherlock_username")
    

sessions_dir = "Не_основные"  


async def find_available_session(status_msg: types.Message):
    os.makedirs(sessions_dir, exist_ok=True)
    sessions = []
    for file in os.listdir(sessions_dir):
        if file.endswith('.session'):
            session_name = file.replace('.session', '')
            json_path = os.path.join(sessions_dir, f"{session_name}.json")
            if not os.path.exists(json_path):
                with open(json_path, 'w') as f:
                    json.dump({"name": session_name, "status": "Свободна", "user_id": None}, f)
            sessions.append({
                "name": session_name,
                "path": os.path.join(sessions_dir, file),
                "api_id": API_ID,
                "api_hash": API_HASH
            })
    if not sessions:
        await status_msg.edit_text("❌ Активные сессии не найдены!")
        return None    
    for session in sessions:
        session_name = session["name"]
        json_path = os.path.join(sessions_dir, f"{session_name}.json")
        try:
            with open(json_path, 'r') as f:
                session_data = json.load(f)
            if session_data.get("status") == "Свободна" and session_data.get("user_id") is None:
                temp_client = None
                try:
                    temp_client = TelegramClient(
                        session=SQLiteSession(session["path"]),
                        api_id=session["api_id"],
                        api_hash=session["api_hash"]
                    )
                    await temp_client.connect()
                    if not await temp_client.is_user_authorized():
                        continue
                except Exception:
                    continue
                finally:
                    if temp_client:
                        await temp_client.disconnect()
                session_data["status"] = "Занята"
                session_data["user_id"] = status_msg.chat.id
                with open(json_path, 'w') as f:
                    json.dump(session_data, f)
                return session
        except Exception:
            continue
    await status_msg.edit_text("❌ Нет доступных сессий!")
    return None

async def connect_to_session(session_info):
    client = TelegramClient(
        session=SQLiteSession(session_info["path"]),
        api_id=session_info["api_id"],
        api_hash=session_info["api_hash"]
    )
    await client.connect()
    return client, session_info["name"]

async def release_session(client, session_name):
    if client:
        await client.disconnect()  
    if session_name:
        json_path = os.path.join(sessions_dir, f"{session_name}.json")
        if os.path.exists(json_path):
            with open(json_path, 'r') as f:
                session_data = json.load(f)
            session_data["status"] = "Свободна"
            session_data["user_id"] = None
            with open(json_path, 'w') as f:
                json.dump(session_data, f)

async def send_sherlock_request(client, username):
    try:
        bot_entity = await client.get_entity("Sherlock1vl_bot")
        await client.send_message(bot_entity, username)
        await asyncio.sleep(5)
        history = await client.get_messages(bot_entity, limit=10)
        for msg in history:
            if not msg.out:
                return msg
        return None
    except Exception:
        return None

@dp.message_handler(state="waiting_for_sherlock_username")
async def process_sherlock_username(message: types.Message, state: FSMContext):
    username = message.text.strip()    
    status_msg = await message.answer("🔍 Идёт поиск...")

    session_info = await find_available_session(status_msg)
    if not session_info:
        await state.finish()
        return

    await state.update_data(current_sherlock_session_name=session_info["name"])

    client, session_name = None, None
    try:
        client, session_name = await connect_to_session(session_info)
        if not client:            
            await status_msg.edit_text("❌ Не удалось подключиться к сессии.")
            await state.finish()
            return        
        bot_response = await send_sherlock_request(client, username)
        if not bot_response:
            await status_msg.edit_text("❌ Бот не ответил на запрос.")
            await release_session(client, session_name)
            await state.finish()
            return

        response_text = getattr(bot_response, 'text', '')
        cleaned_text = response_text.replace('_', '').replace('*', '').replace('`', '')
        
        if "Обнаружен логин" in cleaned_text and "Выберите направление" in cleaned_text:
            markup = InlineKeyboardMarkup(row_width=2)
            markup.add(
                InlineKeyboardButton('ВКонтакте', callback_data=f'sherlock:ВКонтакте:{username}'),
                InlineKeyboardButton('Instagram', callback_data=f'sherlock:Instagram:{username}'),
                InlineKeyboardButton('TikTok', callback_data=f'sherlock:TikTok:{username}'),
                InlineKeyboardButton('Telegram', callback_data=f'sherlock:Telegram:{username}')
            )
            await message.answer(cleaned_text, reply_markup=markup) 
            await status_msg.delete()

        else:
            await status_msg.edit_text(cleaned_text)
            await release_session(client, session_name)
            await state.finish()

    except Exception as e:
        await status_msg.edit_text(f"❌ Ошибка: {str(e)}")
        await release_session(client, session_name)
        await state.finish()
    finally:
        pass

@dp.callback_query_handler(lambda c: c.data.startswith('sherlock:'), state="*")
async def process_sherlock_button(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    data = callback_query.data.split(':')
    button_label = data[1]
    username = data[2]
    user_id = callback_query.from_user.id

    user_data = await state.get_data()
    session_name_from_context = user_data.get('current_sherlock_session_name')

    client, current_session_name = None, None
    session_info = None

    if session_name_from_context:
        json_path = os.path.join(sessions_dir, f"{session_name_from_context}.json")
        if os.path.exists(json_path):
            with open(json_path, 'r') as f:
                session_data = json.load(f)
            
            if session_data.get("status") == "Занята" and session_data.get("user_id") == user_id:
                session_info = {
                    "name": session_name_from_context,
                    "path": os.path.join(sessions_dir, f"{session_name_from_context}.session"),
                    "api_id": API_ID,
                    "api_hash": API_HASH
                }
            else:
                await callback_query.message.edit_text("❌ Ваша предыдущая сессия стала недоступна или была освобождена. Пожалуйста, начните поиск заново.")
                await state.finish()
                return
        else:
            await callback_query.message.edit_text("❌ Ваша предыдущая сессия больше не существует. Пожалуйста, начните поиск заново.")
            await state.finish()
            return
    else:
        await callback_query.message.edit_text("❌ Контекст сессии утерян. Пожалуйста, начните поиск заново.")
        await state.finish()
        return
    
    try:
        client, current_session_name = await connect_to_session(session_info)
        if not client:
            await callback_query.message.edit_text("❌ Не удалось подключиться к вашей привязанной сессии.")
            await release_session(client, current_session_name)
            await state.finish()
            return
        
        bot_entity = await client.get_entity("Sherlock1vl_bot")
        history = await client.get_messages(bot_entity, limit=10)
        target_msg = None
        for msg in history:
            if not msg.out and hasattr(msg, 'reply_markup') and msg.reply_markup and username in msg.text:
                target_msg = msg
                break
        
        if not target_msg:
            await callback_query.message.edit_text("❌ Не найдено сообщение с кнопками для взаимодействия в вашей сессии. Возможно, диалог устарел. Пожалуйста, начните заново.")
            await release_session(client, current_session_name)
            await state.finish()
            return
        
        button_index = -1
        found = False
        for row in target_msg.reply_markup.rows:
            for btn in row.buttons:
                if btn.text == button_label:
                    found = True
                    break
                button_index += 1
            if found:
                break
        
        if not found:
            await callback_query.message.edit_text("❌ Не найдена соответствующая кнопка в исходном сообщении бота.")
            return
        
        await target_msg.click(button_index)
        await asyncio.sleep(5)
        new_history = await client.get_messages(bot_entity, limit=10)
        new_bot_response = None
        for msg in new_history:
            if not msg.out and msg.id != target_msg.id:
                new_bot_response = msg
                break
        
        if not new_bot_response:
            await callback_query.message.edit_text("❌ Бот не ответил после нажатия кнопки.")
            await release_session(client, current_session_name)
            await state.finish()
            return

        response_text = getattr(new_bot_response, 'text', '')
        cleaned_text = response_text.replace('_', '').replace('*', '').replace('`', '')
        markup = None
        # ниже я не доработал забейте хуй
        if "Общая информация" in cleaned_text and "Email" in cleaned_text and "IP адрес" in cleaned_text:
            markup = InlineKeyboardMarkup(row_width=2)
            markup.add(
                InlineKeyboardButton('Общая информация', callback_data=f'sherlock:общая_информация:{username}'),
                InlineKeyboardButton('Email', callback_data=f'sherlock:email:{username}'),
                InlineKeyboardButton('IP адрес', callback_data=f'sherlock:ip_адрес:{username}'),
                InlineKeyboardButton('Геолокация', callback_data=f'sherlock:геолокация:{username}')
            )
        elif "Другая группа кнопок" in cleaned_text:
            markup = InlineKeyboardMarkup(row_width=1)
            markup.add(
                InlineKeyboardButton('Кнопка А', callback_data=f'sherlock:кнопка_а:{username}'),
                InlineKeyboardButton('Кнопка Б', callback_data=f'sherlock:кнопка_б:{username}')
            )
        
        if markup:
            await callback_query.message.answer(cleaned_text, reply_markup=markup)
            await callback_query.message.delete()
        else:
            await callback_query.message.edit_text(cleaned_text)
            await release_session(client, current_session_name)
            await state.finish()
    except Exception as e:
        await callback_query.message.edit_text(f"❌ Произошла ошибка во время взаимодействия: {str(e)}")
        await release_session(client, current_session_name)
        await state.finish()


@dp.callback_query_handler(lambda c: c.data == 'file_search', state='*')
async def file_search_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    msg = await callback_query.message.answer("🔍 Введите запрос для поиска по файлам в базе данных:")
    await state.set_data({"search_message_id": msg.message_id})
    await state.set_state("waiting_for_file_search_query")


def search_data_in_file(file_path: str, query: str) -> list:
    results = []
    query_lower = query.lower()
    file_extension = os.path.splitext(file_path)[1].lower()

    if file_extension == '.txt':
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            for line_num, line in enumerate(f, 1):
                if query_lower in line.lower():
                    results.append({"Файл": os.path.basename(file_path), "Строка": line_num, "Содержимое": line.strip()})
    elif file_extension == '.csv':
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.DictReader(f)
                for row_num, row in enumerate(reader, 1):
                    for key, value in row.items():
                        if value and query_lower in str(value).lower():
                            result_data = {
                                "Файл": os.path.basename(file_path),
                                "Строка CSV": row_num + 1
                            }
                            result_data.update(row)
                            results.append(result_data)
                            break
        except Exception:
            pass
    elif file_extension == '.xlsx':
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(file_path)
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                for row_num, row in enumerate(sheet.iter_rows(), 1):
                    row_matched = False
                    for cell in row:
                        if cell.value and query_lower in str(cell.value).lower():
                            result_data = {
                                "Файл": os.path.basename(file_path),
                                "Лист": sheet_name,
                                "Строка Excel": row_num
                            }
                            row_values = {sheet.cell(row=1, column=col.column).value or f"Колонка_{col.column}": col.value for col in row}
                            result_data.update(row_values)
                            results.append(result_data)
                            row_matched = True
                            break
                    if row_matched and len(results) >= 15:
                        break
                if len(results) >= 15:
                    break
        except Exception:
            pass
    elif file_extension == '.json':
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                data = json.load(f)
                json_string = json.dumps(data, ensure_ascii=False).lower()
                if query_lower in json_string:
                     results.append({"Файл": os.path.basename(file_path), "Содержимое": json_string[:200] + "..."})
        except Exception:
            pass
    elif file_extension == '.xml':
        try:
            from lxml import etree
            tree = etree.parse(file_path)
            for element in tree.iter():
                if element.text and query_lower in element.text.lower():
                    results.append({"Файл": os.path.basename(file_path), "Элемент": element.tag, "Содержимое": element.text.strip()})
        except Exception:
            pass
    elif file_extension == '.html':
        try:
            from bs4 import BeautifulSoup
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                soup = BeautifulSoup(f, 'html.parser')
                text_content = soup.get_text().lower()
                if query_lower in text_content:
                    results.append({"Файл": os.path.basename(file_path), "Содержимое": text_content[:200] + "..."})
        except Exception:
            pass
    return results


async def animate_search(chat_id, message_id):
    dots = ["", ".", "..", "..."]
    while True:
        for dot in dots:
            try:
                await bot.edit_message_text(
                    chat_id=chat_id,
                    message_id=message_id,
                    text=f"<b>🔍 Поиск{dot}</b>",
                    parse_mode="HTML"
                )
                await asyncio.sleep(0.5)
            except:
                return
from googletrans import Translator
@dp.message_handler(state="waiting_for_file_search_query")
async def process_file_search(message: types.Message):
    search_query = message.text.strip()
    database_dir = "База_Данных"
    
    search_msg = await message.answer("<b>🔍 Идет поиск...</b>", parse_mode="HTML")
    
    try:
        all_results = []
        supported_extensions = ('.csv', '.xlsx', '.txt', '.json', '.xml', '.html')
        for root, dirs, files in os.walk(database_dir):
            for file in files:
                if file.lower().endswith(supported_extensions):
                    file_path = os.path.join(root, file)
                    try:
                        with open(file_path, 'r', encoding='utf-8') as f:
                            headers = f.readline().strip().split(',')
                            for line_num, line in enumerate(f, 2):
                                if search_query.lower() in line.lower():
                                    parts = line.strip().split(',')
                                    content = []
                                    for i, part in enumerate(parts):
                                        if part:
                                            header = headers[i] if i < len(headers) else f"Column_{i+1}"
                                            content.append(f"{header}: {part}")
                                    all_results.append({
                                        'file': file,
                                        'content': content
                                    })
                    except Exception as e:
                        print(f"Ошибка чтения файла {file_path}: {e}")
        
        if not all_results:
            await search_msg.edit_text(
                "<b>🔍 По вашему запросу ничего не найдено.</b>",
                parse_mode="HTML"
            )
        elif len(all_results) <= 5:
            response = "<b>───── ⋆⋅☆⋅⋆ ─────</b>\n<blockquote>"
            response += "<b>🔍 Результаты поиска</b>\n"
            response += f"└─ <b>Запрос:</b> <code>{html.escape(search_query)}</code>\n\n"
            
            for result in all_results:
                response += f"📄 <b>Файл:</b> <code>{html.escape(result['file'])}</code>\n"
                for i, item in enumerate(result['content']):
                    if i == len(result['content']) - 1:
                        response += f"              └─ <code>{html.escape(item)}</code>\n"
                    else:
                        response += f"              ├─ <code>{html.escape(item)}</code>\n"
                response += "\n"
            
            response += f"<b>Найдено:</b> {len(all_results)}\n"
            response += "</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>"
            
            await search_msg.edit_text(response, parse_mode="HTML")
        else:
            txt_content = f"Результаты поиска: {search_query}\n"
            txt_content += f"Всего найдено: {len(all_results)} записей\n\n"
            
            for result in all_results:
                txt_content += f"=== Файл: {result['file']} ===\n"
                for item in result['content']:
                    txt_content += f"{item}\n"
                txt_content += "\n"
            temp_file = io.BytesIO(txt_content.encode('utf-8'))
            temp_file.name = f"Результаты поиска {search_query}.txt"            
            await bot.send_document(
                chat_id=message.chat.id,
                document=temp_file,
                caption=f"<b>🔍 Найдено {len(all_results)} результатов по запросу:</b> <code>{html.escape(search_query)}</code>",
                parse_mode="HTML"
            )            
            temp_file.close()
            
            await search_msg.delete()
    
    except Exception as e:
        await search_msg.edit_text(
            f"❌ <b>Ошибка поиска:</b> <code>{html.escape(str(e))}</code>",
            parse_mode="HTML"
        )
    
from aiogram.types import InputFile

DATA_FILE = 'query_response_data.json'

def load_data():
    try:
        with open(DATA_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return {"queries": {}, "responses": {}, "unique_users_per_response": {}, "user_stats": {}}

def save_data(data):
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=4, ensure_ascii=False)

def update_interaction_data(user_id, query_text, response_data, matched_values):
    data = load_data()
    
    if str(user_id) not in data["user_stats"]:
        data["user_stats"][str(user_id)] = {"queries": 0, "interests": 0}
    data["user_stats"][str(user_id)]["queries"] += 1

    if query_text not in data["queries"]:
        data["queries"][query_text] = []
    if str(user_id) not in data["queries"][query_text]:
        data["queries"][query_text].append(str(user_id))

    for item in response_data:
        if isinstance(item, dict):
            response_values = tuple(sorted((k, v) for k, v in item.items() if k != 'database'))
            response_key = str(response_values)

            if response_key not in data["responses"]:
                data["responses"][response_key] = {"data": item, "users": []}
            if str(user_id) not in data["responses"][response_key]["users"]:
                data["responses"][response_key]["users"].append(str(user_id))
                data["user_stats"][str(user_id)]["interests"] += 1

            for key, value in item.items():
                normalized_value = str(value).strip().lower()
                if normalized_value:
                    if normalized_value not in data["unique_users_per_response"]:
                        data["unique_users_per_response"][normalized_value] = {"matched_data": item, "users": []}
                    if str(user_id) not in data["unique_users_per_response"][normalized_value]["users"]:
                        data["unique_users_per_response"][normalized_value]["users"].append(str(user_id))
        else:
            response_key = str(item)
            if response_key not in data["responses"]:
                data["responses"][response_key] = {"data": item, "users": []}
            if str(user_id) not in data["responses"][response_key]["users"]:
                data["responses"][response_key]["users"].append(str(user_id))
                data["user_stats"][str(user_id)]["interests"] += 1

            normalized_value = str(item).strip().lower()
            if normalized_value:
                if normalized_value not in data["unique_users_per_response"]:
                    data["unique_users_per_response"][normalized_value] = {"matched_data": item, "users": []}
                if str(user_id) not in data["unique_users_per_response"][normalized_value]["users"]:
                    data["unique_users_per_response"][normalized_value]["users"].append(str(user_id))

    for value in matched_values:
        normalized_value = str(value).strip().lower()
        if normalized_value:
            if normalized_value not in data["unique_users_per_response"]:
                data["unique_users_per_response"][normalized_value] = {"matched_data": None, "users": []}
            if str(user_id) not in data["unique_users_per_response"][normalized_value]["users"]:
                data["unique_users_per_response"][normalized_value]["users"].append(str(user_id))

    save_data(data)

def get_interest_count_for_value(value):
    data = load_data()
    normalized_value = str(value).strip().lower()
    if normalized_value in data["unique_users_per_response"]:
        return len(data["unique_users_per_response"][normalized_value]["users"])
    return 0

async def send_large_message(bot, chat_id, text, original_query, sources_list=None, user_id=None, interest_overall_count=0):
    if len(text) <= MAX_MESSAGE_LENGTH:
        await bot.send_message(chat_id, text, parse_mode="HTML")
    else:
        file_content = f"Результаты по запросу: {original_query}\n\n"
        if sources_list:
            file_content += "Список источников:\n"
            for source in sources_list:
                file_content += f"- {source}\n"
            file_content += "\n"
        
        file_content += text.replace('<b>', '').replace('</b>', '').replace('<code>', '').replace('</code>', '').replace('<blockquote>', '').replace('</blockquote>', '')
        
        file_content += f"\nДанными интересовались: {interest_overall_count} пользователь(ей)"
        
        file = io.BytesIO(file_content.encode('utf-8'))
        file.name = f"results_{original_query}.txt"
        caption = f"<b>🔍 Результаты по запросу:</b> <code>{original_query}</code>"
        if sources_list:
            caption += "\n<b>📁 Источники:</b> " + ", ".join([f"<code>{s}</code>" for s in sources_list])
        
        caption += f"\n\n<b>Данными интересовались:</b> {interest_overall_count} пользователь(ей)"
        
        await bot.send_document(chat_id, InputFile(file), caption=caption, parse_mode="HTML")

def format_list_results(results, original_query, for_file=False, user_id=None):
    sources = {}
    unknown_count = 0
    sources_list = []
    response_data_for_tracking = []
    matched_values_for_tracking = set()

    for data_item in results:
        if isinstance(data_item, dict):
            source = data_item.pop('database', None)
            if not source:
                unknown_count += 1
                source = f"Неизвестный источник #{unknown_count}"
            if source not in sources:
                sources[source] = []
                sources_list.append(source)
            sources[source].append(data_item)
            response_data_for_tracking.append(data_item)
            for value in data_item.values():
                matched_values_for_tracking.add(str(value))
        else:
            if "Другие данные" not in sources:
                sources["Другие данные"] = []
                sources_list.append("Другие данные")
            sources["Другие данные"].append(data_item)
            response_data_for_tracking.append(data_item)
            matched_values_for_tracking.add(str(data_item))

    if user_id:
        update_interaction_data(user_id, original_query, response_data_for_tracking, list(matched_values_for_tracking))

    data = load_data()
    current_response_interested_users = set()
    for value in matched_values_for_tracking:
        normalized_value = str(value).strip().lower()
        if normalized_value in data["unique_users_per_response"]:
            current_response_interested_users.update(data["unique_users_per_response"][normalized_value]["users"])
    interest_overall_count = len(current_response_interested_users)


    formatted = ""
    if for_file:
        formatted = f"Результаты по запросу: {original_query}\n"
        for source, items in sources.items():
            formatted += f"Источник: {source}\n"
            for item in items:
                if isinstance(item, dict):
                    for key, value in item.items():
                        formatted += f"{key}: {value}\n"
                        count_for_field = get_interest_count_for_value(value)
                        if count_for_field > 0:
                            formatted += f"(Проявили интерес: {count_for_field})\n"
                else:
                    formatted += f"{str(item)}\n"
                    interest_count = get_interest_count_for_value(item)
                    if interest_count > 0:
                        formatted += f"(Проявили интерес: {interest_count})\n"
                formatted += "\n"
        
        formatted += f"Данными интересовались: {interest_overall_count} пользователь(ей)\n"

        return formatted, sources_list, interest_overall_count
    else:
        formatted = """<b>🔍 Результаты по запросу: </b><code>{}</code>""".format(original_query)
        for source, items in sources.items():
            formatted += """\n<b>📁 Источник:</b> <code>{}</code>""".format(source)
            for item in items:
                formatted += "\n<blockquote>"
                if isinstance(item, dict):
                    for key, value in item.items():
                        formatted += "├─ <b>{}:</b> {}\n".format(key, value)
                        count_for_field = get_interest_count_for_value(value)
                        if count_for_field > 0:
                            formatted += "   └─ (Проявили интерес: {})\n".format(count_for_field)
                else:
                    formatted += "├─ {}\n".format(str(item))
                    interest_count = get_interest_count_for_value(item)
                    if interest_count > 0:
                        formatted += "   └─ (Проявили интерес: {})\n".format(interest_count)
                formatted += "</blockquote>"
        
        formatted += f"\n\n<b>Данными интересовались:</b> {interest_overall_count} пользователь(ей)"
        
        return formatted, sources_list, interest_overall_count

def format_probe_results(results, original_query, for_file=False, user_id=None):
    sources = {}
    sources_list = []
    response_data_for_tracking = []
    matched_values_for_tracking = set()

    for idx, data_item in results.items():
        source = data_item.pop('database', 'Неизвестный источник')
        if source not in sources:
            sources[source] = []
            sources_list.append(source)
        sources[source].append(data_item)
        response_data_for_tracking.append(data_item)
        for value in data_item.values():
            matched_values_for_tracking.add(str(value))

    if user_id:
        update_interaction_data(user_id, original_query, response_data_for_tracking, list(matched_values_for_tracking))

    data = load_data()
    current_response_interested_users = set()
    for value in matched_values_for_tracking:
        normalized_value = str(value).strip().lower()
        if normalized_value in data["unique_users_per_response"]:
            current_response_interested_users.update(data["unique_users_per_response"][normalized_value]["users"])
    interest_overall_count = len(current_response_interested_users)


    formatted = ""
    if for_file:
        formatted = f"Результаты по запросу: {original_query}\n"
        for source, items in sources.items():
            formatted += f"Источник: {source}\n"
            for data_item in items:
                for key, value in data_item.items():
                    formatted += f"{key}: {value}\n"
                    count_for_field = get_interest_count_for_value(value)
                    if count_for_field > 0:
                        formatted += f"(Проявили интерес: {count_for_field})\n"
                formatted += "\n"
        
        formatted += f"Данными интересовались: {interest_overall_count} пользователь(ей)\n"

        return formatted, sources_list, interest_overall_count
    else:
        formatted = """<b>🔍 Результаты по запросу: </b><code>{}</code>""".format(original_query)
        for source, items in sources.items():
            formatted += """\n<b>📁 Источник: </b><code>{}</code>""".format(source)
            for data_item in items:
                formatted += "\n<blockquote>"
                for key, value in data_item.items():
                    formatted += "├─ <b>{}:</b> {}\n".format(key, value)
                    count_for_field = get_interest_count_for_value(value)
                    if count_for_field > 0:
                        formatted += "   └─ (Проявили интерес: {})\n".format(count_for_field)
                formatted += "</blockquote>"
        
        formatted += f"\n\n<b>Данными интересовались:</b> {interest_overall_count} пользователь(ей)"
        
        return formatted, sources_list, interest_overall_count


        
@dp.message_handler(state='waiting_probe_query')
async def process_probe_query(message: types.Message, state: FSMContext):
    query = message.text.strip()
    user_data = await state.get_data()
    search_type = user_data.get('search_type', 'regular')
    
    search_msg = await message.reply("<b>🔍 Поиск</b>", parse_mode="HTML")
    task = asyncio.create_task(animate_search(search_msg.chat.id, search_msg.message_id))
    
    try:
        if search_type == 'extended':
            result = await make_api_request(f"extended_search/{query}")
        else:
            result = await make_api_request(f"search/{query}")
        
        task.cancel()
        
        if isinstance(result, dict) and 'error' in result:
            await bot.edit_message_text(
                chat_id=search_msg.chat.id,
                message_id=search_msg.message_id,
                text="<b>❌ Ошибка: {}</b>".format(result['error']),
                parse_mode="HTML"
            )
        elif isinstance(result, dict) and 'result' in result:
            if isinstance(result['result'], (dict, list)):
                if isinstance(result['result'], dict):
                    response_text, sources_list = format_probe_results(result['result'], query, for_file=False)
                    file_content, _ = format_probe_results(result['result'], query, for_file=True)
                else:
                    response_text, sources_list = format_list_results(result['result'], query, for_file=False)
                    file_content, _ = format_list_results(result['result'], query, for_file=True)
                
                await send_large_message(
                    chat_id=search_msg.chat.id,
                    text=response_text,
                    original_query=query,
                    sources_list=sources_list
                )
                await bot.delete_message(search_msg.chat.id, search_msg.message_id)
            else:
                await bot.edit_message_text(
                    chat_id=search_msg.chat.id,
                    message_id=search_msg.message_id,
                    text="<b>ℹ️ Неожиданный формат данных</b>",
                    parse_mode="HTML"
                )
        else:
            await bot.edit_message_text(
                chat_id=search_msg.chat.id,
                message_id=search_msg.message_id,
                text="<b>ℹ️ Ничего не найдено</b>",
                parse_mode="HTML"
            )
            
    except Exception as e:
        task.cancel()
        await bot.edit_message_text(
            chat_id=search_msg.chat.id,
            message_id=search_msg.message_id,
            text="<b>⚠️ Ошибка: {}</b>".format(str(e)),
            parse_mode="HTML"
        )
    
    await state.finish()

    
@dp.callback_query_handler(lambda c: c.data == 'extract_users', state='*')
async def extract_users_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()

    with open('База_Бота/users.txt', 'r', encoding='utf-8') as file:
        users_data = file.read()
    user_count = len(users_data.splitlines())
    document = types.InputFile('База_Бота/users.txt')
    await callback_query.message.answer_document(document)
    await callback_query.message.answer(f'📝В файле содержится {user_count} пользователей.')

@dp.callback_query_handler(lambda c: c.data == 'stats', state='*')
async def stats_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    with open('База_Бота/users.txt', 'r', encoding='utf-8') as file:
        lines = file.readlines()
        total_users = len(lines)
        active_users = sum(1 for line in lines if 'id' not in line)
    await callback_query.message.answer(f'📊Статистика:\n\n👤Всего пользователей: {total_users}\n✅Активных пользователей: {active_users}')



@dp.callback_query_handler(lambda c: c.data == 'send_message', state='*')
async def send_message_start(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer('Введите текст сообщения:')
    await SendMessage.text.set()

@dp.message_handler(state=SendMessage.text)
async def process_text(message: types.Message, state: FSMContext):
    async with state.proxy() as data:
        data['text'] = message.text
        data['entities'] = message.entities or message.caption_entities
        
    markup = InlineKeyboardMarkup(row_width=2)
    btn_yes = InlineKeyboardButton('Да', callback_data='yes')
    btn_no = InlineKeyboardButton('Нет', callback_data='no')
    markup.add(btn_yes, btn_no)
    await message.answer('Хотите добавить фото или видео?', reply_markup=markup)
    await SendMessage.media_type.set()

@dp.callback_query_handler(lambda c: c.data in ['yes', 'no'], state=SendMessage.media_type)
async def process_media_type(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    async with state.proxy() as data:
        if callback_query.data == 'yes':
            await callback_query.message.answer('Отправьте фото или видео:')
            await SendMessage.media.set()
        else:
            await send_message_to_users(data['text'], data.get('entities'), None, None)
            await state.finish()
            await callback_query.message.answer('✅Сообщение отправлено всем пользователям.')

@dp.message_handler(content_types=['photo', 'video'], state=SendMessage.media)
async def process_media(message: types.Message, state: FSMContext):
    async with state.proxy() as data:
        if message.photo:
            data['media_type'] = 'photo'
            data['media'] = message.photo[-1].file_id
        elif message.video:
            data['media_type'] = 'video'
            data['media'] = message.video.file_id
            
        await send_message_to_users(
            message.caption if message.caption else data['text'],
            message.caption_entities if message.caption else data.get('entities'),
            data['media_type'],
            data['media']
        )
        await state.finish()
        await message.answer('✅Сообщение отправлено всем пользователям.')

async def send_message_to_users(text, entities, media_type, media_id):
    with open('База_Бота/users.txt', 'r', encoding='utf-8') as file:
        for line in file:
            user_id = line.split()[0]
            try:
                if media_type == 'photo':
                    await bot.send_photo(
                        user_id, 
                        media_id, 
                        caption=text,
                        caption_entities=entities,
                        parse_mode=None  
                    )
                elif media_type == 'video':
                    await bot.send_video(
                        user_id, 
                        media_id, 
                        caption=text,
                        caption_entities=entities,
                        parse_mode=None  
                    )
                else:
                    await bot.send_message(
                        user_id, 
                        text,
                        entities=entities,
                        parse_mode=None  
                    )
            except Exception as e:
                logging.error(f'Error sending message to user {user_id}: {e}')
    
@dp.callback_query_handler(lambda c: c.data == 'demolition', state='*')
async def demolition_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    snos_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📫Email-snos-меню📫</b> - Меню для сноса через почту
🚨<b>Botnen-snos🚨</b> - снос через репорты
<b>📢Канал-snos📢</b> - снос канала через репорты</blockquote>
<b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    markup = InlineKeyboardMarkup(row_width=2)
    btn_email_complaint = InlineKeyboardButton('📫Email-snos-меню📫', callback_data='Email_snos_men')  
    btn_report_message = InlineKeyboardButton('🚨Botnen-snos🚨', callback_data='report_message')
    btn_channel_demolition = InlineKeyboardButton('📢Канал-snos📢', callback_data='channel_demolition')  
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='to_start')  
    
    markup.add(btn_email_complaint)
    markup.add(btn_report_message, btn_channel_demolition) 
    markup.add(btn_back)
    
    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=snos_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=snos_message,
            reply_markup=markup,
            parse_mode="HTML"
        )

async def get_subscription_end_time(user_id: int):
    try:
        with open("База_Бота/paid_users.txt", "r", encoding="utf-8") as file:
            for line in file:
                line = line.strip()
                if not line:
                    continue
                parts = line.split(',')
                if len(parts) == 2 and int(parts[0]) == user_id:
                    if parts[1] == "forever":
                        return "forever"
                    return datetime.strptime(parts[1], "%Y-%m-%d %H:%M:%S")
    except FileNotFoundError:
        print("Файл paid_users.txt не найден.")
    except Exception as e:
        print(f"Ошибка при чтении файла: {e}")
    return None


@dp.callback_query_handler(lambda c: c.data == 'spam_menu', state='*')
async def spam_menu_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()

    try:
        with open('База_Бота/emails.json', 'r') as f:
            emails_data = json.load(f)                
        email_count = len(emails_data.keys())  
        client_count = len(clients)  

        spam_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📧 Управление спамом</b>
📊 <b>Количество почт:</b> {email_count}
📊 <b>Количество клиентов:</b> {client_count}
🔥 <b>Spam-code</b> - Отправляет код входа
📧 <b>Email-spam</b> - Отправляет спам на почту</blockquote>
<b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    except FileNotFoundError:
        spam_message = "❌ Файл `emails.json` не найден!"
        email_count = 0
    except json.JSONDecodeError:
        spam_message = "❌ Ошибка в формате `emails.json`!"
        email_count = 0

    markup = InlineKeyboardMarkup(row_width=2)
    btn_spam_code = InlineKeyboardButton('🔥Spam-code🔥', callback_data='spam_code')
    btn_email_spam = InlineKeyboardButton('📧 Email-spam📧', callback_data='email_spam')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='to_start')
    markup.add(btn_spam_code, btn_email_spam)
    markup.add(btn_back)

    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=spam_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=spam_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    
class EmailSpamStates(StatesGroup):
    waiting_for_receiver = State()
    waiting_for_subject = State()
    waiting_for_body = State()
    waiting_for_count = State()

def send_spam_email(receiver, sender_email, sender_password, subject, body):
    domain = sender_email.split('@')[1]
    if domain not in smtp_servers:
        return False
    smtp_server, smtp_port = smtp_servers[domain]
    try:
        msg = MIMEMultipart()
        msg['From'] = sender_email
        msg['To'] = receiver
        msg['Subject'] = subject
        msg.attach(MIMEText(body, 'plain'))
        with smtplib.SMTP(smtp_server, smtp_port) as server:
            server.starttls()
            server.login(sender_email, sender_password)
            server.sendmail(sender_email, receiver, msg.as_string())
            time.sleep(3)
        return True
    except Exception as e:
        return False

def load_banned_users():
    banned_users_file = os.path.join('База_Бота', 'banned_users.txt')
    try:
        with open(banned_users_file, 'r') as file:
            return set(map(int, file.read().splitlines()))
    except FileNotFoundError:
        return set()

def save_banned_users(banned_users):
    banned_users_file = os.path.join('База_Бота', 'banned_users.txt')
    with open(banned_users_file, 'w') as file:
        for user_id in banned_users:
            file.write(f'{user_id}\n')

@dp.message_handler(state=EmailSpamStates.waiting_for_receiver)
async def process_receiver_email(message: types.Message, state: FSMContext):
    await state.update_data(receiver=message.text)
    await EmailSpamStates.next()
    await message.answer("📝 Введите тему письма:")

@dp.message_handler(state=EmailSpamStates.waiting_for_subject)
async def process_subject(message: types.Message, state: FSMContext):
    await state.update_data(subject=message.text)
    await EmailSpamStates.next()
    await message.answer("📝 Введите текст письма:")

@dp.message_handler(state=EmailSpamStates.waiting_for_body)
async def process_body(message: types.Message, state: FSMContext):
    await state.update_data(body=message.text)
    await EmailSpamStates.next()
    await message.answer("🔢 Введите количество отправок:")


@dp.message_handler(state=EmailSpamStates.waiting_for_count)
async def process_count(message: types.Message, state: FSMContext):
    try:
        count = int(message.text)
        if count <= 0:
            await message.answer("❌ <b>Количество отправок должно быть больше 0.</b>", parse_mode="HTML")
            return
        try:
            with open('База_Бота/emails.json', 'r') as f:
                senders = json.load(f)
            if not senders:
                await message.answer("❌ <b>Файл emails.json пуст или не содержит почт.</b>", parse_mode="HTML")
                return
        except FileNotFoundError:
            await message.answer("❌ <b>Файл emails.json не найден.</b>", parse_mode="HTML")
            return
        except json.JSONDecodeError:
            await message.answer("❌ <b>Ошибка в формате файла emails.json.</b>", parse_mode="HTML")
            return

        data = await state.get_data()
        receiver = data.get('receiver')
        subject = data.get('subject')
        body = data.get('body')

        status_message = await message.answer("⏳ <b>Подготовка к отправке...</b>", parse_mode="HTML")

        successful = 0
        failed = 0

        for i in range(count):
            sender_email, sender_password = random.choice(list(senders.items()))
            status = send_spam_email(receiver, sender_email, sender_password, subject, body)

            if status:
                successful += 1
            else:
                failed += 1

            await status_message.edit_text(
                f"<b>───── ⋆⋅☆⋅⋆ ─────</b>\n"
                f"<blockquote>"
                f"<b>📤 Отправитель:</b> <code>{sender_email}</code>\n"
                f"<b>📥 Цель:</b> <code>{receiver}</code>\n"
                f"<b>📝 Тема:</b> <code>{subject}</code>\n"
                f"<b>📄 Текст:</b> <code>{html.escape(body[:50])}...</code>\n"
                f"<b>👀 Статус:</b> {'✅<b>Удачно</b>' if status else '❌<b>Не удачно</b>'}\n"
                f"<b>📩 Прогресс:</b> <i>{i + 1}/{count}</i>\n"
                f"</blockquote>"
                f"<b>───── ⋆⋅☆⋅⋆ ─────</b>",
                parse_mode="HTML"
            )

        await status_message.edit_text(
            f"<b>───── ⋆⋅☆⋅⋆ ─────</b>\n"
            f"<blockquote><b>📬 Итоговый отчет</b>\n"
            f"<b>🎯 Цель:</b> <code>{receiver}</code>\n"
            f"<b>✅ Удачно:</b> <code>{successful}</code>\n"
            f"<b>❌ Не удачно:</b> <code>{failed}</code>\n"
            f"<b>📝 Тема:</b> <code>{subject}</code>\n"
                f"<b>📄 Текст:</b> <code>{html.escape(body[:50])}...</code>\n</blockquote>"
                f"<b>───── ⋆⋅☆⋅⋆ ─────</b>",
                parse_mode="HTML"
            )

        await state.finish()
    except ValueError:
        await message.answer("❌ <b>Пожалуйста, введите корректное число.</b>", parse_mode="HTML")



class SpamCodeStates(StatesGroup):
    waiting_for_numbers = State()


@dp.message_handler(state=SpamCodeStates.waiting_for_numbers)
async def process_spam_code_input(message: types.Message, state: FSMContext):
    try:
        lines = message.text.splitlines()
        phone_numbers = []
        for line in lines:
            line = line.strip()
            if line: 
                phone_number, num_sendings = line.split()
                phone_numbers.append((phone_number, int(num_sendings)))
        if phone_numbers:
            await process_numbers(message, phone_numbers)
            await state.finish()  
        else:
            await message.reply("Список номеров пуст.")
    except ValueError:
        await message.reply('❌ Неверный формат ввода. Используйте формат: +79991234567 10')


async def process_numbers(message, phone_numbers):
    message = await bot.send_message(message.chat.id, "⏳ Начинаем отправку кодов...")
    message_id = message.message_id
    overall_summary = "📊 Итоги отправки кодов\n"

    for phone_number, num_sendings in phone_numbers:
        summary = await send_code_requests(phone_number, num_sendings, message.chat.id, message_id, bot)
        overall_summary += summary

    await bot.edit_message_text(overall_summary, message.chat.id, message_id)

async def send_code_requests(phone_number, num_sendings, chat_id, message_id, bot):
    successful_sendings = 0
    failed_sendings = 0
    start_time = asyncio.get_event_loop().time()

    if not re.match(r'^\+?[1-9]\d{10,12}$', phone_number):
        await bot.edit_message_text(
            f"""<b>📱 Номер {html.escape(phone_number)}:</b>
Ошибка: Неправильный формат""",
            chat_id,
            message_id,
            parse_mode="HTML"
        )
        return

    for i in range(num_sendings):
        client_data = random.choice(clients)
        client = None
        try:
            client = TelegramClient(client_data["name"], client_data["api_id"], client_data["api_hash"])
            await client.connect()
            await client.send_code_request(phone_number)
            successful_sendings += 1
            status = "✅ <b>Удачно</b>"
        except ValueError as e:
            failed_sendings += 1
            status = f"❌ <b>Ошибка:</b> Клиент '{html.escape(client_data.get('name', 'неизвестный'))}' зарегистрирован неправильно"
        except Exception as e:
            failed_sendings += 1
            status = f"❌ <b>Не удачно:</b> {html.escape(str(e))}"
        finally:
            if client:
                await client.disconnect()
                client.session.delete()

        await bot.edit_message_text(
            f"""<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📱 Номер:</b> <code>{html.escape(phone_number)}</code>
<b>👤 Клиент:</b> <code>{html.escape(client_data.get('name', 'неизвестный'))}</code>
<b>📤 Статус:</b> {status}
<b>📊 Прогресс:</b> <i>{successful_sendings + failed_sendings}/{num_sendings}</i>
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>""",
            chat_id,
            message_id,
            parse_mode="HTML"
        )
        await asyncio.sleep(1)

    end_time = asyncio.get_event_loop().time()
    elapsed_time = end_time - start_time
    total_time_str = "{:.2f}".format(elapsed_time)

    await bot.edit_message_text(
        f"""<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📱 Номер:</b> <code>{html.escape(phone_number)}</code>
<b>✅ Удачно:</b> <code>{successful_sendings}</code>
<b>❌ Неудачно:</b> <code>{failed_sendings}</code>
<b>⏱️ Время:</b> <code>{total_time_str} сек.</code></blockquote>
<b>───── ⋆⋅☆⋅⋆ ─────</b>""",
        chat_id,
        message_id,
        parse_mode="HTML"
    )

from aiogram import types
from aiogram.dispatcher import FSMContext
import json

import shutil
from telethon.tl.types import User
from telethon.errors import AuthKeyDuplicatedError, SessionPasswordNeededError

@dp.callback_query_handler(lambda c: c.data == 'email_menu', state='*')
async def email_menu_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    email_message = """
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📧 Управление почтами Gmail</b>

📥 <b>Добавить почту</b> - Добавить новую почту в систему.
🗑 <b>Удалить почту</b> - Удалить почту из системы.
📋 <b>Список почт</b> - Просмотреть все добавленные почты.
🔙 <b>Назад</b> - Вернуться в админ панель.
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    
    markup = InlineKeyboardMarkup(row_width=2)
    btn_add_email = InlineKeyboardButton('📥 Добавить почту', callback_data='add_email')
    btn_delete_email = InlineKeyboardButton('🗑 Удалить почту', callback_data='delete_email')
    btn_list_emails = InlineKeyboardButton('📋 Список почт', callback_data='list_emails')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='admin_panel')
    
    markup.add(btn_add_email, btn_delete_email)
    markup.add(btn_list_emails)
    markup.add(btn_back)

    await bot.edit_message_caption(
        chat_id=callback_query.message.chat.id,
        message_id=callback_query.message.message_id,
        caption=email_message,
        reply_markup=markup,
        parse_mode="HTML"
    )

@dp.callback_query_handler(lambda c: c.data == 'admin_panel', state='*')
async def admin_panel_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    admin_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>👋 Приветствую тебя админ!</b>
👤 <b>.session(осн)</b> - Меню Основные session
👤 <b>.sesion(Не_осн)</b> - Меню Не_основные session
👑 <b>Админы</b> - Просмотр списка админов.
🎫 <b>Промокоды</b> - Управление промокодами.
📧 <b>Gmail</b> - Управление почтами
⚙️ <b>API</b> - Управление Api для пробива
🔄 <b>Обновить код</b> - Перезапустить бота с обновлениями.
🔙 <b>Назад</b> - Вернуться в главное меню.
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""

    markup = InlineKeyboardMarkup(row_width=2)        
    btn_user_menu = InlineKeyboardButton('User Menu', callback_data='user_menu')
    btn_create_account = InlineKeyboardButton('👤.session(осн)', callback_data='sessionOsn_menu')
    btn_create_non_main = InlineKeyboardButton('👤.sesion(Не_осн)', callback_data='create_non_menu')
    btn_view_admins = InlineKeyboardButton('👑Админы', callback_data='view_admins')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='to_start')
    btn_promocodes = InlineKeyboardButton('🎫 Промокоды', callback_data='promocodes_menu')
    btn_update = InlineKeyboardButton('🔄 Обновить код', callback_data='update_code')
    btn_email = InlineKeyboardButton('📧 Gmail', callback_data='email_menu')
        
    markup.add(btn_create_account, btn_create_non_main, btn_view_admins, btn_user_menu)
    markup.add(btn_promocodes, 
              InlineKeyboardButton('⚙️API', callback_data='api_settings'),
              btn_update, btn_email)
    markup.add(btn_back)

    await bot.edit_message_caption(
        chat_id=callback_query.message.chat.id,
        message_id=callback_query.message.message_id,
        caption=admin_message,
        reply_markup=markup,
        parse_mode="HTML"
    )

@dp.callback_query_handler(lambda c: c.data == 'add_email', state='*')
async def add_email_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await bot.send_message(
        callback_query.from_user.id,
        "Введите одну или несколько почт с паролями (каждая с новой строки) в формате:\n\n"
        "<code>email@gmail.com:password</code>\n"
        "<code>email2@yahoo.com:password123</code>\n"
        "<code>email3@mail.ru:qwerty</code>",
        parse_mode="HTML"
    )
    await state.set_state("waiting_for_email")

@dp.message_handler(state="waiting_for_email")
async def process_email(message: types.Message, state: FSMContext):
    emails_file = os.path.join('База_Бота', 'emails.json')
    existing_emails = {}
    if os.path.exists(emails_file):
        with open(emails_file, "r", encoding="utf-8") as f:
            try:
                existing_emails = json.load(f)
            except json.JSONDecodeError:
                existing_emails = {}
    input_lines = message.text.strip().split('\n')
    added_emails = []
    duplicate_emails = []
    invalid_format = []
    for line in input_lines:
        line = line.strip()
        if not line:
            continue
        try:
            if ":" not in line:
                raise ValueError
            email, password = line.split(":", 1)
            email = email.strip()
            password = password.strip()
            if "@" not in email or "." not in email:
                invalid_format.append(line)
                continue
            if email in existing_emails:
                duplicate_emails.append(email)
                continue
            existing_emails[email] = password
            added_emails.append(email)
        except ValueError:
            invalid_format.append(line)
    with open(emails_file, "w", encoding="utf-8") as f:
        json.dump(existing_emails, f, ensure_ascii=False, indent=4)
    report = []
    if added_emails:
        report.append("✅ <b>Успешно добавлены:</b>")
        report.extend(f"• <code>{email}</code>" for email in added_emails)
    if duplicate_emails:
        report.append("\n❌ <b>Уже существуют:</b>")
        report.extend(f"• <code>{email}</code>" for email in duplicate_emails)
    if invalid_format:
        report.append("\n⚠️ <b>Неверный формат:</b>")
        report.extend(f"• <code>{line}</code>" for line in invalid_format)
    if not report:
        await message.answer("❌ Не получено ни одной validной почты для добавления.")
    else:
        await message.answer("\n".join(report), parse_mode="HTML")
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'list_emails', state='*')
async def list_emails_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    
    emails_file = "База_Бота/emails.json"
    if not os.path.exists(emails_file):
        await bot.send_message(callback_query.from_user.id, "📭 Список почт пуст.")
        return
    
    with open(emails_file, "r", encoding="utf-8") as f:
        emails = json.load(f)
    
    if not emails:
        await bot.send_message(callback_query.from_user.id, "📭 Список почт пуст.")
        return
    
    message_text = "📧 <b>Список добавленных почт:</b>\n\n"
    for i, (email, password) in enumerate(emails.items(), 1):
        message_text += f"{i}. <code>{email}</code>\n"
    
    await bot.send_message(callback_query.from_user.id, message_text, parse_mode="HTML")

@dp.callback_query_handler(lambda c: c.data == 'delete_email', state='*')
async def delete_email_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    
    emails_file = "База_Бота/emails.json"
    if not os.path.exists(emails_file):
        await bot.send_message(callback_query.from_user.id, "📭 Список почт пуст. Нечего удалять.")
        return
    
    with open(emails_file, "r", encoding="utf-8") as f:
        emails = json.load(f)
    
    if not emails:
        await bot.send_message(callback_query.from_user.id, "📭 Список почт пуст. Нечего удалять.")
        return
    
    markup = InlineKeyboardMarkup(row_width=2)
    for email in emails.keys():
        markup.add(InlineKeyboardButton(email, callback_data=f"delete_email_{email}"))
    
    await bot.send_message(
        callback_query.from_user.id,
        "Выберите почту для удаления:",
        reply_markup=markup
    )


@dp.callback_query_handler(lambda c: c.data.startswith('exec_delete_'))
async def process_email_deletion(callback_query: types.CallbackQuery):
    try:
        await callback_query.answer()
        email = callback_query.data.replace("exec_delete_", "")
        
        emails_file = "База_Бота/emails.json"
        
        with open(emails_file, "r", encoding="utf-8") as f:
            emails = json.load(f)
        
        if email in emails:
            del emails[email]
            
            with open(emails_file, "w", encoding="utf-8") as f:
                json.dump(emails, f, ensure_ascii=False, indent=4)
            
            await callback_query.message.delete()
            await bot.send_message(
                callback_query.from_user.id,
                f"✅ Почта <code>{email}</code> успешно удалена!",
                parse_mode="HTML"
            )
        else:
            await callback_query.message.edit_text(
                f"❌ Почта <code>{email}</code> не найдена.",
                parse_mode="HTML"
            )
            
    except Exception as e:
        logger.error(f"Ошибка при удалении почты: {e}")
        await callback_query.message.edit_text(
            f"⚠️ Произошла ошибка при удалении почты: {str(e)}",
            parse_mode="HTML"
        )

@dp.callback_query_handler(lambda c: c.data.startswith('delete_email_'))
async def ask_confirm_delete(callback_query: types.CallbackQuery):
    email = callback_query.data.replace("delete_email_", "")
    
    markup = InlineKeyboardMarkup()
    markup.add(
        InlineKeyboardButton("✅ Да, удалить", callback_data=f"exec_delete_{email}"),
        InlineKeyboardButton("❌ Отмена", callback_data="cancel_delete")
    )    
    message = await callback_query.message.answer(
        f"Удалить почту <code>{email}</code>?",
        reply_markup=markup,
        parse_mode="HTML"
    )
    await callback_query.message.delete()

@dp.callback_query_handler(lambda c: c.data == "cancel_delete")
async def cancel_delete(callback_query: types.CallbackQuery):
    await callback_query.message.delete()

@dp.callback_query_handler(lambda c: c.data.startswith('delete_email_'))
async def ask_confirm_delete(callback_query: types.CallbackQuery):
    email = callback_query.data.replace("delete_email_", "")
    
    markup = InlineKeyboardMarkup()
    markup.add(
        InlineKeyboardButton("✅ Да, удалить", callback_data=f"exec_delete_{email}"),
        InlineKeyboardButton("❌ Отмена", callback_data="cancel_delete")
    )    
    message = await callback_query.message.answer(
        f"Удалить почту <code>{email}</code>?",
        reply_markup=markup,
        parse_mode="HTML"
    )
    await callback_query.message.delete()

@dp.callback_query_handler(lambda c: c.data == "cancel_delete")
async def cancel_delete(callback_query: types.CallbackQuery):
    await callback_query.message.delete()
    
    

@dp.callback_query_handler(lambda c: c.data == 'user_menu', state='*')
async def demolition_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()

    if os.path.exists("База_Бота/users.txt"):
        with open("База_Бота/users.txt", "r") as file:
            user_ids = [line.strip() for line in file.readlines() if line.strip()]
            user_count = len(user_ids)
    else:
        user_count = 0
    user_menu_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📊 Статистика:</b> Пользователей использует бот: <code>{user_count}</code>
🚫 <b>Бан</b> - Заблокировать пользователя.
👥 <b>Статистика</b> - Просмотр статистики бота.
👑 <b>Vip</b> - Защита от сноса.
🔍 <b>Чек пользователя</b> - Проверить информацию о пользователе.
⏳ <b>Подписка</b> - Управление подписками пользователей.
 </blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    markup = InlineKeyboardMarkup(row_width=2)
    btn_statsit = InlineKeyboardButton('👥 Статистика', callback_data='statsit')
    btn_banis = InlineKeyboardButton('🚫 Бан', callback_data='banis_user')        
    btn_privat = InlineKeyboardButton('👑 Vip', callback_data='privat')
    btn_check_user = InlineKeyboardButton('🔍 Чек пользователя', callback_data='check_user')
    btn_user = InlineKeyboardButton('⏳ Подписка', callback_data='user')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='admin_panel')
    markup.add(btn_statsit, btn_banis, btn_privat, btn_check_user, btn_user) 
    markup.add(btn_back)
    
    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=user_menu_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=user_menu_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
def write_private_users(data):
    with open("База_Бота/private_users.txt", "w", encoding="utf-8") as f:
        ids_str = ",".join(map(str, data["ids"])) if data["ids"] else ""
        usernames_str = ",".join(f"@{uname}" for uname in data["usernames"]) if data["usernames"] else ""
        f.write(f"{ids_str}\n{usernames_str}\n")
private_users = read_private_users()

async def is_private_user(text, private_users):
    words = text.split()
    for word in words:
        if word.isdigit() and int(word) in private_users["ids"]:
            return True
        if word.startswith('@') and word[1:] in private_users["usernames"]:
            return True
    return False
    
class CheckUser(StatesGroup):
    waiting_for_user_id = State()

@dp.callback_query_handler(lambda c: c.data == 'check_user')
async def check_user_callback(callback_query: types.CallbackQuery):
    await callback_query.answer()
    await bot.send_message(
        chat_id=callback_query.message.chat.id,
        text="Введите ID пользователя для проверки:")
    await CheckUser.waiting_for_user_id.set()

@dp.message_handler(state=CheckUser.waiting_for_user_id)
async def process_user_id(message: types.Message, state: FSMContext):
    try:
        user_id = int(message.text)
        try:
            user = await bot.get_chat(user_id)
            subscription_end = await get_subscription_end_time(user_id)
            remaining_time = await get_remaining_time(user_id)
            
            if subscription_end == "forever":
                subscription_status = "✅ <b>Активна (Навсегда)</b>"
                subscription_end_formatted = "∞"
            elif subscription_end and subscription_end > datetime.now():
                subscription_status = "✅ <b>Активна</b>"
                subscription_end_formatted = subscription_end.strftime("%Y-%m-%d %H:%M:%S")
            else:
                subscription_status = "❌ <b>Не активна</b>"
                subscription_end_formatted = "<i>Нет активной подписки</i>"
            
            private_users = read_private_users()
            is_vip = user.id in private_users["ids"] or (user.username and user.username in private_users["usernames"])
            vip_status = "✅ <b>Да</b>" if is_vip else "❌ <b>Нет</b>"
            
            if os.path.exists("База_Бота/users.txt"):
                with open("База_Бота/users.txt", "r") as file:
                    user_ids = [line.strip() for line in file.readlines() if line.strip()]
                    is_in_users = "✅ <b>Да</b>" if str(user.id) in user_ids else "❌ <b>Нет</b>"
            else:
                is_in_users = "❌ <b>Нет (файл не существует)</b>"
            
            is_banned = user_id in banned_users
            ban_button_text = "🔓 Разбанить" if is_banned else "🚫 Забанить"
            ban_button_data = f"unban_user:{user_id}" if is_banned else f"ban_user:{user_id}"
            
            user_info = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🔍 Информация о пользователе</b>

<b>🆔 ID:</b> <code>{user.id}</code>
<b>👤 Имя:</b> <code>{html.escape(user.first_name)}</code>
<b>👤 Фамилия:</b> <code>{html.escape(user.last_name) if user.last_name else 'не указана'}</code>
<b>👤 Юзернейм:</b> @{html.escape(user.username) if user.username else 'отсутствует'}

<b>🔐 Подписка:</b> {subscription_status}
<b>💰 Подписка до:</b> <code>{subscription_end_formatted}</code>
<b>⏳ Оставшееся время:</b> <code>{remaining_time}</code>
<b>🌟 VIP статус:</b> {vip_status}
<b>📝 В базе бота:</b> {is_in_users}
<b>🚫 Статус бана:</b> {"✅ <b>Да</b>" if is_banned else "❌ <b>Нет</b>"}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
            markup = InlineKeyboardMarkup(row_width=2)
            btn_ban = InlineKeyboardButton(ban_button_text, callback_data=ban_button_data)
            btn_privat = InlineKeyboardButton('👑 Vip', callback_data=f'toggle_vip:{user_id}')
            btn_user = InlineKeyboardButton('⏳ Подписка', callback_data=f'user:{user_id}')
            
            
            markup.add(btn_ban, btn_privat, btn_user)
            
            await bot.send_message(
                chat_id=message.chat.id,
                text=user_info,
                reply_markup=markup,
                parse_mode="HTML"
            )
            await state.finish()
            
        except Exception as e:
            await message.reply(f"Ошибка: пользователь с ID {user_id} не найден. Ошибка: {str(e)}")
            await state.finish()
    except ValueError:
        await message.reply("Пожалуйста, введите корректный ID пользователя (только цифры).")
        return

@dp.callback_query_handler(lambda c: c.data.startswith('user:'))
async def handle_user_subscription(callback_query: types.CallbackQuery):
    user_id = int(callback_query.data.split(':')[1])
    
    markup = InlineKeyboardMarkup(row_width=1)
    btn_add = InlineKeyboardButton('➕ Добавить', callback_data=f'add_sub:{user_id}')
    btn_remove = InlineKeyboardButton('➖ Удалить', callback_data=f'remove_sub:{user_id}')
    btn_change = InlineKeyboardButton('✏️ Изменить', callback_data=f'change_sub:{user_id}')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data=f'sub_back:{user_id}')
    
    markup.add(btn_add, btn_remove, btn_change, btn_back)
    
    await bot.edit_message_reply_markup(
        chat_id=callback_query.message.chat.id,
        message_id=callback_query.message.message_id,
        reply_markup=markup
    )
    await callback_query.answer()

@dp.callback_query_handler(lambda c: c.data.startswith('sub_back:'))
async def handle_subscription_back(callback_query: types.CallbackQuery):
    user_id = int(callback_query.data.split(':')[1])
    
    user = await bot.get_chat(user_id)
    subscription_end = await get_subscription_end_time(user_id)
    remaining_time = await get_remaining_time(user_id)
    
    if subscription_end == "forever":
        subscription_status = "✅ <b>Активна (Навсегда)</b>"
        subscription_end_formatted = "∞"
    elif subscription_end and subscription_end > datetime.now():
        subscription_status = "✅ <b>Активна</b>"
        subscription_end_formatted = subscription_end.strftime("%Y-%m-%d %H:%M:%S")
    else:
        subscription_status = "❌ <b>Не активна</b>"
        subscription_end_formatted = "<i>Нет активной подписки</i>"
    
    is_banned = user_id in banned_users
    ban_button_text = "🔓 Разбанить" if is_banned else "🚫 Забанить"
    ban_button_data = f"unban_user:{user_id}" if is_banned else f"ban_user:{user_id}"
    
    markup = InlineKeyboardMarkup(row_width=2)
    btn_ban = InlineKeyboardButton(ban_button_text, callback_data=ban_button_data)
    btn_privat = InlineKeyboardButton('👑 Vip', callback_data=f'toggle_vip:{user_id}')
    btn_user = InlineKeyboardButton('⏳ Подписка', callback_data=f'user:{user_id}')    
    
    markup.add(btn_ban, btn_privat, btn_user)
    
    await bot.edit_message_reply_markup(
        chat_id=callback_query.message.chat.id,
        message_id=callback_query.message.message_id,
        reply_markup=markup
    )
    await callback_query.answer()

@dp.callback_query_handler(lambda c: c.data.startswith('add_sub:'))
async def handle_add_subscription(callback_query: types.CallbackQuery, state: FSMContext):
    user_id = int(callback_query.data.split(':')[1])
    exists, status = check_user_subscription_exist(user_id)
    
    if exists:
        await callback_query.answer(f"ℹ️ Пользователь {user_id} уже существует! Статус: {status}.", show_alert=True)
        return
    
    await SubStates.add_sub_date.set()
    async with state.proxy() as data:
        data['user_id'] = user_id
    
    await bot.send_message(
        chat_id=callback_query.message.chat.id,
        text="📅 Введите срок подписки ('forever' или число дней.часов):"
    )
    await callback_query.answer()

@dp.callback_query_handler(lambda c: c.data.startswith('remove_sub:'))
async def handle_remove_subscription(callback_query: types.CallbackQuery, state: FSMContext):
    user_id = int(callback_query.data.split(':')[1])
    exists, _ = check_user_subscription_exist(user_id)
    
    if not exists:
        await callback_query.answer(f"❌ Пользователь с ID {user_id} не найден.", show_alert=True)
        return
    
    await SubStates.remove_sub.set()
    async with state.proxy() as data:
        data['user_id'] = user_id
    
    try:
        with open("База_Бота/paid_users.txt", "r") as file:
            lines = file.readlines()
        
        with open("База_Бота/paid_users.txt", "w") as file:
            deleted = False
            for line in lines:
                if not line.startswith(f"{user_id},"):
                    file.write(line)
                else:
                    deleted = True
            
            if deleted:
                await bot.send_message(
                    chat_id=callback_query.message.chat.id,
                    text=f"✅ Пользователь {user_id} успешно удалён!"
                )
            else:
                await bot.send_message(
                    chat_id=callback_query.message.chat.id,
                    text=f"❌ Пользователь с ID {user_id} не найден."
                )
    except Exception as e:
        await bot.send_message(
            chat_id=callback_query.message.chat.id,
            text=f"❌ Ошибка при удалении: {e}"
        )
    
    await state.finish()
    await callback_query.answer()

@dp.callback_query_handler(lambda c: c.data.startswith('change_sub:'))
async def handle_change_subscription(callback_query: types.CallbackQuery, state: FSMContext):
    user_id = int(callback_query.data.split(':')[1])
    exists, status = check_user_subscription_exist(user_id)
    
    if not exists:
        await callback_query.answer(f"❌ Пользователь с ID {user_id} не найден.", show_alert=True)
        return
    
    await SubStates.change_sub_date.set()
    async with state.proxy() as data:
        data['user_id'] = user_id
        paid_users = read_paid_users()
        date_part = paid_users.get(str(user_id), "")
        
        if date_part.lower() == 'forever':
            await bot.send_message(
                chat_id=callback_query.message.chat.id,
                text=f"🕒 Текущий статус: вечная подписка"
            )
        else:
            await bot.send_message(
                chat_id=callback_query.message.chat.id,
                text=f"🕒 Текущая дата окончания: {date_part}"
            )
    
    await bot.send_message(
        chat_id=callback_query.message.chat.id,
        text="📅 Введите новое время ('forever' или число дней.часов):"
    )
    await callback_query.answer()

@dp.message_handler(state=SubStates.add_sub_date)
async def process_add_sub_date(message: types.Message, state: FSMContext):
    time_input = message.text.strip()
    formatted_date = parse_sub_time_input(time_input)
    
    if formatted_date is None:
        await message.reply("❌ Неверный формат! Введите 'forever' или число (напр. 1, 0.5, 1.12)")
        return
    
    async with state.proxy() as data:
        user_id = data['user_id']
        
        try:
            with open("База_Бота/paid_users.txt", "a") as file:
                file.write(f"{user_id},{formatted_date}\n")
            
            if formatted_date == 'forever':
                await message.reply(f"✅ Новый пользователь {user_id} добавлен навсегда!")
            else:
                await message.reply(f"✅ Новый пользователь {user_id} добавлен до {formatted_date}!")
        except Exception as e:
            await message.reply(f"❌ Ошибка: {e}")
    
    await state.finish()

@dp.message_handler(state=SubStates.change_sub_date)
async def process_change_sub_date(message: types.Message, state: FSMContext):
    time_input = message.text.strip()
    formatted_date = parse_sub_time_input(time_input)
    
    if formatted_date is None:
        await message.reply("❌ Неверный формат! Введите 'forever' или число (напр. 1, 0.5, 1.12)")
        return
    
    async with state.proxy() as data:
        user_id = data['user_id']
        
        try:
            with open("База_Бота/paid_users.txt", "r") as file:
                lines = file.readlines()
            
            with open("База_Бота/paid_users.txt", "w") as file:
                updated = False
                for line in lines:
                    if line.startswith(f"{user_id},"):
                        file.write(f"{user_id},{formatted_date}\n")
                        updated = True
                    else:
                        file.write(line)
                
                if updated:
                    if formatted_date == 'forever':
                        await message.reply(f"✅ Пользователь {user_id} изменён на вечную подписку!")
                    else:
                        await message.reply(f"✅ Время для пользователя {user_id} изменено до {formatted_date}!")
                else:
                    await message.reply(f"❌ Пользователь с ID {user_id} не найден.")
        except Exception as e:
            await message.reply(f"❌ Ошибка: {e}")
    
    await state.finish()

def check_user_subscription_exist(user_id):
    try:
        with open("База_Бота/paid_users.txt", "r") as file:
            for line in file:
                if line.startswith(f"{user_id},"):
                    date_part = line.split(',')[1].strip()
                    if date_part.lower() == 'forever':
                        return True, "вечная подписка"
                    else:
                        return True, f"до {date_part}"
        return False, "не найден"
    except FileNotFoundError:
        return False, "файл не существует"

def parse_sub_time_input(time_input):
    if time_input.lower() == 'forever':
        return 'forever'
    
    try:
        if '.' in time_input:
            days, hours = time_input.split('.')
            total_hours = float(days) * 24 + float(hours)
        else:
            total_hours = float(time_input) * 24
        
        end_date = datetime.now() + timedelta(hours=total_hours)
        return end_date.strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None

@dp.callback_query_handler(lambda c: c.data.startswith('toggle_vip:'))
async def toggle_vip_status(callback_query: types.CallbackQuery):
    user_id = int(callback_query.data.split(':')[1])
    private_users = read_private_users()
    is_vip = user_id in private_users["ids"]
    
    if is_vip:
        private_users["ids"].remove(user_id)
        action = "удален"
    else:
        private_users["ids"].append(user_id)
        action = "добавлен"
    
    write_private_users(private_users)
    
    try:
        user = await bot.get_chat(user_id)
        subscription_end = await get_subscription_end_time(user_id)
        remaining_time = await get_remaining_time(user_id)
        
        if subscription_end == "forever":
            subscription_status = "✅ <b>Активна (Навсегда)</b>"
            subscription_end_formatted = "∞"
        elif subscription_end and subscription_end > datetime.now():
            subscription_status = "✅ <b>Активна</b>"
            subscription_end_formatted = subscription_end.strftime("%Y-%m-%d %H:%M:%S")
        else:
            subscription_status = "❌ <b>Не активна</b>"
            subscription_end_formatted = "<i>Нет активной подписки</i>"
        
        private_users = read_private_users()
        is_vip = user.id in private_users["ids"] or (user.username and user.username in private_users["usernames"])
        vip_status = "✅ <b>Да</b>" if is_vip else "❌ <b>Нет</b>"
        
        if os.path.exists("База_Бота/users.txt"):
            with open("База_Бота/users.txt", "r") as file:
                user_ids = [line.strip() for line in file.readlines() if line.strip()]
                is_in_users = "✅ <b>Да</b>" if str(user.id) in user_ids else "❌ <b>Нет</b>"
        else:
            is_in_users = "❌ <b>Нет (файл не существует)</b>"
        
        is_banned = user_id in banned_users
        ban_button_text = "🔓 Разбанить" if is_banned else "🚫 Забанить"
        ban_button_data = f"unban_user:{user_id}" if is_banned else f"ban_user:{user_id}"
        
        user_info = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🔍 Информация о пользователе</b>

<b>🆔 ID:</b> <code>{user.id}</code>
<b>👤 Имя:</b> <code>{html.escape(user.first_name)}</code>
<b>👤 Фамилия:</b> <code>{html.escape(user.last_name) if user.last_name else 'не указана'}</code>
<b>👤 Юзернейм:</b> @{html.escape(user.username) if user.username else 'отсутствует'}

<b>🔐 Подписка:</b> {subscription_status}
<b>💰 Подписка до:</b> <code>{subscription_end_formatted}</code>
<b>⏳ Оставшееся время:</b> <code>{remaining_time}</code>
<b>🌟 VIP статус:</b> {vip_status}
<b>📝 В базе бота:</b> {is_in_users}
<b>🚫 Статус бана:</b> {"✅ <b>Да</b>" if is_banned else "❌ <b>Нет</b>"}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
        markup = InlineKeyboardMarkup(row_width=2)
        btn_ban = InlineKeyboardButton(ban_button_text, callback_data=ban_button_data)
        btn_privat = InlineKeyboardButton('👑 Vip', callback_data=f'toggle_vip:{user_id}')
        btn_user = InlineKeyboardButton('⏳ Подписка', callback_data=f'user:{user_id}')                
        markup.add(btn_ban, btn_privat, btn_user)
        
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=user_info,
            reply_markup=markup,
            parse_mode="HTML"
        )
        
        await callback_query.answer(f"VIP статус {action} для пользователя {user_id}")
    except Exception as e:
        await callback_query.answer(f"Ошибка при обновлении статуса VIP: {str(e)}", show_alert=True)

@dp.callback_query_handler(lambda c: c.data.startswith(('ban_user:', 'unban_user:')))
async def process_ban_unban(callback_query: types.CallbackQuery):
    action, user_id = callback_query.data.split(':')
    user_id = int(user_id)
    
    if action == 'ban_user':
        if user_id in banned_users:
            await callback_query.answer('🚫 Пользователь уже забанен')
            return
        
        banned_users.add(user_id)
        save_banned_users(banned_users)
        await callback_query.answer('✅ Пользователь забанен')
        try:
            await bot.send_message(user_id, '📢 Администратор посчитал ваш аккаунт подозрительным и вы были забанены 📢')
        except Exception as e:
            logging.error(f'Error sending ban message to user {user_id}: {e}')
            
    elif action == 'unban_user':
        if user_id not in banned_users:
            await callback_query.answer('🚫 Пользователь не забанен')
            return
        
        banned_users.remove(user_id)
        save_banned_users(banned_users)
        await callback_query.answer('✅ Пользователь разбанен')
        try:
            await bot.send_message(user_id, '📢 Ваш аккаунт был разбанен администратором 📢')
        except Exception as e:
            logging.error(f'Error sending unban message to user {user_id}: {e}')
    
    try:
        user = await bot.get_chat(user_id)
        subscription_end = await get_subscription_end_time(user_id)
        remaining_time = await get_remaining_time(user_id)
        
        if subscription_end == "forever":
            subscription_status = "✅ <b>Активна (Навсегда)</b>"
            subscription_end_formatted = "∞"
        elif subscription_end and subscription_end > datetime.now():
            subscription_status = "✅ <b>Активна</b>"
            subscription_end_formatted = subscription_end.strftime("%Y-%m-%d %H:%M:%S")
        else:
            subscription_status = "❌ <b>Не активна</b>"
            subscription_end_formatted = "<i>Нет активной подписки</i>"
        
        private_users = read_private_users()
        is_vip = user.id in private_users["ids"] or (user.username and user.username in private_users["usernames"])
        vip_status = "✅ <b>Да</b>" if is_vip else "❌ <b>Нет</b>"
        
        if os.path.exists("База_Бота/users.txt"):
            with open("База_Бота/users.txt", "r") as file:
                user_ids = [line.strip() for line in file.readlines() if line.strip()]
                is_in_users = "✅ <b>Да</b>" if str(user.id) in user_ids else "❌ <b>Нет</b>"
        else:
            is_in_users = "❌ <b>Нет (файл не существует)</b>"
        
        is_banned = user_id in banned_users
        ban_button_text = "🔓 Разбанить" if is_banned else "🚫 Забанить"
        ban_button_data = f"unban_user:{user_id}" if is_banned else f"ban_user:{user_id}"
        
        user_info = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🔍 Информация о пользователе</b>

<b>🆔 ID:</b> <code>{user.id}</code>
<b>👤 Имя:</b> <code>{html.escape(user.first_name)}</code>
<b>👤 Фамилия:</b> <code>{html.escape(user.last_name) if user.last_name else 'не указана'}</code>
<b>👤 Юзернейм:</b> @{html.escape(user.username) if user.username else 'отсутствует'}

<b>🔐 Подписка:</b> {subscription_status}
<b>💰 Подписка до:</b> <code>{subscription_end_formatted}</code>
<b>⏳ Оставшееся время:</b> <code>{remaining_time}</code>
<b>🌟 VIP статус:</b> {vip_status}
<b>📝 В базе бота:</b> {is_in_users}
<b>🚫 Статус бана:</b> {"✅ <b>Да</b>" if is_banned else "❌ <b>Нет</b>"}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
        markup = InlineKeyboardMarkup(row_width=2)
        btn_ban = InlineKeyboardButton(ban_button_text, callback_data=ban_button_data)
        btn_privat = InlineKeyboardButton('👑 Vip', callback_data=f'privat:{user_id}')
        btn_user = InlineKeyboardButton('⏳ Подписка', callback_data=f'user:{user_id}')
       
        
        markup.add(btn_ban, btn_privat, btn_user)
        
        await callback_query.message.edit_text(
            text=user_info,
            reply_markup=markup,
            parse_mode="HTML"
        )
        
    except Exception as e:
        await callback_query.answer(f'Ошибка при обновлении информации: {str(e)}', show_alert=True)

@dp.message_handler(state=BanState.waiting_for_ban_confirmation)
async def confirm_ban_user(message: types.Message, state: FSMContext):
    user_id = message.text
    if user_id.isdigit():
        user_id = int(user_id)
        if user_id in banned_users:
            await message.answer(f'🚫 Пользователь с ID {user_id} уже забанен.')
        else:
            banned_users.add(user_id)
            save_banned_users(banned_users)  
            await message.answer(f'✅ Пользователь с ID {user_id} забанен.')
            try:
                await bot.send_message(user_id, '📢 Администратор посчитал ваш аккаунт подозрительным и вы были забанены 📢')
            except Exception as e:
                logging.error(f'Ошибка отправки сообщения о бане пользователю {user_id}: {e}')
            await show_user_info_after_action(message, user_id)
    else:
        await message.answer('❌ Неверный формат ID. Пожалуйста, введите числовой ID.')
    await state.finish()

@dp.message_handler(state=BanState.waiting_for_unban_confirmation)
async def confirm_unban_user(message: types.Message, state: FSMContext):
    user_id = message.text
    if user_id.isdigit():
        user_id = int(user_id)
        if user_id not in banned_users:
            await message.answer(f'🚫 Пользователь с ID {user_id} не забанен.')
        else:
            banned_users.remove(user_id)
            save_banned_users(banned_users)  
            await message.answer(f'✅ Пользователь с ID {user_id} разбанен.')
            try:
                await bot.send_message(user_id, '📢 Ваш аккаунт был разбанен администратором 📢')
            except Exception as e:
                logging.error(f'Ошибка отправки сообщения о разбане пользователю {user_id}: {e}')
            await show_user_info_after_action(message, user_id)
    else:
        await message.answer('❌ Неверный формат ID. Пожалуйста, введите числовой ID.')
    await state.finish()

async def show_user_info_after_action(message: types.Message, user_id: int):
    try:
        user = await bot.get_chat(user_id)
        subscription_end = await get_subscription_end_time(user_id)
        remaining_time = await get_remaining_time(user_id)
        
        if subscription_end == "forever":
            subscription_status = "✅ <b>Активна (Навсегда)</b>"
            subscription_end_formatted = "∞"
        elif subscription_end and subscription_end > datetime.now():
            subscription_status = "✅ <b>Активна</b>"
            subscription_end_formatted = subscription_end.strftime("%Y-%m-%d %H:%M:%S")
        else:
            subscription_status = "❌ <b>Не активна</b>"
            subscription_end_formatted = "<i>Нет активной подписки</i>"
        
        private_users = read_private_users()
        is_vip = user.id in private_users["ids"] or (user.username and user.username in private_users["usernames"])
        vip_status = "✅ <b>Да</b>" if is_vip else "❌ <b>Нет</b>"
        
        if os.path.exists("База_Бота/users.txt"):
            with open("База_Бота/users.txt", "r") as file:
                user_ids = [line.strip() for line in file.readlines() if line.strip()]
                is_in_users = "✅ <b>Да</b>" if str(user.id) in user_ids else "❌ <b>Нет</b>"
        else:
            is_in_users = "❌ <b>Нет (файл не существует)</b>"
        
        is_banned = user_id in banned_users
        ban_button_text = "🔓 Разбанить" if is_banned else "🚫 Забанить"
        ban_button_data = f"unban_user:{user_id}" if is_banned else f"ban_user:{user_id}"
        
        user_info = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🔍 Информация о пользователе</b>

<b>🆔 ID:</b> <code>{user.id}</code>
<b>👤 Имя:</b> <code>{html.escape(user.first_name)}</code>
<b>👤 Фамилия:</b> <code>{html.escape(user.last_name) if user.last_name else 'не указана'}</code>
<b>👤 Юзернейм:</b> @{html.escape(user.username) if user.username else 'отсутствует'}

<b>🔐 Подписка:</b> {subscription_status}
<b>💰 Подписка до:</b> <code>{subscription_end_formatted}</code>
<b>⏳ Оставшееся время:</b> <code>{remaining_time}</code>
<b>🌟 VIP статус:</b> {vip_status}
<b>📝 В базе бота:</b> {is_in_users}
<b>🚫 Статус бана:</b> {"✅ <b>Да</b>" if is_banned else "❌ <b>Нет</b>"}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
        markup = InlineKeyboardMarkup(row_width=2)
        btn_ban = InlineKeyboardButton(ban_button_text, callback_data=ban_button_data)
        btn_privat = InlineKeyboardButton('👑 Vip', callback_data=f'privat:{user_id}')
        btn_user = InlineKeyboardButton('⏳ Подписка', callback_data=f'user:{user_id}')
        
        markup.add(btn_ban, btn_privat, btn_user)
        
        await bot.send_message(
            chat_id=message.chat.id,
            text=user_info,
            reply_markup=markup,
            parse_mode="HTML"
        )
    except Exception as e:
        await message.answer(f'Ошибка при отображении информации: {str(e)}')


@dp.callback_query_handler(lambda c: c.data == 'update_code')
async def update_code_callback(callback_query: types.CallbackQuery):
    await callback_query.answer("🔄 Обновление кода... Бот будет перезапущен и запущен одновременно", show_alert=True)
    await bot.send_message(callback_query.message.chat.id, "🔧 Идет процесс обновления, пожалуйста подождите...")
    await bot.close()
    os.system("python обновление.py")
    os._exit(0)

from aiogram.types import CallbackQuery  
from datetime import datetime
from aiogram.dispatcher.filters.state import State, StatesGroup
import json

NON_MAIN_API_ID = 2040
NON_MAIN_API_HASH = 'b18441a1ff607e10a989891a5462e627'

class CreateNonMainAccountStates(StatesGroup):
    phone = State()
    code = State()
    password = State()

@dp.callback_query_handler(lambda c: c.data == 'create_non_main_account', state='*')
async def create_non_main_account_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await bot.send_message(callback_query.from_user.id, '📱 Введите номер телефона (с кодом страны, можно с +):')
    await CreateNonMainAccountStates.phone.set()

@dp.message_handler(state=CreateNonMainAccountStates.phone)
async def process_non_main_phone_step(message: types.Message, state: FSMContext):
    user_id = message.from_user.id
    if user_id in banned_users:
        await message.answer('📢Администратор посчитал ваш аккаунт подозрительным и вы были забанены📢')
        return
    
    phone = message.text.replace('+', '') 
    if not phone or not phone.isdigit():
        await message.answer('❌ Введите корректный номер телефона.')
        return
    
    session_dir = "Не_основные"
    if not os.path.exists(session_dir):
        os.makedirs(session_dir)
    
    session_name = f"session_{phone}"
    session_path = os.path.join(session_dir, session_name)
    
    json_path = os.path.join(session_dir, f"{session_name}.json")
    if not os.path.exists(json_path):
        with open(json_path, 'w') as f:
            json.dump({
                "status": "Свободна",
                "phone": phone,
                "session_name": session_name,
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }, f, indent=4)
    
    client = TelegramClient(session_path, api_id=NON_MAIN_API_ID, api_hash=NON_MAIN_API_HASH)
    
    await client.connect()
    if not await client.is_user_authorized():
        try:
            result = await client.send_code_request(phone)
            phone_code_hash = result.phone_code_hash
            async with state.proxy() as data:
                data['phone'] = phone
                data['phone_code_hash'] = phone_code_hash
                data['session_dir'] = session_dir
                data['session_name'] = session_name
                data['attempts'] = 0
                data['password_attempts'] = 0
            await message.answer('📩 Введите код подтверждения:', reply_markup=create_code_keyboard())
            await CreateNonMainAccountStates.code.set()
        except errors.PhoneNumberInvalidError:
            await message.answer('❌ Неверный номер телефона. Пожалуйста, попробуйте еще раз.')
            await state.finish()
        except Exception as e:
            await message.answer(f'❌ Ошибка при отправке кода: {str(e)}')
            await state.finish()
        finally:
            await client.disconnect()
    else:
        await message.answer('❌ Аккаунт уже авторизован.')
        await state.finish()
        await client.disconnect()

user_code_input = {}

@dp.callback_query_handler(lambda c: c.data.startswith('code_'), state=CreateNonMainAccountStates.code)
async def process_code_buttons(callback_query: CallbackQuery, state: FSMContext):
    action = callback_query.data.split('_')[1]
    user_id = callback_query.from_user.id
    current_code = user_code_input.get(user_id, "")
    
    if action == "clear":
        current_code = ""
    elif action == "confirm":
        async with state.proxy() as data:
            data['code'] = current_code
        await callback_query.message.edit_text(f"🔐 Введён код: {current_code}\nПроверяем...")
        await process_non_main_code_step(callback_query.message, state)
        return
    elif action.isdigit():
        if len(current_code) < 5:
            current_code += action
    
    user_code_input[user_id] = current_code
    
    await callback_query.message.edit_text(
        f"📩 Введите код подтверждения: {'*' * len(current_code)}",
        reply_markup=create_code_keyboard()
    )  
    await callback_query.answer()

@dp.message_handler(state=CreateNonMainAccountStates.code)
async def process_non_main_code_step(message: types.Message, state: FSMContext):
    async with state.proxy() as data:
        code = data.get('code', message.text)
        phone = data['phone']
        phone_code_hash = data['phone_code_hash']
        session_dir = data['session_dir']
        session_name = data['session_name']
        attempts = data['attempts']
    
    if not code or len(code) != 5:
        await message.answer('❌ Введите корректный код подтверждения (5 цифр).', reply_markup=create_code_keyboard())
        return
    
    session_path = os.path.join(session_dir, session_name)
    client = TelegramClient(session_path, api_id=NON_MAIN_API_ID, api_hash=NON_MAIN_API_HASH)
    
    await client.connect()
    try:
        await client.sign_in(phone, code, phone_code_hash=phone_code_hash)
        
        json_path = os.path.join(session_dir, f"{session_name}.json")
        with open(json_path, 'r') as f:
            session_data = json.load(f)
        
        session_data["status"] = "Свободна"
        session_data["auth_date"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        with open(json_path, 'w') as f:
            json.dump(session_data, f, indent=4)
        
        await message.answer(f'✅ Аккаунт успешно создан и сохранен как {session_name}.session в папке "Не_основные"')
        await state.finish()
    except errors.SessionPasswordNeededError:
        await message.answer('🔒 Введите пароль от 2FA:')
        await CreateNonMainAccountStates.password.set()
    except errors.PhoneCodeInvalidError:
        attempts += 1
        if attempts >= 3:
            await message.answer('❌ Превышено количество попыток ввода кода. Начните заново.')
            await state.finish()
        else:
            async with state.proxy() as data:
                data['attempts'] = attempts
            await message.answer(f'❌ Неверный код. Осталось попыток: {3 - attempts}. Введите код снова:', reply_markup=create_code_keyboard())
    except Exception as e:
        await message.answer(f'❌ Ошибка при авторизации: {str(e)}')
        await state.finish()
    finally:
        await client.disconnect()

@dp.message_handler(state=CreateNonMainAccountStates.password)
async def process_non_main_password_step(message: types.Message, state: FSMContext):
    password = message.text
    async with state.proxy() as data:
        phone = data['phone']
        session_dir = data['session_dir']
        session_name = data['session_name']
        password_attempts = data['password_attempts']
    
    session_path = os.path.join(session_dir, session_name)
    client = TelegramClient(session_path, api_id=NON_MAIN_API_ID, api_hash=NON_MAIN_API_HASH)
    
    await client.connect()
    try:
        await client.sign_in(password=password)
        
        json_path = os.path.join(session_dir, f"{session_name}.json")
        with open(json_path, 'r') as f:
            session_data = json.load(f)
        
        session_data["status"] = "Свободна"
        session_data["auth_date"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        with open(json_path, 'w') as f:
            json.dump(session_data, f, indent=4)
        
        await message.answer(f'✅ Аккаунт успешно создан и сохранен как {session_name}.session в папке "Не_основные"')
        await state.finish()
    except errors.PasswordHashInvalidError:
        password_attempts += 1
        if password_attempts >= 3:
            await message.answer('❌ Превышено количество попыток ввода пароля. Начните заново.')
            await state.finish()
        else:
            async with state.proxy() as data:
                data['password_attempts'] = password_attempts
            await message.answer(f'❌ Неверный пароль. Осталось попыток: {3 - password_attempts}. Введите пароль снова:')
    except Exception as e:
        await message.answer(f'❌ Ошибка при вводе пароля: {str(e)}')
        await state.finish()
    finally:
        await client.disconnect()

PRIVATE_USERS_FILE = 'База_Бота/private_users.txt'



@dp.callback_query_handler(lambda c: c.data == 'privat', state='*')
async def privat_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    private_users = await update_and_validate_users(bot)
    
    has_ids = bool(private_users["ids"])
    has_usernames = bool(private_users["usernames"])
    
    if not has_ids and not has_usernames:
        users_list = """
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>👥 Список пользователей под приватом:</b>
<i>Список пуст</i>
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    else:
        users_list = """
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>👥 Список пользователей под приватом:</b>
"""
        if has_ids:
            users_list += f"\n<b>🆔 IDs:</b> <code>{', '.join(map(str, private_users['ids']))}</code>"
        
        if has_usernames:
            users_list += f"\n\n<b>📛 Usernames:</b> {', '.join(f'@{username}' for username in private_users['usernames'])}"
        
        users_list += "</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>"
    
    markup = InlineKeyboardMarkup(row_width=2)
    markup.add(
        InlineKeyboardButton('➕ Добавить привата', callback_data='add_private'),
        InlineKeyboardButton('➖ Удалить привата', callback_data='remove_private')
    )
    markup.add(InlineKeyboardButton('🔙 Назад', callback_data='user_menu'))
    
    try:
        if callback_query.message.photo:
            await bot.edit_message_caption(
                chat_id=callback_query.message.chat.id,
                message_id=callback_query.message.message_id,
                caption=users_list,
                reply_markup=markup,
                parse_mode="HTML"
            )
        else:
            await bot.edit_message_text(
                chat_id=callback_query.message.chat.id,
                message_id=callback_query.message.message_id,
                text=users_list,
                reply_markup=markup,
                parse_mode="HTML"
            )
    except Exception as e:
        print(f"Ошибка при редактировании сообщения: {e}")

@dp.callback_query_handler(lambda c: c.data == 'statsit', state='*')
async def statsit_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()

    if os.path.exists("База_Бота/users.txt"):
        with open("База_Бота/users.txt", "r") as file:
            user_ids = [line.strip() for line in file.readlines() if line.strip()]
            user_count = len(user_ids)
    else:
        user_count = 0

    statsit_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📊 Статистика бота:</b>

<b>👥 Пользователей использует бот:</b> <code>{user_count}</code>

<b>📌 Доступные действия:</b>
📥 <b>Извлечь ID пользователей</b> - Получить список ID пользователей.
📊 <b>Статистика бота</b> - Просмотр общей статистики.
📨 <b>Отправить сообщение</b> - Отправить сообщение всем пользователям.
🔙 <b>Назад</b> - Вернуться в админ-панель.
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""

    markup = InlineKeyboardMarkup(row_width=2)
    btn_extract_users = InlineKeyboardButton('📥 Извлечь ID', callback_data='extract_users')
    btn_stats = InlineKeyboardButton('📊 Статистика', callback_data='stats')
    btn_send_message = InlineKeyboardButton('📨 Рассылка', callback_data='send_message')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='user_menu')
    markup.add(btn_extract_users, btn_stats, btn_send_message)
    markup.add(btn_back)

    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=statsit_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=statsit_message,
            reply_markup=markup,
            parse_mode="HTML"
        )

@dp.callback_query_handler(lambda c: c.data == 'banis_user', state='*')
async def banis_user_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()

    if os.path.exists("База_Бота/banned_users.txt"):
        with open("База_Бота/banned_users.txt", "r", encoding="utf-8") as file:
            banned_users = [line.strip() for line in file.readlines() if line.strip()]
            banned_count = len(banned_users)
            banned_list = "\n• " + "\n• ".join(f"<code>{html.escape(user)}</code>" for user in banned_users)
    else:
        banned_count = 0
        banned_list = "<i>Нет пользователей в бане.</i>"

    ban_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🚫 Управление блокировками</b>

<b>📊 Количество пользователей в бане:</b> <code>{banned_count}</code>

<b>🆔 Список забаненных пользователей:</b>
{banned_list}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""

    markup = InlineKeyboardMarkup(row_width=2)
    btn_ban = InlineKeyboardButton('🚫 Забанить', callback_data='ban_user')
    btn_unban = InlineKeyboardButton('🔓 Разбанить', callback_data='unban_user')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='user_menu')
    markup.add(btn_ban, btn_unban)
    markup.add(btn_back)

    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=ban_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=ban_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    
@dp.callback_query_handler(lambda c: c.data == 'user', state='*')
async def user_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    
    paid_users_count = 0
    if os.path.exists('База_Бота/paid_users.txt'):
        with open('База_Бота/paid_users.txt', 'r') as file:
            paid_users_count = len([line for line in file if line.strip()])
    
    mes_text = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>⏳ Управление подписками</b>

<b>📊 Пользователей с подпиской:</b> <code>{paid_users_count}</code>

<b>📌 Доступные действия:</b>
➕ <b>Добавить</b> - нового пользователя
🗑️ <b>Удалить</b> - из системы
🕒 <b>Изменить время</b> - подписки
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    
    markup = InlineKeyboardMarkup(row_width=2)
    btn_add = InlineKeyboardButton('➕ Добавить', callback_data='add_user')
    btn_delete = InlineKeyboardButton('🗑️ Удалить', callback_data='delete_user') 
    btn_time = InlineKeyboardButton('🕒 Время', callback_data='change_time')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='user_menu')
    markup.add(btn_add, btn_delete, btn_time)
    markup.add(btn_back)

    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=mes_text,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=mes_text,
            reply_markup=markup,
            parse_mode="HTML"
        )
    
from aiogram.dispatcher.filters.state import State, StatesGroup
from datetime import datetime, timedelta
import re

class Form(StatesGroup):
    user_id = State()
    date = State()
    change_time_user_id = State()
    new_date = State()
    delete_user_id = State()

def read_paid_users():
    users = {}
    try:
        with open("База_Бота/paid_users.txt", "r") as file:
            for line in file:
                line = line.strip()
                if not line:
                    continue
                try:
                    user_id, date = line.split(',', 1)
                    users[user_id] = date.strip()
                except ValueError:
                    continue
    except FileNotFoundError:
        pass
    return users

def check_user_subscription(user_id):
    paid_users = read_paid_users()
    if str(user_id) not in paid_users:
        return (False, "не найден")
    
    date_str = paid_users[str(user_id)]
    if date_str.lower() == "forever":
        return (True, "вечная")
    
    try:
        end_time = datetime.strptime(date_str, "%Y-%m-%d %H:%M:%S")
        if datetime.now() < end_time:
            return (True, "активна")
        else:
            return (True, "просрочена")
    except ValueError:
        return (False, "ошибка формата даты")

def parse_time_input(time_input):
    time_input = time_input.strip().lower()
    
    if time_input == 'forever':
        return 'forever'
    
    try:
        duration = float(time_input)
        days = int(duration)
        hours = int((duration - days) * 24)
        expiry_date = datetime.now() + timedelta(days=days, hours=hours)
        return expiry_date.strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None

@dp.callback_query_handler(lambda c: c.data == 'add_user')
async def process_callback_add_user(callback_query: types.CallbackQuery):
    await bot.answer_callback_query(callback_query.id)
    await Form.user_id.set()
    await bot.send_message(callback_query.from_user.id, "🆔 Введите ID нового пользователя:")

@dp.callback_query_handler(lambda c: c.data == 'change_time')
async def process_callback_change_time(callback_query: types.CallbackQuery):
    await bot.answer_callback_query(callback_query.id)
    await Form.change_time_user_id.set()
    await bot.send_message(callback_query.from_user.id, "🆔 Введите ID пользователя для изменения времени:")

@dp.callback_query_handler(lambda c: c.data == 'delete_user')
async def process_callback_delete_user(callback_query: types.CallbackQuery):
    await bot.answer_callback_query(callback_query.id)
    await Form.delete_user_id.set()
    await bot.send_message(callback_query.from_user.id, "🆔 Введите ID пользователя для удаления:")

@dp.message_handler(state=Form.user_id)
async def process_user_id(message: types.Message, state: FSMContext):
    user_id = message.text.strip()
    exists, status = check_user_subscription(user_id)
    
    if exists:
        await state.finish()
        await message.reply(f"ℹ️ Пользователь {user_id} уже существует! Статус: {status}.")
        return
    
    async with state.proxy() as data:
        data['user_id'] = user_id
    
    await Form.date.set()
    await message.reply("📅 Введите срок подписки ('forever' или число дней.часов):")

@dp.message_handler(state=Form.change_time_user_id)
async def process_change_time_user_id(message: types.Message, state: FSMContext):
    user_id = message.text.strip()
    exists, status = check_user_subscription(user_id)
    
    if not exists:
        await state.finish()
        await message.reply(f"❌ Пользователь с ID {user_id} не найден.")
        return
    
    async with state.proxy() as data:
        data['user_id'] = user_id
        paid_users = read_paid_users()
        date_part = paid_users.get(str(user_id), "")
        
        if date_part.lower() == 'forever':
            await message.reply(f"🕒 Текущий статус: вечная подписка")
        else:
            await message.reply(f"🕒 Текущая дата окончания: {date_part}")
    
    await Form.new_date.set()
    await message.reply("📅 Введите новое время ('forever' или число дней.часов):")

@dp.message_handler(state=Form.delete_user_id)
async def process_delete_user_id(message: types.Message, state: FSMContext):
    user_id = message.text.strip()
    exists, _ = check_user_subscription(user_id)
    
    if not exists:
        await state.finish()
        await message.reply(f"❌ Пользователь с ID {user_id} не найден.")
        return
    
    try:
        with open("База_Бота/paid_users.txt", "r") as file:
            lines = file.readlines()
        
        with open("База_Бота/paid_users.txt", "w") as file:
            deleted = False
            for line in lines:
                if not line.startswith(f"{user_id},"):
                    file.write(line)
                else:
                    deleted = True
            
            if deleted:
                await message.reply(f"✅ Пользователь {user_id} успешно удалён!")
            else:
                await message.reply(f"❌ Пользователь с ID {user_id} не найден.")
    except Exception as e:
        await message.reply(f"❌ Ошибка при удалении: {e}")
    
    await state.finish()

@dp.message_handler(state=Form.date)
async def process_new_user_date(message: types.Message, state: FSMContext):
    time_input = message.text.strip()
    formatted_date = parse_time_input(time_input)
    
    if formatted_date is None:
        await message.reply("❌ Неверный формат! Введите 'forever' или число (напр. 1, 0.5, 1.12)")
        return
    
    async with state.proxy() as data:
        user_id = data['user_id']
        
        try:
            with open("База_Бота/paid_users.txt", "a") as file:
                file.write(f"{user_id},{formatted_date}\n")
            
            if formatted_date == 'forever':
                await message.reply(f"✅ Новый пользователь {user_id} добавлен навсегда!")
            else:
                await message.reply(f"✅ Новый пользователь {user_id} добавлен до {formatted_date}!")
        except Exception as e:
            await message.reply(f"❌ Ошибка: {e}")
    
    await state.finish()

@dp.message_handler(state=Form.new_date)
async def process_change_date(message: types.Message, state: FSMContext):
    time_input = message.text.strip()
    formatted_date = parse_time_input(time_input)
    
    if formatted_date is None:
        await message.reply("❌ Неверный формат! Введите 'forever' или число (напр. 1, 0.5, 1.12)")
        return
    
    async with state.proxy() as data:
        user_id = data['user_id']
        
        try:
            with open("База_Бота/paid_users.txt", "r") as file:
                lines = file.readlines()
            
            with open("База_Бота/paid_users.txt", "w") as file:
                updated = False
                for line in lines:
                    if line.startswith(f"{user_id},"):
                        file.write(f"{user_id},{formatted_date}\n")
                        updated = True
                    else:
                        file.write(line)
                
                if updated:
                    if formatted_date == 'forever':
                        await message.reply(f"✅ Пользователь {user_id} изменён на вечную подписку!")
                    else:
                        await message.reply(f"✅ Время для пользователя {user_id} изменено до {formatted_date}!")
                else:
                    await message.reply(f"❌ Пользователь с ID {user_id} не найден.")
        except Exception as e:
            await message.reply(f"❌ Ошибка: {e}")
    
    await state.finish()

            
@dp.callback_query_handler(lambda c: c.data == 'view_admins', state='*')
async def view_admins_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    markup = InlineKeyboardMarkup(row_width=2)
    btn_add_admin = InlineKeyboardButton('➕ Добавить админа', callback_data='add_admin')
    btn_remove_admin = InlineKeyboardButton('➖ Удалить админа', callback_data='remove_admin')
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='admin_panel')
    
    if admin_chat_ids:
        admins_list = "\n".join([
            f"👤 <code>{admin_id}</code>"
            for admin_id in admin_chat_ids
        ])
        admin_message = f"""<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🛡️ Список администраторов</b>
{admins_list}
<b>📊 Всего: {len(admin_chat_ids)}</b>
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
        markup.row(btn_add_admin, btn_remove_admin)
        markup.row(btn_back)
    else:
        admin_message = """<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🛡️ Список администраторов</b>
❌ Список пуст
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
        markup.row(btn_add_admin)
        markup.row(btn_back)

    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=admin_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=admin_message,
            reply_markup=markup,
            parse_mode="HTML"
        )

@dp.callback_query_handler(lambda c: c.data == 'add_admin', state='*')
async def add_admin_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("Введите ID пользователя, которого хотите назначить администратором:")
    await state.set_state("wait_for_admin_id")

@dp.message_handler(state="wait_for_admin_id")
async def process_admin_id(message: types.Message, state: FSMContext):
    user_id = message.text
    if user_id.isdigit():
        if user_id not in admin_chat_ids:
            admin_chat_ids.append(user_id)
            await message.answer(f"✅Пользователь с ID {user_id} успешно добавлен в список администраторов.")
            await bot.send_message(user_id, "📢Вы были назначены администратором.📢")
        else:
            await message.answer(f"❌Пользователь с ID {user_id} уже является администратором.❌")
    else:
        await message.answer("❌Некорректный ID. Пожалуйста, введите числовой ID.❌")
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'remove_admin', state='*')
async def remove_admin_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("Введите ID пользователя, которого хотите удалить из списка администраторов:")
    await state.set_state("wait_for_remove_admin_id")

@dp.message_handler(state="wait_for_remove_admin_id")
async def process_remove_admin_id(message: types.Message, state: FSMContext):
    user_id = message.text
    if user_id in admin_chat_ids:
        admin_chat_ids.remove(user_id)
        await message.answer(f"✅Пользователь с ID {user_id} успешно удален из списка администраторов.✅")
    else:
        await message.answer(f"❌Пользователь с ID {user_id} не найден в списке администраторов.❌")
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'back_to_main_menu', state='*')
async def back_to_main_menu_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    markup = InlineKeyboardMarkup(row_width=2)
    btn_support = InlineKeyboardButton('📩 Написать поддержку📩', callback_data='support')
    btn_demolition = InlineKeyboardButton('☠️Snos☠️', callback_data='demolition')  
    btn_restore_account = InlineKeyboardButton('🔄Реснуть акк🔄', callback_data='restore_account')
    btn_my_time = InlineKeyboardButton('👤Профиль👤', callback_data='my_time')
    btn_spam_menu = InlineKeyboardButton('🔥Spam🔥', callback_data='spam_menu')
    btn_osint = InlineKeyboardButton('🔍Osint🔎', callback_data='osint')
    btn_instruction = InlineKeyboardButton('📚 Инструкция', callback_data='instruction') 
    markup.add(btn_demolition, btn_osint)
    markup.add(btn_spam_menu)
    markup.add(btn_support, btn_restore_account)
    markup.add(btn_my_time, btn_instruction)    
    if str(callback_query.from_user.id) in admin_chat_ids:
        btn_admin_panel = InlineKeyboardButton('🛠Админ панель🛠', callback_data='admin_panel')
        markup.add(btn_admin_panel)
    await callback_query.message.edit_reply_markup(reply_markup=markup)
    
@dp.callback_query_handler(lambda c: c.data == 'add_private', state='*')
async def add_private_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("➕ Введите ID или @username пользователя для добавления в приват:")
    await state.set_state("waiting_for_private_add")

async def get_user_info(user_input: str):
    try:
        if user_input.isdigit():
            user = await bot.get_chat(user_input)
            return user.id, user.username
        else:
            username = user_input.lstrip('@')
            user = await bot.get_chat(f"@{username}")
            return user.id, user.username
    except Exception:
        return None, None
        
async def update_and_validate_users(bot: Bot):
    private_users = read_private_users()
    updated_data = {"ids": [], "usernames": []}
    valid_usernames = set()
    for user_id in private_users["ids"]:
        try:
            user = await bot.get_chat(user_id)
            updated_data["ids"].append(user_id)
            if user.username:
                clean_username = user.username.lstrip('@')
                valid_usernames.add(clean_username)
        except Exception as e:
            print(f"Ошибка при получении данных пользователя {user_id}: {e}")
            updated_data["ids"].append(user_id)
    for username in private_users["usernames"]:
        try:
            user = await bot.get_chat(f"@{username}")
            if user.id in private_users["ids"]:
                valid_usernames.add(username.lstrip('@'))
        except Exception:
            continue
    
    updated_data["usernames"] = list(valid_usernames)
    
    if ((set(private_users["ids"]) != set(updated_data["ids"])) or 
        (set(private_users["usernames"]) != set(updated_data["usernames"]))):
        write_private_users(updated_data)
    
    return updated_data

@dp.message_handler(state="waiting_for_private_add")
async def process_add_private(message: types.Message, state: FSMContext):
    user_input = message.text.strip()
    private_users = read_private_users()
    user_id, username = await get_user_info(user_input)
    clean_username = username.lstrip('@') if username else None
    
    added_ids = []
    added_usernames = []
    if user_id is not None and user_id not in private_users["ids"]:
        private_users["ids"].append(user_id)
        added_ids.append(str(user_id))
    if clean_username is not None and clean_username not in private_users["usernames"]:
        private_users["usernames"].append(clean_username)
        added_usernames.append(clean_username)
    
    if user_id is None and clean_username is None:
        if user_input.isdigit():
            user_id = int(user_input)
            if user_id not in private_users["ids"]:
                private_users["ids"].append(user_id)
                added_ids.append(str(user_id))
        else:
            username = user_input.lstrip('@')
            if username not in private_users["usernames"]:
                private_users["usernames"].append(username)
                added_usernames.append(username)
    
    if added_ids or added_usernames:
        write_private_users(private_users)
        response = "✅ Добавлено:"
        if added_ids:
            response += f"\nID: {', '.join(added_ids)}"
        if added_usernames:
            response += f"\nUsername: @{', @'.join(added_usernames)}"
        await message.answer(response)
    else:
        response = "❌ Уже есть в привате:"
        if user_id and user_id in private_users["ids"]:
            response += f"\nID: {user_id}"
        if clean_username and clean_username in private_users["usernames"]:
            response += f"\nUsername: @{clean_username}"
        await message.answer(response)
    
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'remove_private', state='*')
async def remove_private_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("➖ Введите ID или @username пользователя для удаления из привата:")
    await state.set_state("waiting_for_private_remove")

@dp.message_handler(state="waiting_for_private_remove")
async def process_remove_private(message: types.Message, state: FSMContext):
    user_input = message.text.strip()
    private_users = read_private_users()
    user_id, username = await get_user_info(user_input)
    clean_username = username.lstrip('@') if username else None
    
    removed_ids = []
    removed_usernames = []
    if user_id is not None and user_id in private_users["ids"]:
        private_users["ids"].remove(user_id)
        removed_ids.append(str(user_id))
    if clean_username is not None and clean_username in private_users["usernames"]:
        private_users["usernames"].remove(clean_username)
        removed_usernames.append(clean_username)
    if not removed_ids and not removed_usernames:
        if user_input.isdigit():
            user_id = int(user_input)
            if user_id in private_users["ids"]:
                private_users["ids"].remove(user_id)
                removed_ids.append(str(user_id))
        else:
            username = user_input.lstrip('@')
            if username in private_users["usernames"]:
                private_users["usernames"].remove(username)
                removed_usernames.append(username)
    
    if removed_ids or removed_usernames:
        write_private_users(private_users)
        response = "✅ Удалено:"
        if removed_ids:
            response += f"\nID: {', '.join(removed_ids)}"
        if removed_usernames:
            response += f"\nUsername: @{', @'.join(removed_usernames)}"
        await message.answer(response)
    else:
        response = "❌ Не найдено в привате:"
        if user_input.isdigit():
            response += f"\nID: {user_input}"
        else:
            response += f"\nUsername: @{user_input.lstrip('@')}"
        await message.answer(response)
    
    await state.finish()             

@dp.callback_query_handler(lambda c: c.data == 'ban_user', state='*')
async def ban_user_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer('📝Введите ID пользователя, которого хотите забанить:')
    await BanState.waiting_for_ban_user_id.set()

@dp.message_handler(state=BanState.waiting_for_ban_user_id)
async def ban_user_input(message: types.Message, state: FSMContext):
    user_id = message.text
    if user_id.isdigit():
        user_id = int(user_id)
        if user_id in banned_users:
            await message.answer(f'🚫 Пользователь с ID {user_id} уже забанен.')
        else:
            banned_users.add(user_id)
            save_banned_users(banned_users)  
            await message.answer(f'✅ Пользователь с ID {user_id} забанен.')
            try:
                await bot.send_message(user_id, '📢Администратор посчитал ваш аккаунт подозрительным и вы были забанены📢')
            except Exception as e:
                logging.error(f'Error sending ban message to user {user_id}: {e}')
    else:
        await message.answer('❌ Неверный формат ID. Пожалуйста, введите числовой ID.')
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'unban_user', state='*')
async def unban_user_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer('📝Введите ID пользователя, которого хотите разбанить:')
    await BanState.waiting_for_unban_user_id.set()

@dp.message_handler(state=BanState.waiting_for_unban_user_id)
async def unban_user_input(message: types.Message, state: FSMContext):
    user_id = message.text
    if user_id.isdigit():
        user_id = int(user_id)
        if user_id not in banned_users:
            await message.answer(f'🚫 Пользователь с ID {user_id} не забанен.')
        else:
            banned_users.remove(user_id)
            save_banned_users(banned_users)  
            await message.answer(f'✅ Пользователь с ID {user_id} разбанен.')
            try:
                await bot.send_message(user_id, '📢Ваш аккаунт был разбанен администратором📢')
            except Exception as e:
                logging.error(f'Error sending unban message to user {user_id}: {e}')
    else:
        await message.answer('❌ Неверный формат ID. Пожалуйста, введите числовой ID.')
    await state.finish()
    
from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup, InputMediaPhoto
from datetime import datetime

#--------



from aiogram.dispatcher.filters.state import State, StatesGroup

class PromoStates(StatesGroup):
    promo_name = State()
    promo_end_date = State()
    promo_duration = State()
    promo_activations = State()
    edit_promo_select = State()
    edit_promo_choice = State()
    edit_promo_new_name = State()
    edit_promo_new_end_date = State()
    edit_promo_new_duration = State()
    edit_promo_new_activations = State()


@dp.callback_query_handler(lambda c: c.data == 'promocodes_menu', state='*')
async def promocodes_menu_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()

    try:
        promocodes = read_promocodes()
        promo_count = len(promocodes)
        
        if promo_count > 0:
            promo_list = []
            for promo in promocodes:
                name = promo.get("promo_code", "Без названия")
                activations = promo.get("activations", "0")
                end_date = promo.get("end_date", "Не указано")
                activated_count = len(promo.get("activated_users", []))
                
                if activations.lower() == "бесконечно":
                    remaining = "∞"
                elif activations.isdigit():
                    total = int(activations)
                    remaining = max(0, total - activated_count)
                else:
                    remaining = "?"
                
                promo_list.append(
                    f"• <code>{html.escape(name)}</code> (до {end_date}): "
                    f"<b>{remaining}</b> из {activations if activations.isdigit() else activations.lower()}"
                )
        else:
            promo_list = ["<i>Промокоды отсутствуют.</i>"]
    except Exception as e:
        print(f"Ошибка при чтении промокодов: {e}")
        promo_count = 0
        promo_list = ["<i>Ошибка при загрузке промокодов</i>"]

    promo_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>🎫 Управление промокодами</b>

<b>📊 Всего промокодов:</b> <code>{promo_count}</code>

<b>📝 Список промокодов:</b>
{"<br>".join(promo_list)}
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""

    markup = InlineKeyboardMarkup(row_width=2)
    markup.add(
        InlineKeyboardButton('🎫 Создать', callback_data='create_promo'),
        InlineKeyboardButton('❌ Удалить', callback_data='delete_promo'),
        InlineKeyboardButton('✏️ Редактировать', callback_data='edit_promo')
    )
    markup.add(InlineKeyboardButton('🔙 Назад', callback_data='admin_panel'))

    try:
        if callback_query.message.photo:
            await bot.edit_message_caption(
                chat_id=callback_query.message.chat.id,
                message_id=callback_query.message.message_id,
                caption=promo_message,
                reply_markup=markup,
                parse_mode="HTML"
            )
        else:
            await bot.edit_message_text(
                chat_id=callback_query.message.chat.id,
                message_id=callback_query.message.message_id,
                text=promo_message,
                reply_markup=markup,
                parse_mode="HTML"
            )
    except Exception as e:
        print(f"Ошибка при редактировании сообщения: {e}")

@dp.message_handler(state=PromoStates.promo_name)
async def process_promo_name(message: types.Message, state: FSMContext):
    promo_name = message.text
    promocodes = read_promocodes()
    
    if any(p["promo_code"] == promo_name for p in promocodes):
        await message.answer("❌ Промокод с таким названием уже создан. Введите другое название:")
        return
    
    await state.update_data(promo_name=promo_name)
    await message.answer("⏳ Введите срок действия промокода (в днях или дате ГГГГ-ММ-ДД ЧЧ:ММ:СС):")
    await state.set_state(PromoStates.promo_end_date)

@dp.callback_query_handler(lambda c: c.data == 'edit_promo', state='*')
async def edit_promo_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("✏️ Введите название промокода для редактирования:")
    await state.set_state("edit_promo_select")

@dp.message_handler(state="edit_promo_select")
async def process_edit_promo_select(message: types.Message, state: FSMContext):
    promo_name = message.text
    promocodes = read_promocodes()
    
    promo = next((p for p in promocodes if p["promo_code"] == promo_name), None)
    if not promo:
        await message.answer("❌ Промокод не найден.")
        await state.finish()
        return
    
    await state.update_data(edit_promo=promo_name)
    
    markup = InlineKeyboardMarkup(row_width=2)
    markup.add(
        InlineKeyboardButton("Название", callback_data="edit_promo_name"),
        InlineKeyboardButton("Срок действия", callback_data="edit_promo_end_date"),
        InlineKeyboardButton("Время подписки", callback_data="edit_promo_duration"),
        InlineKeyboardButton("Активации", callback_data="edit_promo_activations"),
        InlineKeyboardButton("Отмена", callback_data="promocodes_menu")
    )
    
    await message.answer(
        f"✏️ Выберите что изменить в промокоде <code>{promo_name}</code>:",
        reply_markup=markup,
        parse_mode="HTML"
    )
    await state.set_state("edit_promo_choice")

@dp.callback_query_handler(state="edit_promo_choice")
async def process_edit_promo_choice(callback_query: types.CallbackQuery, state: FSMContext):
    choice = callback_query.data
    await callback_query.answer()
    
    if choice == "edit_promo_name":
        await callback_query.message.answer("✏️ Введите новое название промокода:")
        await state.set_state("edit_promo_new_name")
    elif choice == "edit_promo_end_date":
        await callback_query.message.answer("✏️ Введите новый срок действия (в днях или ГГГГ-ММ-ДД ЧЧ:ММ:СС):")
        await state.set_state("edit_promo_new_end_date")
    elif choice == "edit_promo_duration":
        await callback_query.message.answer("✏️ Введите новое время подписки (дни.часы или forever):")
        await state.set_state("edit_promo_new_duration")
    elif choice == "edit_promo_activations":
        await callback_query.message.answer("✏️ Введите новое количество активаций (число или бесконечно):")
        await state.set_state("edit_promo_new_activations")
    else:
        await state.finish()
        await promocodes_menu_callback(callback_query, state)

@dp.message_handler(state="edit_promo_new_name")
async def process_edit_promo_new_name(message: types.Message, state: FSMContext):
    new_name = message.text
    data = await state.get_data()
    promocodes = read_promocodes()
    
    if any(p["promo_code"] == new_name for p in promocodes):
        await message.answer("❌ Промокод с таким названием уже существует.")
        return
    
    for promo in promocodes:
        if promo["promo_code"] == data["edit_promo"]:
            promo["promo_code"] = new_name
            break
    
    write_promocodes(promocodes)
    await message.answer(f"✅ Название промокода успешно изменено на <code>{new_name}</code>", parse_mode="HTML")
    await state.finish()

@dp.message_handler(state="edit_promo_new_end_date")
async def process_edit_promo_new_end_date(message: types.Message, state: FSMContext):
    try:
        if message.text.isdigit():
            end_date = datetime.now() + timedelta(days=int(message.text))
        else:
            end_date = datetime.strptime(message.text, "%Y-%m-%d %H:%M:%S")
        
        data = await state.get_data()
        promocodes = read_promocodes()
        
        for promo in promocodes:
            if promo["promo_code"] == data["edit_promo"]:
                promo["end_date"] = end_date.strftime("%Y-%m-%d %H:%M:%S")
                break
        
        write_promocodes(promocodes)
        await message.answer(f"✅ Срок действия промокода изменён на <code>{end_date.strftime('%Y-%m-%d %H:%M:%S')}</code>", parse_mode="HTML")
        await state.finish()
    except ValueError:
        await message.answer("❌ Неверный формат. Введите количество дней или дату в формате ГГГГ-ММ-ДД ЧЧ:ММ:СС.")

@dp.message_handler(state="edit_promo_new_duration")
async def process_edit_promo_new_duration(message: types.Message, state: FSMContext):
    duration = parse_duration_input(message.text)
    if duration is None:
        await message.answer("❌ Неверный формат. Введите дни.часы (1.5) или 'forever'")
        return
    
    data = await state.get_data()
    promocodes = read_promocodes()
    
    duration_str = "forever" if duration == "forever" else format_duration(duration)
    
    for promo in promocodes:
        if promo["promo_code"] == data["edit_promo"]:
            promo["duration"] = duration_str
            break
    
    write_promocodes(promocodes)
    await message.answer(f"✅ Время подписки изменено на <code>{duration_str}</code>", parse_mode="HTML")
    await state.finish()

@dp.message_handler(state="edit_promo_new_activations")
async def process_edit_promo_new_activations(message: types.Message, state: FSMContext):
    data = await state.get_data()
    promocodes = read_promocodes()
    
    for promo in promocodes:
        if promo["promo_code"] == data["edit_promo"]:
            promo["activations"] = message.text
            break
    
    write_promocodes(promocodes)
    await message.answer(f"✅ Количество активаций изменено на <code>{message.text}</code>", parse_mode="HTML")
    await state.finish()


def read_promocodes():
    if not os.path.exists("База_Бота/promocodes.json"):
        return []
    try:
        with open("База_Бота/promocodes.json", "r", encoding="utf-8") as file:
            return json.load(file)
    except:
        return []

def write_promocodes(promocodes):
    with open("База_Бота/promocodes.json", "w", encoding="utf-8") as file:
        json.dump(promocodes, file, ensure_ascii=False, indent=4)

def read_paid_users():
    paid_users = {}
    if not os.path.exists("База_Бота/paid_users.txt"):
        return paid_users
        
    with open("База_Бота/paid_users.txt", "r", encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line:
                continue
            user_id, time_str = line.split(",", 1)
            paid_users[user_id] = time_str
    return paid_users

def write_paid_users(paid_users):
    with open("База_Бота/paid_users.txt", "w", encoding="utf-8") as file:
        for user_id, time_str in paid_users.items():
            file.write(f"{user_id},{time_str}\n")

def parse_duration_input(duration_input):
    if duration_input.lower() == "forever":
        return "forever"
    try:
        if "." in duration_input:
            days, hours = map(float, duration_input.split("."))
            return timedelta(days=days, hours=hours)
        else:
            return timedelta(days=int(duration_input))
    except:
        return None

def format_duration(duration):
    if duration == "forever":
        return "forever"
    total_seconds = duration.total_seconds()
    days = int(total_seconds // 86400)
    hours = int((total_seconds % 86400) // 3600)
    return f"{days}.{hours:02d}"

@dp.callback_query_handler(lambda c: c.data == 'create_promo', state='*')
async def create_promo_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("🎫 Введите название промокода:")
    await state.set_state(PromoStates.promo_name)

@dp.message_handler(state=PromoStates.promo_end_date)
async def process_promo_end_date(message: types.Message, state: FSMContext):
    try:
        if message.text.isdigit():
            end_date = datetime.now() + timedelta(days=int(message.text))
        else:
            end_date = datetime.strptime(message.text, "%Y-%m-%d %H:%M:%S")
        await state.update_data(promo_end_date=end_date)
        await message.answer("⏳ Введите время, которое даёт промокод (в днях или днях.часах, например 1.5 = 1 день и 12 часов, или 'forever'):")
        await state.set_state(PromoStates.promo_duration)
    except ValueError:
        await message.answer("❌ Неверный формат. Введите количество дней или дату в формате ГГГГ-ММ-ДД ЧЧ:ММ:СС.")
        return

@dp.message_handler(state=PromoStates.promo_duration)
async def process_promo_duration(message: types.Message, state: FSMContext):
    duration = parse_duration_input(message.text)
    if duration is None:
        await message.answer("❌ Неверный формат. Введите количество дней (например 7) или дни.часы (например 1.5) или 'forever'")
        return
    await state.update_data(promo_duration=duration)
    await message.answer("🔢 Введите количество активаций (число или 'бесконечно'):")
    await state.set_state(PromoStates.promo_activations)

@dp.message_handler(state=PromoStates.promo_activations)
async def process_promo_activations(message: types.Message, state: FSMContext):
    data = await state.get_data()
    promo_name = data['promo_name']
    promo_end_date = data['promo_end_date'].strftime("%Y-%m-%d %H:%M:%S")
    promo_duration = data['promo_duration']
    
    duration_str = "forever" if promo_duration == "forever" else format_duration(promo_duration)
    activations = message.text
    
    promocodes = read_promocodes()
    promocodes.append({
        "promo_code": promo_name,
        "end_date": promo_end_date,
        "duration": duration_str,
        "activations": activations,
        "activated_users": []
    })
    write_promocodes(promocodes)
    
    duration_display = "вечное" if duration_str == "forever" else duration_str
    await message.answer(
        f"<b>───── ⋆⋅☆⋅⋆ ─────</b>\n<blockquote><b>🎉 Промокод успешно создан!</b>\n"
        f"<b>🎫 Название:</b> <code>{promo_name}</code>\n"
        f"<b>⏳ Срок действия до:</b> <code>{promo_end_date}</code>\n"
        f"<b>⏳ Время, которое даёт:</b> <code>{duration_display}</code>\n"
        f"<b>🔢 Активаций:</b> <code>{activations}</code></blockquote>\n<b>───── ⋆⋅☆⋅⋆ ─────</b>",
        parse_mode="HTML"
    )
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'delete_promo', state='*')
async def delete_promo_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("❌ Введите название промокода для удаления:")
    await state.set_state("delete_promo_name")

@dp.message_handler(state="delete_promo_name")
async def process_delete_promo_name(message: types.Message, state: FSMContext):
    promo_name = message.text
    promocodes = read_promocodes()
    updated_promocodes = [p for p in promocodes if p['promo_code'] != promo_name]
    
    if len(updated_promocodes) == len(promocodes):
        await message.answer(f"❌ Промокод '{promo_name}' не найден.")
    else:
        write_promocodes(updated_promocodes)
        await message.answer(f"❌ Промокод '{promo_name}' успешно удалён.")
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == 'activate_promo', state='*')
async def activate_promo_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    await callback_query.message.answer("Введите промокод для активации:")
    await state.set_state("activate_promo_name")

@dp.message_handler(state="activate_promo_name")
async def process_activate_promo_name(message: types.Message, state: FSMContext):
    promo_name = message.text
    user_id = str(message.from_user.id)
    current_time = datetime.now()
    
    promocodes = read_promocodes()
    promo = next((p for p in promocodes if p["promo_code"] == promo_name), None)
    
    if not promo:
        await message.answer("❌ Промокод не найден.")
        await state.finish()
        return
    
    if user_id in promo["activated_users"]:
        await message.answer("❌ Вы уже активировали этот промокод.")
        await state.finish()
        return
    
    end_date = datetime.strptime(promo["end_date"], "%Y-%m-%d %H:%M:%S")
    if current_time > end_date:
        await message.answer("❌ Срок действия промокода истёк.")
        await state.finish()
        return
    
    activations_left = -1 if promo["activations"].lower() == "бесконечно" else int(promo["activations"])
    if activations_left == 0:
        await message.answer("❌ Лимит активаций исчерпан.")
        await state.finish()
        return
    
    if activations_left > 0:
        activations_left -= 1
        promo["activations"] = str(activations_left) if activations_left > 0 else "0"
    
    promo["activated_users"].append(user_id)
    write_promocodes(promocodes)
    
    paid_users = read_paid_users()
    duration = promo["duration"]
    
    if duration == "forever":
        paid_users[user_id] = "forever"
    else:
        days, hours = map(float, duration.split("."))
        duration_td = timedelta(days=days, hours=hours)
        new_end_time = current_time + duration_td
        
        if user_id in paid_users:
            if paid_users[user_id] == "forever":
                pass
            else:
                user_end = datetime.strptime(paid_users[user_id], "%Y-%m-%d %H:%M:%S")
                if new_end_time > user_end:
                    paid_users[user_id] = new_end_time.strftime("%Y-%m-%d %H:%M:%S")
        else:
            paid_users[user_id] = new_end_time.strftime("%Y-%m-%d %H:%M:%S")
    
    write_paid_users(paid_users)
    
    duration_display = "вечное" if duration == "forever" else duration
    activations_display = "бесконечно" if activations_left == -1 else activations_left
    
    keyboard = InlineKeyboardMarkup()
    keyboard.add(InlineKeyboardButton("Запуск", callback_data="send_welcome"))
    
    await message.answer(
        f"<b>───── ⋆⋅☆⋅⋆ ─────</b>\n<blockquote>🎉 Промокод '{promo_name}' активирован!\n"
        f"👤 Пользователь: {user_id}\n"
        f"⏳ Получено время: {duration_display}\n"
        f"🔢 Осталось активаций: {activations_display}</blockquote>\n<b>───── ⋆⋅☆⋅⋆ ─────</b>",
        reply_markup=keyboard,
        parse_mode="HTML"
    )
    await state.finish()



#--------




async def get_remaining_time(user_id):
    if str(user_id) in admin_chat_ids:
        return "∞ (Администратор)"
    if not os.path.exists('База_Бота/paid_users.txt'):
        return "Нет доступа"
    try:
        with open('База_Бота/paid_users.txt', 'r', encoding='utf-8') as file:
            for line in file:
                line = line.strip()
                if not line:
                    continue
                parts = line.split(',')
                if len(parts) != 2:
                    continue
                paid_user_id, expiry_time_str = parts
                if paid_user_id == str(user_id):
                    if expiry_time_str == "forever":
                        return "∞ (Навсегда)"
                    expiry_time = datetime.strptime(expiry_time_str, '%Y-%m-%d %H:%M:%S')
                    remaining_time = expiry_time - datetime.now()
                    if remaining_time.total_seconds() > 0:
                        days = remaining_time.days
                        hours, remainder = divmod(remaining_time.seconds, 3600)
                        minutes, seconds = divmod(remainder, 60)
                        return f"{days} дней, {hours} часов, {minutes} минут, {seconds} секунд"
                    else:
                        return "Время истекло"
    except Exception as e:
        print(f"Ошибка при чтении файла paid_users.txt: {e}")
    return "Нет доступа"

@dp.callback_query_handler(lambda c: c.data == 'to_start')
async def process_callback_back_to_start(callback_query: types.CallbackQuery):
    user_id = callback_query.from_user.id
    first_name = callback_query.from_user.first_name or ""
    last_name = callback_query.from_user.last_name or ""
    username = f"@{callback_query.from_user.username}" if callback_query.from_user.username else f"id{callback_query.from_user.id}"
    
    welcome_message = PAID_WELCOME_BANNER.format(
        first_name=first_name,
        last_name=last_name,
        username=username
    )
    
    markup = InlineKeyboardMarkup(row_width=2)
    btn_support = InlineKeyboardButton('📩 Написать поддержку📩', callback_data='support')
    btn_demolition = InlineKeyboardButton('☠️Snos☠️', callback_data='demolition')  
    btn_restore_account = InlineKeyboardButton('🔄Pеснуть акк🔄', callback_data='restore_account')
    btn_my_time = InlineKeyboardButton('👤Профиль👤', callback_data='my_time')
    btn_spam_menu = InlineKeyboardButton('🔥Spam🔥', callback_data='spam_menu')
    btn_osint = InlineKeyboardButton('🔍Osint🔎', callback_data='osint')
    btn_instruction = InlineKeyboardButton('📚 Инструкция', callback_data='instruction') 
    markup.add(btn_demolition, btn_osint)
    markup.add(btn_spam_menu)
    markup.add(btn_support, btn_restore_account)
    markup.add(btn_my_time, btn_instruction)  
    
    if str(user_id) in admin_chat_ids:  
        btn_admin_panel = InlineKeyboardButton('🛠 Админ панель', callback_data='admin_panel')
        markup.add(btn_admin_panel)
    
    await bot.answer_callback_query(callback_query.id)
    
    if callback_query.message.photo:
        photo = callback_query.message.photo[-1].file_id
        media = InputMediaPhoto(media=photo, caption=welcome_message, parse_mode="HTML")
        await bot.edit_message_media(
            chat_id=user_id,
            message_id=callback_query.message.message_id,
            media=media,
            reply_markup=markup
        )
    else:
        await bot.edit_message_text(
            chat_id=user_id,
            message_id=callback_query.message.message_id,
            text=welcome_message,
            reply_markup=markup,
            parse_mode="HTML"
        )

@dp.callback_query_handler(lambda c: c.data == 'instruction')
async def process_callback_instruction(callback_query: types.CallbackQuery):
    await callback_query.answer(
        "⚠️ Внимательно ознакомьтесь с инструкцией и правилами ниже",
        show_alert=True
    )
    
    instruction_message = """<b>───── ⋆⋅☆⋅⋆ ─────</b>\n<blockquote>
<b>📚 ПОЛНАЯ ИНСТРУКЦИЯ И ПРАВИЛА</b>  

<u>🔧 Основные функции:</u>  

<code>☠️ Snos</code> – полное удаление аккаунта/канала Telegram  
• Botnet  
• Gmail  

<code>🔍 Osint</code> – сбор информации  
• Шерлок  
• Вектор  
• Файлы БД  
• VK OSINT  

<code>🔥 Spam</code> – инструменты спам-рассылки  
• Gmail  
• Фул кодами на ТГ  

<u>⚖️ ПРАВИЛА ИСПОЛЬЗОВАНИЯ:</u>  

1. Запрещено: 
   • Проводить нагрузку на бота  
   • Спамить поддержку  
   • Предлагать рекламу  

2. Обязательно:
   • Соблюдать правила  
   • Ознакомиться с инструкцией  

<u>❗ ПОСЛЕДСТВИЯ НАРУШЕНИЙ:</u>  

• Первое нарушение – предупреждение  
• Повторное – временная блокировка/уменьшение времени подписки  
• Серьёзные нарушения – полный бан  

Для продолжения выберите нужный раздел в меню.  
Скрыть инструкцию – кнопка ниже.  
</blockquote>\n<b>───── ⋆⋅☆⋅⋆ ─────</b>"""
    
 
    markup = InlineKeyboardMarkup()
    markup.add(InlineKeyboardButton("❌ Скрыть", callback_data="hide_instruction"))
    
    try:
        sent_message = await bot.send_message(
            chat_id=callback_query.from_user.id,
            text=instruction_message,
            parse_mode="HTML",
            reply_markup=markup,
            disable_web_page_preview=True
        )
        
        await state.update_data(instruction_msg_id=sent_message.message_id)
        
    except Exception as e:
        print(f"Ошибка при отправке инструкции: {e}")
        await callback_query.answer(
            "⚠️ Не удалось отправить инструкцию. Попробуйте позже.",
            show_alert=True
        )

@dp.callback_query_handler(lambda c: c.data == 'hide_instruction')
async def hide_instruction(callback_query: types.CallbackQuery):
    try:
        await bot.delete_message(
            chat_id=callback_query.from_user.id,
            message_id=callback_query.message.message_id
        )
        await callback_query.answer("Инструкция скрыта", show_alert=False)
    except Exception as e:
        print(f"Ошибка при удалении сообщения: {e}")
        await callback_query.answer("Не удалось скрыть инструкцию", show_alert=False)
              
        
session_dir = "Session"
if not os.path.exists(session_dir):
    os.makedirs(session_dir)

for client in clients:
    client_folder = os.path.join(session_dir, client["name"])
    if not os.path.exists(client_folder):
        os.makedirs(client_folder)

def get_random_client():
    return random.choice(clients)

def create_code_keyboard():
    keyboard = InlineKeyboardMarkup(row_width=3)
    keyboard.row(
        InlineKeyboardButton("1", callback_data="code_1"),
        InlineKeyboardButton("2", callback_data="code_2"),
        InlineKeyboardButton("3", callback_data="code_3")
    )
    keyboard.row(
        InlineKeyboardButton("4", callback_data="code_4"),
        InlineKeyboardButton("5", callback_data="code_5"),
        InlineKeyboardButton("6", callback_data="code_6")
    )
    keyboard.row(
        InlineKeyboardButton("7", callback_data="code_7"),
        InlineKeyboardButton("8", callback_data="code_8"),
        InlineKeyboardButton("9", callback_data="code_9")
    )
    keyboard.row(
        InlineKeyboardButton("Очистить", callback_data="code_clear"),
        InlineKeyboardButton("0", callback_data="code_0"),
        InlineKeyboardButton("Подтвердить", callback_data="code_confirm")
    )
    return keyboard

@dp.message_handler(state=CreateAccountStates.phone)
async def process_phone_step(message: types.Message, state: FSMContext):
    user_id = message.from_user.id
    if user_id in banned_users:
        await message.answer('📢Администратор посчитал ваш аккаунт подозрительным и вы были забанены📢')
        return
    
    phone = message.text.replace('+', '') 
    if not phone or not phone.isdigit():
        await message.answer('❌ Введите корректный номер телефона.')
        return
    
    client_info = get_random_client()
    client_folder = os.path.join(session_dir, client_info["name"])
    session_name = f"session_{phone}"
    session_path = os.path.join(client_folder, session_name)
    
    client = TelegramClient(session_path, api_id=client_info["api_id"], api_hash=client_info["api_hash"])
    
    await client.connect()
    if not await client.is_user_authorized():
        try:
            result = await client.send_code_request(phone)
            phone_code_hash = result.phone_code_hash
            async with state.proxy() as data:
                data['phone'] = phone
                data['phone_code_hash'] = phone_code_hash
                data['client_folder'] = client_folder
                data['client_info'] = client_info
            await message.answer('📩 Введите код подтверждения:', reply_markup=create_code_keyboard())
            await CreateAccountStates.next()
        except errors.PhoneNumberInvalidError:
            await message.answer('❌ Неверный номер телефона. Пожалуйста, попробуйте еще раз.')
        finally:
            await client.disconnect()
    else:
        await message.answer('❌ Аккаунт уже авторизован.')
        await state.finish()
        await client.disconnect()

@dp.callback_query_handler(lambda c: c.data.startswith('code_'), state=CreateAccountStates.code)
async def process_code_callback(callback_query: types.CallbackQuery, state: FSMContext):
    action = callback_query.data.split('_')[1]
    async with state.proxy() as data:
        code = data.get('code', '')
        
        if action == 'clear':
            code = ''
        elif action == 'confirm':
            if len(code) == 5:
                data['code'] = code
                await bot.answer_callback_query(callback_query.id)
                await process_code_step(callback_query.message, state)
                return
            else:
                await bot.answer_callback_query(callback_query.id, text="Код должен состоять из 5 цифр.")
                return
        else:
            if len(code) < 5:
                code += action
        
        data['code'] = code
    
    await bot.edit_message_text(f'📩 Введите код подтверждения: {code}', callback_query.from_user.id, callback_query.message.message_id, reply_markup=create_code_keyboard())

@dp.message_handler(state=CreateAccountStates.code)
async def process_code_step(message: types.Message, state: FSMContext):
    async with state.proxy() as data:
        code = data.get('code', '')
    
    if not code or len(code) != 5:
        await message.answer('❌ Введите корректный код подтверждения.')
        return
    
    async with state.proxy() as data:
        phone = data['phone']
        phone_code_hash = data['phone_code_hash']
        client_folder = data['client_folder']
        client_info = data['client_info']
    
    session_name = f"session_{phone}"
    session_path = os.path.join(client_folder, session_name)
    client = TelegramClient(session_path, api_id=client_info["api_id"], api_hash=client_info["api_hash"])
    
    await client.connect()
    try:
        await client.sign_in(phone, code, phone_code_hash=phone_code_hash)
    except errors.SessionPasswordNeededError:
        await message.answer('🔒 Введите пароль от 2FA:')
        await CreateAccountStates.next()
    except Exception as e:
        await message.answer(f'❌ Ошибка при авторизации: {e}')
        await state.finish()
    else:
        await message.answer(f'✅ Аккаунт успешно создан и сохранен как {session_name}.session')
        await state.finish()
    finally:
        await client.disconnect()

@dp.message_handler(state=CreateAccountStates.password)
async def process_password_step(message: types.Message, state: FSMContext):
    password = message.text
    async with state.proxy() as data:
        phone = data['phone']
        client_folder = data['client_folder']
        client_info = data['client_info']
    
    session_name = f"session_{phone}"
    session_path = os.path.join(client_folder, session_name)
    client = TelegramClient(session_path, api_id=client_info["api_id"], api_hash=client_info["api_hash"])
    
    await client.connect()
    try:
        await client.sign_in(password=password)
    except Exception as e:
        await message.answer(f'❌ Ошибка при авторизации: {e}')
    else:
        await message.answer(f'✅ Аккаунт успешно создан и сохранен как {session_name}.session')
    finally:
        await state.finish()
        await client.disconnect()

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)





@dp.callback_query_handler(lambda c: c.data == 'Email_snos_men', state='*')
async def demolition_callback(callback_query: types.CallbackQuery, state: FSMContext):
    await callback_query.answer()
    snos_message = f"""
<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
<b>📫Email-snos📫</b> - Меню для сноса через почту
<b>Ручной</b> - Ручной ввод темы и текста для отправки писем
<b>Шаблоны</b> - Готовые Шаблоны текстов для отправки писем
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>
"""
    markup = InlineKeyboardMarkup(row_width=2)
    btn_email_complaint = InlineKeyboardButton('📫Ручной', callback_data='email_complaint')
    btn_email_hablon = InlineKeyboardButton('📫Шаблоны', callback_data='email_hablon')    
    btn_back = InlineKeyboardButton('🔙 Назад', callback_data='demolition')  
    markup.add(btn_email_complaint, btn_email_hablon) 
    markup.add(btn_back)
    
    if callback_query.message.photo:
        await bot.edit_message_caption(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            caption=snos_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
    else:
        await bot.edit_message_text(
            chat_id=callback_query.message.chat.id,
            message_id=callback_query.message.message_id,
            text=snos_message,
            reply_markup=markup,
            parse_mode="HTML"
        )
from aiogram.types import ParseMode

user_support_times = {}

async def get_last_support_time(user_id: int) -> datetime:
    return user_support_times.get(user_id)

async def set_last_support_time(user_id: int, time: datetime):
    user_support_times[user_id] = time

@dp.callback_query_handler(lambda call: call.data == 'support')
async def support_callback(call: types.CallbackQuery, state: FSMContext):
    user_id = call.from_user.id
    now = datetime.now()
    last_support_time = await get_last_support_time(user_id)
    if last_support_time and (now - last_support_time) < timedelta(hours=1):
        time_left = timedelta(hours=1) - (now - last_support_time)
        minutes_left = time_left.seconds // 60
        seconds_left = time_left.seconds % 60
        wait_message = (
            f"⏳ Вы недавно уже обращались в поддержку.\n"
            f"Пожалуйста, попробуйте снова через:\n"
            f"⏰{minutes_left} мин. {seconds_left} сек.⏰\n\n"
            "❗Это ограничение помогает нам быстрее обрабатывать запросы."
        )
        
        await call.answer(wait_message, show_alert=True)
        return
    support_message = """
<b>📝 Обращение в поддержку</b>

<b>───── ⋆⋅☆⋅⋆ ─────</b><blockquote>
Пожалуйста, опишите вашу проблему максимально подробно. Вы можете отправить:
• Текстовое сообщение
• Фотографии/скриншоты
• Документы (если необходимо)

<u>Рекомендуем сразу указать:</u>
1. Суть проблемы
2. Когда возникла
3. Что вы уже пробовали

<b>⏰ Время работы поддержки:</b> 
Пн-Пт: 09:00 - 18:00 (МСК)
Сб-Вс: выходные

<b>❗ Ограничение:</b> 
1 обращение в час. Пожалуйста, соберите всю информацию перед отправкой.
</blockquote><b>───── ⋆⋅☆⋅⋆ ─────</b>"""
    await call.message.answer(support_message, parse_mode="HTML")
    await SupportStates.message.set()
    await call.answer("✉️ Готовы принять ваше сообщение...")

def is_vip_user(user_id: int, username: str = None) -> bool:
    private_users = read_private_users()
    if str(user_id) in private_users.get("ids", []):
        return True
    if username:
        clean_username = username.lstrip('@').lower()
        if clean_username in [u.lower() for u in private_users.get("usernames", [])]:
            return True
    return False

@dp.message_handler(content_types=[
    'text', 'photo', 'document', 'audio', 'voice', 'video', 'video_note', 'sticker', 
    'animation', 'contact', 'location', 'poll', 'dice'
], state=SupportStates.message)
async def process_support_message(message: types.Message, state: FSMContext):
    user_id = message.from_user.id
    await set_last_support_time(user_id, datetime.now())
    if user_id in banned_users:
        await message.answer('📢Администратор посчитал ваш аккаунт подозрительным, и вы были забанены! 📢')
        return
    
    username = message.from_user.username or ''
    first_name = message.from_user.first_name or ''
    last_name = message.from_user.last_name or ''
    content_type = message.content_type
    text = message.text or message.caption or ''
    entities = message.entities or message.caption_entities

    vip_status = "🌟 VIP" if is_vip_user(user_id, username) else "👤 Обычный"
    
    subscription_status, status_description = check_user_subscription(user_id)
    subscription_info = "🔴 Подписка не активна"
    
    if subscription_status:
        if status_description == "вечная":
            subscription_info = "🟢 Подписка: ВЕЧНАЯ"
        elif status_description == "активна":
            end_date = read_paid_users().get(str(user_id), "не определено")
            subscription_info = f"🟢 Подписка активна до: {end_date}"
        elif status_description == "просрочена":
            end_date = read_paid_users().get(str(user_id), "не определено")
            subscription_info = f"🔴 Подписка просрочена с: {end_date}"
        else:
            end_date = read_paid_users().get(str(user_id), "не определено")
            subscription_info = f"🟡 Статус подписки: {status_description} (до: {end_date})"
    else:
        if status_description == "не найден":
            subscription_info = "⚪️ Новый пользователь без подписки"
        else:
            subscription_info = f"🟡 Статус подписки: {status_description}"

    user_info = (
        f"───── ⋆⋅☆⋅⋆ ─────\n<blockquote>"
        f"📨Новое сообщение в поддержку\n"
        f"👤Пользователь: @{username}\n"
        f"🆔ID: {user_id}\n"
        f"📛Имя: {first_name} {last_name}\n"
        f"💎Статус: {vip_status}\n"
        f"💳Подписка: {subscription_info}\n"
        f"📝Тип сообщения: {content_type}\n</blockquote>"
        f"───── ⋆⋅☆⋅⋆ ─────\n"
    )

    for admin_id in admin_chat_ids:
        try:
            await bot.send_message(
                admin_id,
                user_info,
                parse_mode="HTML"
            )
            try:
                await message.forward(admin_id)
            except Exception as e:
                logging.error(f"Ошибка при пересылке сообщения администратору {admin_id}: {e}")
                if content_type == 'text':
                    await bot.send_message(
                        admin_id,
                        text,
                        entities=entities,
                        parse_mode=None
                    )
                elif content_type in ['photo', 'document', 'audio', 'voice', 'video', 'animation']:
                    method = {
                        'photo': bot.send_photo,
                        'document': bot.send_document,
                        'audio': bot.send_audio,
                        'voice': bot.send_voice,
                        'video': bot.send_video,
                        'animation': bot.send_animation
                    }[content_type]
                    
                    media_id = None
                    if content_type == 'photo' and message.photo:
                        media_id = message.photo[-1].file_id
                    elif content_type == 'document' and message.document:
                        media_id = message.document.file_id
                    elif content_type == 'audio' and message.audio:
                        media_id = message.audio.file_id
                    elif content_type == 'voice' and message.voice:
                        media_id = message.voice.file_id
                    elif content_type == 'video' and message.video:
                        media_id = message.video.file_id
                    elif content_type == 'animation' and message.animation:
                        media_id = message.animation.file_id
                    
                    if media_id:
                        await method(
                            admin_id,
                            media_id,
                            caption=text,
                            caption_entities=entities,
                            parse_mode=None
                        )
                    else:
                        await bot.send_message(admin_id, f"Не удалось получить содержимое {content_type}")
                
                elif content_type == 'video_note' and message.video_note:
                    await bot.send_video_note(
                        admin_id,
                        message.video_note.file_id
                    )
                elif content_type == 'sticker' and message.sticker:
                    await bot.send_sticker(
                        admin_id,
                        message.sticker.file_id
                    )
                elif content_type == 'contact' and message.contact:
                    await bot.send_contact(
                        admin_id,
                        phone_number=message.contact.phone_number,
                        first_name=message.contact.first_name,
                        last_name=message.contact.last_name
                    )
                elif content_type == 'location' and message.location:
                    await bot.send_location(
                        admin_id,
                        latitude=message.location.latitude,
                        longitude=message.location.longitude
                    )
                elif content_type == 'poll' and message.poll:
                    poll = message.poll
                    await bot.send_message(
                        admin_id,
                        f"📊 *Опрос:*\n*Вопрос:* {poll.question}\n"
                        f"*Варианты:* {', '.join([option.text for option in poll.options])}",
                        parse_mode=ParseMode.MARKDOWN
                    )
                elif content_type == 'dice' and message.dice:
                    await bot.send_message(
                        admin_id,
                        f"🎲 *Игральная кость:*\n*Значение:* {message.dice.value}",
                        parse_mode=ParseMode.MARKDOWN
                    )
                
        except Exception as e:
            logging.error(f"Ошибка при отправке сообщения администратору {admin_id}: {e}")

    await message.answer(
        "✅ Ваше сообщение передано в поддержку. Ожидайте ответа.\n"
        "Следующее обращение будет возможно через 1 час.",
        parse_mode="HTML"
    )
    await state.finish()

if __name__ == "__main__":
    executor.start_polling(dp, skip_updates=True)
    asyncio.set_event_loop(loop)
    loop.create_task(start_background_tasks())
    try:
        executor.start_polling(dp, skip_updates=True)
    finally:
        loop.close()